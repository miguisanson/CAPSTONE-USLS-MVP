import os
import re
import json
import tempfile
import unittest
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from unittest.mock import MagicMock, patch
from openpyxl import load_workbook

_DB_FILE = tempfile.NamedTemporaryFile(suffix=".sqlite3", delete=False)
_DB_FILE.close()
os.environ["DATABASE_URL"] = f"sqlite:///{_DB_FILE.name}"

from app import (  # noqa: E402
    AcademicTerm,
    AdviserDocumentApproval,
    AwolCase,
    Course,
    CourseOffering,
    CourseOfferingPlan,
    CurriculumOffering,
    CourseRecord,
    DefenseVerdict,
    DocumentCheck,
    Faculty,
    FacultyCoursePreference,
    FacultyWorkingHour,
    Form1Endorsement,
    GraduationEndorsement,
    MonitoringSheetUpload,
    MonitoringValidationIssue,
    PanelAssignment,
    PracticumRecord,
    Program,
    ADVISER_APPROVAL_DOCUMENTS,
    ANALYTICS_REPORTS,
    CourseworkReport,
    OnboardingReview,
    StudentCurriculumTag,
    StudyPlanDraft,
    analytics_payload,
    available_curriculum_versions,
    build_coursework_report,
    build_onboarding_review,
    generate_study_plan_draft,
    onboarding_checklist_state,
    tag_student_curriculum,
    canonical_owner_role,
    completion_attrition_report,
    residency_watchlist_report,
    RESEARCH_DEFENSE_RESULT_ITEMS,
    RESEARCH_DEFENSE_SCHEDULE_ITEMS,
    RESEARCH_GATE_DEFENSE_TYPES,
    ResearchCase,
    ResearchEvidenceFile,
    ResidencyEnrollment,
    ScheduleRequest,
    Student,
    StudentMonitoringFlag,
    SubjectEnrollment,
    StudentRequestAttachment,
    Task,
    TermEnrollment,
    TransactionLog,
    UserAccount,
    WorkflowMessage,
    WithdrawalApplication,
    REQUEST_UPLOAD_ROOT,
    MONITORING_UPLOAD_ROOT,
    UPLOAD_ROOT,
    app,
    awol_residency_roster_payload,
    compute_course_audit,
    curriculum_offerings_for_term,
    db,
    detected_research_progress,
    defense_availability_context,
    earliest_defense_date,
    enrollment_integrity_payload,
    enrollment_subject_states,
    ensure_authoritative_curricula,
    google_calendar_free_window_count,
    get_active_term,
    ensure_workflow_demo_students,
    student_current_course_year,
    student_priority,
    sync_automatic_awol_statuses,
    graduation_eligibility,
    graduation_candidate_payload,
    import_ac_monitoring,
    ensure_demo_accounts,
    panel_roles_for_student,
    practicum_roster_payload,
    practicum_eligibility,
    program_allows_practicum,
    required_documents_for_gate,
    research_requirement_presentation,
    resolve_source_flags_after_monitoring_upload,
    seed_database,
    submitted_request_students,
    task_dict,
    workflow_approvals_payload,
    generate_password_hash,
)


class BpmWorkflowSimulationTests(unittest.TestCase):
    @classmethod
    def tearDownClass(cls):
        try:
            with app.app_context():
                db.session.remove()
                db.engine.dispose()
            os.unlink(_DB_FILE.name)
        except FileNotFoundError:
            pass

    def setUp(self):
        app.config.update(TESTING=True)
        self.created_files = []
        with app.app_context():
            db.drop_all()
            db.create_all()
            self.program = Program(code="BPM", name="BPM Psychology Program", college="Graduate School", has_practicum=True)
            db.session.add(self.program)
            db.session.flush()
            self.course = Course(program_id=self.program.id, code="BPM-501", title="Completion Course", units=3, category="Major")
            db.session.add(self.course)
            self.student = Student(
                student_number="GS-2026-TEST",
                first_name="Workflow",
                last_name="Student",
                email="workflow@example.test",
                program_id=self.program.id,
                entry_year=2025,
                academic_year_entry="25-26",
                year_level="1",
                current_stage="Final Defense",
                standing="Active",
            )
            db.session.add(self.student)
            db.session.flush()
            self.staff = self._account("staff", "staff@example.test")
            self.academic = self._account("academic_coordinator", "academic@example.test")
            self.research = self._account("research_coordinator", "research@example.test")
            self.dean = self._account("dean", "dean@example.test")
            self.program_id = self.program.id
            self.course_id = self.course.id
            self.student_id = self.student.id
            self.staff_id = self.staff.id
            self.academic_id = self.academic.id
            self.research_id = self.research.id
            self.dean_id = self.dean.id
            db.session.commit()

    def tearDown(self):
        for path in self.created_files:
            path.unlink(missing_ok=True)

    def _account(self, role, email):
        account = UserAccount(
            email=email,
            full_name=f"Test {role.title()}",
            password_hash=generate_password_hash("test-password"),
            role=role,
            active=True,
        )
        db.session.add(account)
        db.session.flush()
        return account

    def _attachment(self, request_type, suffix):
        item = StudentRequestAttachment(
            student_id=self.student_id,
            request_type=request_type,
            original_name=f"{suffix}.pdf",
            stored_name=f"{request_type}-{suffix}.pdf",
        )
        db.session.add(item)
        db.session.flush()
        return item

    def _role_client(self, account_id, role):
        client = app.test_client()
        with client.session_transaction() as session:
            session["account_id"] = account_id
            session["role"] = role
        return client

    def _dean_client(self):
        return self._role_client(self.dean_id, "dean")

    def _staff_client(self):
        return self._role_client(self.staff_id, "staff")

    def _academic_client(self):
        return self._role_client(self.academic_id, "academic_coordinator")

    def _research_client(self):
        return self._role_client(self.research_id, "research_coordinator")

    def _transition(self, client, slug, payload, expected=200):
        response = client.post(f"/api/transactions/{slug}", json=payload)
        self.assertEqual(response.status_code, expected, response.get_json())
        return response

    def _complete_research_gate_defenses(self, student):
        faculty = Faculty(
            name=f"{student.last_name} Panel Chair",
            college="Graduate School",
            role="Faculty",
            specialization="Graduate Research",
            email=f"panel-{student.id}@example.test",
            eligible_roles="Panel Chair,Content Specialist,Method Specialist,External Panel",
        )
        db.session.add(faculty)
        db.session.flush()
        if not Form1Endorsement.query.filter_by(student_id=student.id).first():
            db.session.add(Form1Endorsement(
                student_id=student.id,
                coordinator_name="Test Academic Coordinator",
                signature_data="data:image/png;base64,test",
                status="Endorsed",
            ))
            db.session.flush()

        evidence_index = 0

        for gate, defense_type in RESEARCH_GATE_DEFENSE_TYPES.items():
            system_items = {
                "Recommended panel set",
                RESEARCH_DEFENSE_SCHEDULE_ITEMS[gate],
                RESEARCH_DEFENSE_RESULT_ITEMS[gate],
            }
            for item_name in required_documents_for_gate(gate):
                doc = DocumentCheck(
                    student_id=student.id,
                    gate=gate,
                    item_name=item_name,
                    status="Complete",
                    evidence_reference="Test completion fixture",
                )
                db.session.add(doc)
                db.session.flush()
                if item_name in system_items:
                    continue
                if item_name in {"Academic Coordinator endorsement/e-signature", "Ethics clearance status and date"}:
                    continue
                file_count = 3 if item_name == "Three concept papers" else 1
                for file_number in range(file_count):
                    evidence_index += 1
                    evidence = ResearchEvidenceFile(
                        student_id=student.id,
                        document_check_id=doc.id,
                        original_name=f"{item_name} {file_number + 1}.pdf",
                        stored_name=f"test-research-{student.id}-{evidence_index}.pdf",
                        mime_type="application/pdf",
                    )
                    db.session.add(evidence)
                    db.session.flush()
                    db.session.add(AdviserDocumentApproval(
                        student_id=student.id,
                        evidence_file_id=evidence.id,
                        faculty_id=faculty.id,
                        adviser_name=faculty.name,
                        student_name=student.name,
                        document_name=item_name,
                        signature_data="data:image/png;base64,test",
                        status="Signed",
                    ))

            chair_assignment = None
            for panel_role in panel_roles_for_student(student):
                assignment = PanelAssignment(
                    student_id=student.id,
                    faculty_id=faculty.id,
                    gate=gate,
                    panel_role=panel_role,
                    score=100,
                    eligibility_note="Test panel assignment",
                )
                db.session.add(assignment)
                db.session.flush()
                if panel_role == "Panel Chair":
                    chair_assignment = assignment

            schedule = ScheduleRequest(
                student_id=student.id,
                preferred_date=date(2026, 8, 1),
                defense_type=defense_type,
                mode="In person",
                venue="Graduate School Conference Room",
                status="Confirmed",
                matched_count=len(panel_roles_for_student(student)),
                required_forms_status="Complete",
            )
            db.session.add(schedule)
            db.session.flush()
            db.session.add(DefenseVerdict(
                student_id=student.id,
                schedule_request_id=schedule.id,
                panel_assignment_id=chair_assignment.id,
                faculty_id=faculty.id,
                gate=gate,
                defense_type=defense_type,
                research_title="Completed research gate test",
                chair_name=faculty.name,
                result="Passed",
                defense_date=schedule.preferred_date,
            ))

    def test_practicum_complete_and_incomplete_certificate_loop(self):
        with app.app_context():
            self.student = db.session.get(Student, self.student_id)
            moa = self._attachment("practicum", "moa")
            certificates = self._attachment("practicum", "certificates")
            record = PracticumRecord(
                student_id=self.student.id,
                moa_attachment_id=moa.id,
                moa_uploaded=True,
                moa_status="Uploaded",
                practicum_site="Partner Site",
                required_hours=200,
                completed_hours=120,
                certificate_attachment_id=certificates.id,
                document_status="Pending Review",
                certificate_count=1,
                status="MOA Submitted",
            )
            db.session.add(record)
            db.session.commit()

            staff = self._staff_client()
            academic = self._academic_client()
            self._transition(staff, "practicum", {
                "student_id": self.student.id,
                "status": "MOA Under Review",
                "workflow_comment": "Please review the signed MOA with the Academic Coordinator.",
            })
            self.assertEqual(record.status, "MOA Under Review")
            forward_message = WorkflowMessage.query.filter_by(
                transaction_slug="practicum",
                student_id=self.student_id,
                recipient_role="Academic Coordinator",
                action_type="forward",
                new_status="MOA Under Review",
            ).order_by(WorkflowMessage.id.desc()).first()
            self.assertIsNotNone(forward_message)
            self.assertEqual(forward_message.visibility, "internal")
            self.assertIn("Please review the signed MOA", forward_message.comment)
            audit = TransactionLog.query.filter_by(
                transaction_slug="practicum",
                student_id=self.student_id,
                new_status="MOA Under Review",
            ).order_by(TransactionLog.id.desc()).first()
            self.assertIn("Reviewer comment: Please review the signed MOA", audit.notes)
            self.assertTrue(Task.query.filter_by(student_id=self.student_id, owner_role="Academic Coordinator").filter(Task.title.contains("practicum MOA")).first())
            self._transition(staff, "practicum", {"student_id": self.student.id, "status": "Practicum In Progress"}, 400)
            self._transition(academic, "practicum", {"student_id": self.student.id, "status": "Practicum In Progress"})
            self.assertEqual(record.status, "Practicum In Progress")

            record.status = "Hours Incomplete"
            db.session.commit()
            self._transition(staff, "practicum", {"student_id": self.student.id, "status": "Documents Under Review"})
            self.assertEqual(record.status, "Documents Under Review")
            self._transition(academic, "practicum", {
                "student_id": self.student.id,
                "status": "Additional Certificates Requested",
                "return_reason": "Upload the signed completion certificate.",
            })
            self.assertEqual(record.status, "Additional Certificates Requested")

            record.completed_hours = 200
            record.status = "Documents Submitted"
            db.session.commit()
            self._transition(staff, "practicum", {"student_id": self.student.id, "status": "Documents Under Review"})
            self._transition(academic, "practicum", {"student_id": self.student.id, "status": "Completed", "document_status": "Verified"})
            self._transition(academic, "practicum", {"student_id": self.student.id, "status": "Report Sent to Dean", "document_status": "Verified"})
            self.assertEqual(record.status, "Report Sent to Dean")
            dean_item = next(item for item in workflow_approvals_payload()["pending"] if item["type"] == "practicum" and item["id"] == record.id)
            self.assertEqual(dean_item["timeline"][9]["label"], "Status Report Sent to Dean")
            self.assertEqual(dean_item["timeline"][9]["state"], "current")

            response = self._dean_client().post(f"/api/approvals/workflow/practicum/{record.id}/decide", json={"decision": "review"})
            self.assertEqual(response.status_code, 200)
            db.session.refresh(record)
            self.assertEqual(record.status, "Dean Reviewed")

    def test_withdrawal_approval_sequence_and_denial_branch(self):
        with app.app_context():
            self.student = db.session.get(Student, self.student_id)
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
            )
            other_course = Course(
                program_id=self.program_id,
                code="BPM-502",
                title="Unaffected Course",
                units=3,
                category="Major",
            )
            db.session.add_all([term, other_course])
            db.session.flush()
            selected_enrollment = SubjectEnrollment(
                student_id=self.student.id,
                course_id=self.course_id,
                term_id=term.id,
                status="Enrolled",
            )
            other_enrollment = SubjectEnrollment(
                student_id=self.student.id,
                course_id=other_course.id,
                term_id=term.id,
                status="Enrolled",
            )
            selected_record = CourseRecord(
                student_id=self.student.id,
                course_id=self.course_id,
                status="Enrolled",
                term_label=term.label,
                grade_status="No Grade",
            )
            db.session.add_all([selected_enrollment, other_enrollment, selected_record])
            db.session.add(CurriculumOffering(
                program_id=self.program_id,
                academic_year="2026-2027",
                semester="1st Semester",
                course_id=self.course_id,
                added_by="Withdrawal official-offering fixture",
            ))
            db.session.flush()
            request_file = self._attachment("withdrawal", "request")
            application = WithdrawalApplication(
                student_id=self.student.id,
                reason="Personal",
                effective_term=term.label,
                subject_enrollment_id=selected_enrollment.id,
                request_attachment_id=request_file.id,
                status="Submitted to GS Staff",
            )
            db.session.add(application)
            db.session.commit()
            self.assertFalse(any(item["id"] == application.id and item["type"] == "withdrawal" for item in workflow_approvals_payload()["pending"]))

            staff = self._staff_client()
            self._transition(staff, "withdrawal", {
                "student_id": self.student.id,
                "workflow_action": "forward_to_dean",
                "workflow_comment": "Recorded the request and forwarded the signed application to the Dean.",
            })
            dean_item = next(item for item in workflow_approvals_payload()["pending"] if item["id"] == application.id and item["type"] == "withdrawal")
            self.assertEqual(dean_item["workflow_status"], "Dean Review")

            return_response = self._dean_client().post(
                f"/api/approvals/workflow/withdrawal/{application.id}/decide",
                json={"decision": "return", "note": "This path is not in the withdrawal BPM."},
            )
            self.assertEqual(return_response.status_code, 400, return_response.get_json())
            response = self._dean_client().post(f"/api/approvals/workflow/withdrawal/{application.id}/decide", json={"decision": "approve"})
            self.assertEqual(response.status_code, 200)
            db.session.refresh(application)
            db.session.refresh(selected_enrollment)
            self.assertEqual(application.status, "Approved - Awaiting Subject Tag")
            self.assertEqual(application.registrar_status, "Pending Subject Tag")
            self.assertEqual(selected_enrollment.status, "Enrolled")

            blocked_report = staff.post(
                "/api/withdrawal/approved.xlsx",
                json={"application_ids": [application.id]},
            )
            self.assertEqual(blocked_report.status_code, 409, blocked_report.get_json())

            self._transition(staff, "withdrawal", {
                "student_id": self.student.id,
                "workflow_action": "tag_subject_withdrawn",
                "workflow_comment": "Tag the approved student as withdrawn from the selected subject.",
            })
            db.session.refresh(application)
            db.session.refresh(selected_enrollment)
            db.session.refresh(other_enrollment)
            db.session.refresh(selected_record)
            db.session.refresh(self.student)
            self.assertEqual(application.status, "Subject Tagged - Registrar Preparation")
            self.assertEqual(application.registrar_status, "Pending Excel Export")
            self.assertEqual(selected_enrollment.status, "Withdrawn")
            self.assertIn("no academic grade", selected_enrollment.status_note)
            self.assertEqual(other_enrollment.status, "Enrolled")
            self.assertEqual(selected_record.status, "Not Started")
            self.assertIsNone(selected_record.term_label)
            self.assertIsNone(selected_record.grade_value)
            self.assertEqual(selected_record.grade_status, "No Grade")
            self.assertEqual(self.student.standing, "Active")
            self.assertEqual(self.student.enrollment_tag, "Enrolled")

            enrollment_workspace = self._academic_client().get(
                f"/api/enrollment?program_id={self.program_id}"
                f"&term_id={term.id}&student_id={self.student.id}"
            )
            self.assertEqual(enrollment_workspace.status_code, 200, enrollment_workspace.get_json())
            official_subject = next(
                item
                for item in enrollment_workspace.get_json()["offered_subjects"]
                if item["course_id"] == self.course_id
            )
            self.assertEqual(official_subject["status_label"], "Withdrawn")
            self.assertEqual(official_subject["enrollment_state"], "requires_resolution")

            report = staff.post(
                "/api/withdrawal/approved.xlsx",
                json={"application_ids": [application.id]},
            )
            self.assertEqual(report.status_code, 200, report.get_json())
            self.assertEqual(
                report.content_type,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            workbook = load_workbook(BytesIO(report.data))
            sheet = workbook["Approved Withdrawals"]
            self.assertEqual(sheet["E4"].value, "BPM-501")
            self.assertEqual(sheet["O4"].value, "No academic record / no grade impact")
            db.session.refresh(application)
            self.assertEqual(application.status, "Exported - Ready to Send")
            self.assertEqual(selected_enrollment.status, "Withdrawn")

            self._transition(staff, "withdrawal", {"student_id": self.student.id, "workflow_action": "coordinator_follow_through"}, 400)
            handoff = staff.post(
                "/api/withdrawal/registrar-handoff",
                json={
                    "application_ids": [application.id],
                    "registrar_reference": "Test Registrar delivery",
                },
            )
            self.assertIn(handoff.status_code, {404, 405})
            db.session.refresh(application)
            db.session.refresh(selected_enrollment)
            db.session.refresh(other_enrollment)
            db.session.refresh(selected_record)
            db.session.refresh(self.student)
            self.assertEqual(application.status, "Exported - Ready to Send")
            self.assertEqual(application.registrar_status, "Exported - Ready to Send")
            self.assertEqual(selected_enrollment.status, "Withdrawn")
            self.assertEqual(other_enrollment.status, "Enrolled")
            self.assertEqual(selected_record.status, "Not Started")
            self.assertIsNone(selected_record.term_label)
            self.assertIsNone(selected_record.grade_value)
            self.assertEqual(selected_record.grade_status, "No Grade")
            self.assertEqual(self.student.standing, "Active")
            self.assertEqual(self.student.enrollment_tag, "Enrolled")

            denied_student = Student(
                student_number="BPM-002",
                first_name="Denied",
                last_name="Student",
                email="denied@example.test",
                program_id=self.program_id,
                entry_year=2025,
                current_stage="Coursework",
                standing="Active",
            )
            db.session.add(denied_student)
            db.session.flush()
            denied = WithdrawalApplication(student_id=denied_student.id, status="Dean Review", dean_decision="Pending")
            db.session.add(denied)
            db.session.commit()
            response = self._dean_client().post(
                f"/api/approvals/workflow/withdrawal/{denied.id}/decide",
                json={"decision": "deny", "note": "The request does not meet the approved withdrawal grounds."},
            )
            self.assertEqual(response.status_code, 200)
            db.session.refresh(denied_student)
            self.assertEqual(denied.status, "Denied")
            self.assertEqual(denied_student.standing, "Active")
            self.assertEqual(denied_student.enrollment_tag, "Enrolled")

    def test_graduation_return_resubmit_approve_and_export(self):
        with app.app_context():
            self.student = db.session.get(Student, self.student_id)
            db.session.add(CourseRecord(student_id=self.student.id, course_id=self.course_id, status="Completed"))
            db.session.add(ResearchCase(
                student_id=self.student.id,
                case_type="Thesis",
                title="Completed Research",
                current_gate="Completion Evidence",
                status="Verified Complete",
            ))
            for item in required_documents_for_gate("Completion Evidence"):
                db.session.add(DocumentCheck(student_id=self.student.id, gate="Completion Evidence", item_name=item, status="Complete"))
            db.session.add(PracticumRecord(
                student_id=self.student.id,
                practicum_site="Test Partner Organization",
                supervisor_name="Test Supervisor",
                required_hours=200,
                completed_hours=200,
                document_status="Verified",
                completion_status="Completed and accepted",
                status="Dean Reviewed",
            ))
            application_file = self._attachment("graduation", "signed-application")
            db.session.add(GraduationEndorsement(
                student_id=self.student.id,
                request_attachment_id=application_file.id,
                batch_name="Batch 1 July 2026",
                review_window="AY 2026-2027",
                endorsement_status="For Review",
            ))
            db.session.commit()

            # Graduation readiness now gates on coursework, research, and
            # accepted practicum completion for programs that require it.
            self.assertTrue(graduation_eligibility(self.student)["eligible"])
            staff = self._staff_client()
            academic = self._academic_client()
            research = self._research_client()
            self._transition(staff, "graduation", {
                "student_id": self.student.id,
                "review_window": "AY 2026-2027",
                "endorsement_status": "Coursework Review",
                "workflow_comment": "AC, please validate this candidate's completed coursework.",
            })
            coursework_message = WorkflowMessage.query.filter_by(
                transaction_slug="graduation",
                student_id=self.student_id,
                recipient_role="Academic Coordinator",
                action_type="forward",
                new_status="Coursework Review",
            ).order_by(WorkflowMessage.id.desc()).first()
            self.assertIsNotNone(coursework_message)
            self.assertEqual(coursework_message.visibility, "internal")
            self.assertIn("validate this candidate", coursework_message.comment)
            self.assertTrue(Task.query.filter_by(student_id=self.student_id, owner_role="Academic Coordinator").filter(Task.title.contains("coursework completion")).first())
            self._transition(staff, "graduation", {"student_id": self.student.id, "endorsement_status": "Research Review"}, 400)
            self._transition(academic, "graduation", {"student_id": self.student.id, "endorsement_status": "Research Review"})
            self.assertTrue(Task.query.filter_by(student_id=self.student_id, owner_role="Research Coordinator").filter(Task.title.contains("research completion")).first())
            self._transition(research, "graduation", {"student_id": self.student.id, "endorsement_status": "Eligibility Confirmed"})
            self._transition(staff, "graduation", {"student_id": self.student.id, "endorsement_status": "Endorsement Prepared"})
            self._transition(staff, "graduation", {"student_id": self.student.id, "endorsement_status": "Ready for Dean Review"})
            endorsement = GraduationEndorsement.query.filter_by(student_id=self.student.id).first()
            db.session.commit()
            dean_item = next(item for item in workflow_approvals_payload()["pending"] if item["type"] == "graduation" and item["id"] == endorsement.id)
            self.assertTrue(dean_item["eligibility"]["eligible"])

            response = self._dean_client().post(
                f"/api/approvals/workflow/graduation/{endorsement.id}/decide",
                json={"decision": "return", "note": "Correct the endorsement list before resubmitting."},
            )
            self.assertEqual(response.status_code, 200)
            db.session.refresh(endorsement)
            self.assertEqual(endorsement.endorsement_status, "Returned for Revision")
            self._transition(staff, "graduation", {"student_id": self.student.id, "endorsement_status": "Ready for Dean Review"})
            response = self._dean_client().post(f"/api/approvals/workflow/graduation/{endorsement.id}/decide", json={"decision": "approve"})
            self.assertEqual(response.status_code, 200)

            response = self._dean_client().post("/api/graduation/endorsed.csv", json={"endorsement_ids": [endorsement.id]})
            self.assertEqual(response.status_code, 200)
            db.session.refresh(endorsement)
            self.assertEqual(endorsement.endorsement_status, "Dean Approved")
            self.assertEqual(endorsement.registrar_status, "Exported - Ready to Send")
            # Export is terminal inside the portal; external sending is not tracked here.
            self.assertEqual(
                self._dean_client().post("/api/graduation/registrar-handoff").status_code,
                405,
            )

    def test_student_graduation_application_requires_eligibility_and_signed_pdf(self):
        with app.app_context():
            student_account = self._account("student", "graduation-applicant@example.test")
            student_account.student_id = self.student_id
            db.session.commit()
            student_client = self._role_client(student_account.id, "student")

            student_context = student_client.get("/api/student-portal/context").get_json()
            self.assertEqual(
                student_context["graduation_default_review_window"],
                student_context["graduation_review_windows"][0],
            )
            self.assertRegex(student_context["graduation_default_review_window"], r"^AY \d{4}-\d{4}$")
            staff_context = self._staff_client().get("/api/transactions/graduation/context").get_json()
            self.assertEqual(
                staff_context["graduation_review_windows"],
                student_context["graduation_review_windows"],
            )
            dean_context_response = self._dean_client().get("/api/transactions/graduation/context")
            self.assertEqual(dean_context_response.status_code, 200)
            self.assertEqual(
                dean_context_response.get_json()["graduation_review_windows"],
                student_context["graduation_review_windows"],
            )

            response = student_client.post(
                "/api/student-portal/requests/graduation",
                json={"review_window": "AY 2026-2027"},
            )
            self.assertEqual(response.status_code, 400)
            self.assertIn("Upload the signed graduation application", response.get_json()["error"])

            application_file = self._attachment("graduation", "student-signed-application")
            db.session.commit()
            response = student_client.post(
                "/api/student-portal/requests/graduation",
                json={"review_window": "AY 2026-2027", "attachment_id": application_file.id},
            )
            self.assertEqual(response.status_code, 400)
            self.assertIn("opens only after coursework", response.get_json()["error"])

            db.session.add(CourseRecord(student_id=self.student_id, course_id=self.course_id, status="Completed"))
            db.session.add(ResearchCase(
                student_id=self.student_id,
                case_type="Thesis",
                title="Application-ready research",
                current_gate="Completion Evidence",
                status="Verified Complete",
            ))
            for item in required_documents_for_gate("Completion Evidence"):
                db.session.add(DocumentCheck(
                    student_id=self.student_id,
                    gate="Completion Evidence",
                    item_name=item,
                    status="Complete",
                ))
            db.session.add(PracticumRecord(
                student_id=self.student_id,
                required_hours=200,
                completed_hours=200,
                document_status="Verified",
                completion_status="Completed and accepted",
                status="Dean Reviewed",
            ))
            db.session.commit()

            response = student_client.post(
                "/api/student-portal/requests/graduation",
                json={"review_window": "Custom review window", "attachment_id": application_file.id},
            )
            self.assertEqual(response.status_code, 400)
            self.assertIn("Choose a graduation school year", response.get_json()["error"])

            active_year = student_context["graduation_default_review_window"]
            active_start = int(re.search(r"\d{4}", active_year).group())
            closed_year = f"AY {active_start + 1}-{active_start + 2}"
            response = student_client.post(
                "/api/student-portal/requests/graduation",
                json={"review_window": closed_year, "attachment_id": application_file.id},
            )
            self.assertEqual(response.status_code, 400)
            self.assertIn(f"open only for {active_year}", response.get_json()["error"])

            response = student_client.post(
                "/api/student-portal/requests/graduation",
                json={"review_window": "AY 2026-2027", "attachment_id": application_file.id},
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            endorsement = GraduationEndorsement.query.filter_by(student_id=self.student_id).one()
            self.assertEqual(endorsement.endorsement_status, "For Review")
            self.assertEqual(endorsement.request_attachment_id, application_file.id)
            self.assertIsNone(endorsement.batch_name)

            return_response = self._staff_client().post(
                "/api/transactions/graduation/messages",
                json={
                    "student_id": self.student_id,
                    "action_type": "return",
                    "recipient_role": "Student",
                    "visibility": "student_visible",
                    "template": "Remarks",
                    "comment": "The name on the submitted application does not match the student record.",
                    "require_document_resubmission": True,
                },
            )
            self.assertEqual(return_response.status_code, 200, return_response.get_json())
            db.session.refresh(endorsement)
            self.assertEqual(endorsement.endorsement_status, "Returned for Clarification")
            self.assertIsNone(endorsement.request_attachment_id)
            return_message = WorkflowMessage.query.filter_by(
                transaction_slug="graduation",
                student_id=self.student_id,
                action_type="return",
                status="Open",
            ).one()
            self.assertTrue(return_message.requires_document_resubmission)
            self.assertEqual(return_message.recipient_role, "Student")
            self.assertEqual(return_message.visibility, "student_visible")

            returned_context = student_client.get("/api/student-portal/context").get_json()
            self.assertIsNone(returned_context["graduation_endorsement"]["request_attachment"])
            self.assertIn(
                application_file.id,
                [item["id"] for item in returned_context["graduation_endorsement"]["attachments"]],
            )
            visible_return = next(
                item for item in returned_context["workflow_messages"]
                if item["id"] == return_message.id
            )
            self.assertTrue(visible_return["requires_document_resubmission"])
            self.assertIn("does not match", visible_return["comment"])

            old_file_response = student_client.post(
                "/api/student-portal/requests/graduation",
                json={"review_window": active_year, "attachment_id": application_file.id},
            )
            self.assertEqual(old_file_response.status_code, 400)
            self.assertIn("Upload a new file", old_file_response.get_json()["error"])

            replacement_file = self._attachment("graduation", "corrected-student-application")
            db.session.commit()
            resubmit_response = student_client.post(
                "/api/student-portal/requests/graduation",
                json={"review_window": active_year, "attachment_id": replacement_file.id},
            )
            self.assertEqual(resubmit_response.status_code, 200, resubmit_response.get_json())
            db.session.refresh(endorsement)
            db.session.refresh(return_message)
            self.assertEqual(endorsement.endorsement_status, "For Review")
            self.assertEqual(endorsement.request_attachment_id, replacement_file.id)
            self.assertEqual(return_message.status, "Responded")

    def test_graduation_incomplete_coursework_and_research_are_routed_to_staff(self):
        with app.app_context():
            staff = self._staff_client()
            academic = self._academic_client()
            research = self._research_client()
            application_file = self._attachment("graduation", "incomplete-case-application")
            db.session.add(GraduationEndorsement(
                student_id=self.student_id,
                request_attachment_id=application_file.id,
                batch_name="Batch 3 July 2026",
                review_window="AY 2026-2027",
                endorsement_status="For Review",
            ))
            db.session.commit()

            self._transition(staff, "graduation", {"student_id": self.student_id, "endorsement_status": "Coursework Review"})
            self._transition(academic, "graduation", {"student_id": self.student_id, "endorsement_status": "Research Review"})
            endorsement = GraduationEndorsement.query.filter_by(student_id=self.student_id).first()
            self.assertEqual(endorsement.endorsement_status, "Coursework Incomplete")
            exception_item = next(
                row for row in graduation_candidate_payload()
                if row["student"]["id"] == self.student_id
            )
            self.assertFalse(exception_item["eligibility"]["eligible"])
            self.assertEqual(
                exception_item["endorsement"]["endorsement_status"],
                "Coursework Incomplete",
            )
            self.assertTrue(Task.query.filter_by(student_id=self.student_id, owner_role="Graduate School Staff").filter(Task.title.contains("missing graduation coursework")).first())
            self._transition(research, "graduation", {"student_id": self.student_id, "endorsement_status": "Eligibility Confirmed"}, 400)
            self._transition(staff, "graduation", {"student_id": self.student_id, "endorsement_status": "Not Eligible"}, 400)
            self._transition(staff, "graduation", {
                "student_id": self.student_id,
                "endorsement_status": "Not Eligible",
                "workflow_comment": "Complete the missing coursework before resubmitting your graduation application.",
            })
            not_eligible_notice = WorkflowMessage.query.filter_by(
                transaction_slug="graduation",
                student_id=self.student_id,
                recipient_role="Student",
                action_type="return",
                new_status="Not Eligible",
            ).order_by(WorkflowMessage.id.desc()).first()
            self.assertIsNotNone(not_eligible_notice)
            self.assertEqual(not_eligible_notice.visibility, "student_visible")
            self.assertIn("missing coursework", not_eligible_notice.comment)

            db.session.add(CourseRecord(student_id=self.student_id, course_id=self.course_id, status="Completed"))
            db.session.commit()
            self._transition(staff, "graduation", {"student_id": self.student_id, "endorsement_status": "Coursework Review"})
            self._transition(academic, "graduation", {"student_id": self.student_id, "endorsement_status": "Research Review"})
            self._transition(research, "graduation", {"student_id": self.student_id, "endorsement_status": "Eligibility Confirmed"})
            self.assertEqual(endorsement.endorsement_status, "Research Incomplete")
            self.assertTrue(Task.query.filter_by(student_id=self.student_id, owner_role="Graduate School Staff").filter(Task.title.contains("missing graduation research")).first())
            self._transition(staff, "graduation", {
                "student_id": self.student_id,
                "endorsement_status": "Not Eligible",
                "workflow_comment": "Finish the research gate requirements before resubmitting.",
            })
            self.assertEqual(endorsement.endorsement_status, "Not Eligible")

    def test_graduation_projects_research_workflow_progress_and_uses_it_for_roster_discovery(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            student.current_stage = "Proposal Defense"
            db.session.add(ResearchCase(
                student_id=student.id,
                case_type="Thesis",
                title="Final-stage research evidence",
                current_gate="Final Defense",
                status="Complete",
            ))
            db.session.commit()

            eligibility = graduation_eligibility(student)
            self.assertEqual(
                [stage["name"] for stage in eligibility["research_progress"]["stages"]],
                ["Title Defense", "Proposal Defense", "Final Defense"],
            )
            final_stage = eligibility["research_progress"]["stages"][-1]
            self.assertIn("panel", final_stage)
            self.assertIn("schedule", final_stage)
            self.assertIn("defense", final_stage)
            self.assertFalse(eligibility["research_progress"]["research_gates_complete"])

            self._complete_research_gate_defenses(student)
            db.session.commit()
            eligibility = graduation_eligibility(student)
            self.assertTrue(eligibility["research_progress"]["research_gates_complete"])
            self.assertFalse(eligibility["research_progress"]["completion_evidence"]["complete"])
            self.assertEqual(eligibility["research_status"], "Complete")
            self.assertFalse(eligibility["eligible"])

            roster_ids = {row["student"]["id"] for row in graduation_candidate_payload()}
            self.assertNotIn(student.id, roster_ids)

    def test_student_completion_evidence_uploads_feed_graduation_checklist(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            student.comprehensive_exam_status = "Passed"
            db.session.add(CourseRecord(student_id=student.id, course_id=self.course_id, status="Completed"))
            db.session.add(ResearchCase(
                student_id=student.id,
                case_type="Thesis",
                title="Post-defense completion evidence test",
                current_gate="Final Defense",
                status="Complete",
            ))
            self._complete_research_gate_defenses(student)
            student_account = self._account("student", "student-completion-evidence@example.test")
            student_account.student_id = student.id
            db.session.commit()

            progress = detected_research_progress(student)
            self.assertEqual(progress["stage"], "Completion")
            self.assertEqual(progress["gate"], "Completion Evidence")
            self.assertEqual(progress["milestone"]["student_missing_count"], 5)

            student_client = self._role_client(student_account.id, "student")
            for item_name in required_documents_for_gate("Completion Evidence"):
                response = student_client.post(
                    "/api/student-portal/research-evidence/upload",
                    data={
                        "gate": "Completion Evidence",
                        "item_name": item_name,
                        "file": (BytesIO(b"%PDF-1.4\ncompletion evidence\n%%EOF\n"), f"{item_name}.pdf"),
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(response.status_code, 200, response.get_json())
                stored_name = ResearchEvidenceFile.query.order_by(ResearchEvidenceFile.id.desc()).first().stored_name
                self.created_files.append(REQUEST_UPLOAD_ROOT / stored_name)

            progress = detected_research_progress(student)
            self.assertEqual(progress["stage"], "Completion")
            self.assertTrue(progress["milestone"]["student_uploads_ready"])
            self.assertEqual(progress["milestone"]["student_missing_count"], 0)
            self.assertFalse(graduation_eligibility(student)["research_progress"]["completion_evidence"]["complete"])

            response = student_client.post("/api/student-portal/requests/research-gate", json={
                "student_id": student.id,
                "submitted_package": "Final post-defense completion package uploaded.",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            research_case = ResearchCase.query.filter_by(student_id=student.id).order_by(ResearchCase.opened_at.desc()).first()
            self.assertEqual(research_case.current_gate, "Completion Evidence")
            self.assertEqual(research_case.status, "Awaiting Review")

            self._transition(self._research_client(), "research-gate", {"student_id": student.id})
            db.session.refresh(research_case)
            eligibility = graduation_eligibility(student)
            self.assertEqual(research_case.current_gate, "Completion Evidence")
            self.assertEqual(research_case.status, "Verified Complete")
            self.assertTrue(eligibility["research_progress"]["completion_evidence"]["complete"])
            self.assertEqual(eligibility["research_progress"]["completion_evidence"]["requirements_complete"], 5)

    def test_usls_course_audit_grade_writes_are_disabled(self):
        with app.app_context():
            record = CourseRecord(
                student_id=self.student_id,
                course_id=self.course_id,
                term_label="AY 2026-2027 1st Semester",
                status="Enrolled",
                grade_status="No Grade",
                remarks="Imported from AIMS",
            )
            db.session.add(record)
            db.session.commit()
            academic = self._academic_client()

            response = academic.post("/api/course-audit/roster", json={
                "course_id": self.course_id,
                "term": "AY 2026-2027 1st Semester",
                "statuses": {str(self.student_id): "Completed"},
                "grades": {str(self.student_id): "1.25"},
            })
            self.assertEqual(response.status_code, 409, response.get_json())
            self.assertTrue(response.get_json()["read_only"])
            db.session.refresh(record)
            self.assertEqual(record.status, "Enrolled")
            self.assertEqual(record.grade_status, "No Grade")
            self.assertIsNone(record.grade_value)
            self.assertEqual(record.remarks, "Imported from AIMS")
            self.assertEqual(
                TransactionLog.query.filter_by(
                    transaction_slug="course-audit",
                    student_id=self.student_id,
                ).count(),
                0,
            )

    def test_recommended_year_delay_only_uses_authoritative_curricula(self):
        with app.app_context():
            course = db.session.get(Course, self.course_id)
            student = db.session.get(Student, self.student_id)
            course.recommended_term = "Year 1"
            student.entry_year = date.today().year - 3
            student.current_stage = "Coursework"
            db.session.commit()

            self.assertEqual(student_priority(student)["level"], "On Track")

            ensure_authoritative_curricula()
            mba = Program.query.filter_by(code="MBA").one()
            mba_student = Student(
                student_number="MBA-RISK-TEST",
                first_name="Planned",
                last_name="Student",
                email="mba-risk@example.test",
                program_id=mba.id,
                entry_year=date.today().year - 3,
                current_stage="Coursework",
                standing="Active",
            )
            db.session.add(mba_student)
            db.session.commit()

            assessment = student_priority(mba_student)
            self.assertEqual(assessment["level"], "Delayed")
            self.assertIn("suggested course year", " ".join(assessment["causes"]))

    def test_approved_loa_does_not_create_delay_risk(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            student.current_stage = "LOA"
            student.standing = "On Leave"
            student.enrollment_tag = "LOA"
            db.session.commit()

            assessment = student_priority(student)
            self.assertEqual(assessment["level"], "On Track")
            self.assertEqual(assessment["causes"], [])

    def test_failed_course_record_does_not_create_delay_risk(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            course = db.session.get(Course, self.course_id)
            student.current_stage = "Coursework"
            db.session.add(CourseRecord(
                student_id=student.id,
                course_id=course.id,
                status="Failed",
                grade_status="Failed",
            ))
            db.session.commit()

            assessment = student_priority(student)
            self.assertEqual(assessment["level"], "On Track")
            self.assertEqual(assessment["causes"], [])

    def test_course_adjustments_manual_selection_publishes_to_curriculum_planning(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
                is_active_planning_term=True,
            )
            db.session.add_all([
                term,
                CourseRecord(
                    student_id=self.student_id,
                    course_id=self.course_id,
                    status="Completed",
                    grade_status="Passed",
                ),
            ])
            db.session.commit()

            academic = self._academic_client()
            payload = academic.get(
                f"/api/course-adjustments?program_id={self.program_id}&term_id={term.id}"
            )
            self.assertEqual(payload.status_code, 200, payload.get_json())
            self.assertTrue(payload.get_json()["permissions"]["can_manage"])
            manual_row = next(
                row
                for row in payload.get_json()["demand"]
                if row["course"]["id"] == self.course_id
            )
            self.assertEqual(manual_row["demand_count"], 0)
            self.assertEqual(manual_row["selection_source"], "Manual")

            staff_payload = self._staff_client().get(
                f"/api/course-adjustments?program_id={self.program_id}&term_id={term.id}"
            )
            self.assertFalse(staff_payload.get_json()["permissions"]["can_manage"])

            draft = academic.post("/api/course-adjustments/plan", json={
                "program_id": self.program_id,
                "term_id": term.id,
                "action": "draft",
                "selections": [{
                    "course_id": self.course_id,
                    "offer": True,
                    "section_count": 1,
                    "notes": "Manual coordinator decision",
                }],
            })
            self.assertEqual(draft.status_code, 200, draft.get_json())
            plan_id = draft.get_json()["data"]["latest_plan"]["id"]

            submitted = academic.post("/api/course-adjustments/plan", json={
                "program_id": self.program_id,
                "term_id": term.id,
                "action": "submit",
            })
            self.assertEqual(submitted.status_code, 200, submitted.get_json())

            approved = self._dean_client().post(
                f"/api/approvals/{plan_id}/decide",
                json={"decision": "approve", "note": "Approved for publication"},
            )
            self.assertEqual(approved.status_code, 200, approved.get_json())

            published = academic.post("/api/course-adjustments/plan", json={
                "program_id": self.program_id,
                "term_id": term.id,
                "action": "publish",
            })
            self.assertEqual(published.status_code, 200, published.get_json())
            self.assertIn("Curriculum Planning", published.get_json()["message"])
            offering = CurriculumOffering.query.filter_by(
                program_id=self.program_id,
                academic_year="2026-2027",
                semester="1st Semester",
                course_id=self.course_id,
            ).first()
            self.assertIsNotNone(offering)
            self.assertIn("Course Adjustments plan", offering.added_by)

    def test_one_student_demand_lists_full_delay_impact(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
                is_active_planning_term=True,
            )
            db.session.add(term)
            db.session.flush()
            db.session.add(TermEnrollment(
                student_id=self.student_id,
                term_id=term.id,
                status="Enrolled",
                source_reference="Demand test",
            ))
            db.session.commit()

            response = self._academic_client().get(
                f"/api/course-adjustments?program_id={self.program_id}&term_id={term.id}"
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            row = next(item for item in response.get_json()["demand"] if item["course"]["id"] == self.course_id)
            self.assertEqual(row["demand_count"], 1)
            self.assertEqual(row["demand_status"], "meets_minimum")
            self.assertEqual(row["affected_count"], 1)
            self.assertEqual(row["delayed_count"], 1)
            self.assertEqual(row["affected_students"][0]["student_number"], "GS-2026-TEST")
            self.assertTrue(row["affected_students"][0]["will_be_delayed"])

    def test_subject_needs_report_uses_active_students_without_future_enrollment(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2027-2028 1st Semester",
                start_date=date(2027, 8, 1),
                end_date=date(2027, 12, 15),
                is_active_planning_term=True,
            )
            failed_course = Course(
                program_id=self.program_id,
                code="BPM-502",
                title="Failed Course",
                units=3,
                category="Major",
            )
            db.session.add_all([term, failed_course])
            db.session.flush()
            db.session.add(CourseRecord(
                student_id=self.student_id,
                course_id=failed_course.id,
                status="Failed",
                grade_status="Failed",
            ))
            db.session.commit()

            academic = self._academic_client()
            response = academic.get(
                "/api/course-adjustments/subject-needs-report"
                f"?program_id={self.program_id}&term_id={term.id}"
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            payload = response.get_json()
            rows = {row["course"]["code"]: row for row in payload["rows"]}
            self.assertEqual(payload["summary"]["students_reviewed"], 1)
            self.assertEqual(payload["summary"]["subjects_with_need"], 2)
            self.assertEqual(payload["summary"]["student_subject_needs"], 2)
            self.assertEqual(rows["BPM-501"]["not_taken_count"], 1)
            self.assertEqual(rows["BPM-502"]["not_taken_count"], 1)
            self.assertEqual(rows["BPM-502"]["students"][0]["reason"], "Failed")

            # The target term is still in the future and has no TermEnrollment rows.
            # Demand must nevertheless use the active monitoring population.
            demand = academic.get(
                f"/api/course-adjustments?program_id={self.program_id}&term_id={term.id}"
            )
            self.assertEqual(demand.status_code, 200, demand.get_json())
            missing_row = next(
                row for row in demand.get_json()["demand"]
                if row["course"]["id"] == self.course_id
            )
            self.assertEqual(missing_row["demand_count"], 1)

    def test_monitoring_sheet_rejects_manual_status_updates(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
                is_active_planning_term=True,
            )
            record = CourseRecord(
                student_id=self.student_id,
                course_id=self.course_id,
                status="Enrolled",
                term_label=term.label,
                evidence_reference="Official registrar import",
                grade_value="1.50",
                grade_status="Passed",
                remarks="Imported official remark",
            )
            enrollment = SubjectEnrollment(
                student_id=self.student_id,
                course_id=self.course_id,
                term=term,
                status="Enrolled",
                source_reference="Enrollment fixture",
            )
            db.session.add_all([term, record, enrollment])
            db.session.commit()

            academic = self._academic_client()
            response = academic.post("/api/monitoring/subject-status", json={
                "student_id": self.student_id,
                "course_id": self.course_id,
                "term_id": term.id,
                "status": "Taken",
                "expected_status": "Enrolled",
                "source_reference": "Faculty email dated July 23, 2026",
                "note": "Final requirements were reported complete.",
            })
            self.assertEqual(response.status_code, 409, response.get_json())
            self.assertTrue(response.get_json()["read_only"])

            db.session.refresh(record)
            db.session.refresh(enrollment)
            self.assertEqual(enrollment.status, "Enrolled")
            self.assertEqual(enrollment.source_reference, "Enrollment fixture")
            self.assertEqual(record.status, "Enrolled")
            self.assertEqual(record.term_label, term.label)
            self.assertEqual(record.grade_value, "1.50")
            self.assertEqual(record.grade_status, "Passed")
            self.assertEqual(record.evidence_reference, "Official registrar import")
            self.assertEqual(record.remarks, "Imported official remark")
            student = db.session.get(Student, self.student_id)
            self.assertEqual(len(compute_course_audit(student)["completed"]), 0)

            self.assertEqual(TransactionLog.query.filter_by(
                transaction_slug="monitoring-status", student_id=self.student_id,
            ).count(), 0)

            grid = academic.get(
                f"/api/monitoring/grid?program_id={self.program_id}&term_id={term.id}"
            )
            self.assertEqual(grid.status_code, 200, grid.get_json())
            row = next(item for item in grid.get_json()["students"] if item["id"] == self.student_id)
            self.assertEqual(row["cells"][str(self.course_id)], "Enrolled")
            self.assertEqual(row["operational_cells"][str(self.course_id)], "Enrolled")
            self.assertEqual(row["official_cells"][str(self.course_id)], "Enrolled")
            self.assertEqual(row["grades"][str(self.course_id)], "1.50")

    def test_monitoring_flags_are_typed_audited_and_source_uploads_can_resolve_them(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            staff = self._staff_client()
            invalid = staff.post(f"/api/students/{student.id}/flag-issue", json={
                "note": "The source row appears incomplete.",
            })
            self.assertEqual(invalid.status_code, 400, invalid.get_json())

            created = staff.post(f"/api/students/{student.id}/flag-issue", json={
                "category": "Source data discrepancy",
                "note": "The AY entry does not match the latest AIMS export.",
            })
            self.assertEqual(created.status_code, 200, created.get_json())
            flag_id = created.get_json()["flag"]["id"]
            item = db.session.get(StudentMonitoringFlag, flag_id)
            self.assertEqual(item.status, "Open")
            self.assertEqual(item.source, "Manual")

            upload = MonitoringSheetUpload(
                original_name="corrected-monitoring.xlsx",
                stored_name="corrected-monitoring-fixture.xlsx",
                program_code=student.program.code,
                row_count=1,
                subject_count=0,
                snapshot_json="{}",
                result_json="{}",
            )
            db.session.add(upload)
            db.session.flush()
            self.assertEqual(resolve_source_flags_after_monitoring_upload(student, upload), 1)
            db.session.commit()
            db.session.refresh(item)
            self.assertEqual(item.status, "Resolved")
            self.assertIn("corrected-monitoring.xlsx", item.resolution_note)

    def test_monitoring_class_list_is_synced_from_enrollment_and_read_only(self):
        with app.app_context():
            current_term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
                is_active_planning_term=True,
            )
            future_term = AcademicTerm(
                label="AY 2026-2027 2nd Semester",
                start_date=date(2027, 1, 10),
                end_date=date(2027, 5, 20),
            )
            other_course = Course(
                program_id=self.program_id,
                code="BPM-502",
                title="Unaffected Enrolled Course",
                units=3,
                category="Major",
            )
            db.session.add_all([current_term, future_term, other_course])
            db.session.flush()
            first_record = CourseRecord(
                student_id=self.student_id,
                course_id=self.course_id,
                status="Enrolled",
                term_label=current_term.label,
                grade_value="",
                grade_status="No Grade",
            )
            second_record = CourseRecord(
                student_id=self.student_id,
                course_id=other_course.id,
                status="Enrolled",
                term_label=current_term.label,
            )
            first_enrollment = SubjectEnrollment(
                student_id=self.student_id,
                course_id=self.course_id,
                term_id=current_term.id,
                status="Enrolled",
            )
            second_enrollment = SubjectEnrollment(
                student_id=self.student_id,
                course_id=other_course.id,
                term_id=current_term.id,
                status="Enrolled",
            )
            term_enrollment = TermEnrollment(
                student_id=self.student_id,
                term_id=current_term.id,
                status="Enrolled",
                source_reference="Drop scope fixture",
            )
            db.session.add_all([
                first_record,
                second_record,
                first_enrollment,
                second_enrollment,
                term_enrollment,
            ])
            db.session.add(CurriculumOffering(
                program_id=self.program_id,
                academic_year="2026-2027",
                semester="1st Semester",
                course_id=self.course_id,
                added_by="Class list fixture",
            ))
            db.session.add(CurriculumOffering(
                program_id=self.program_id,
                academic_year="2026-2027",
                semester="1st Semester",
                course_id=other_course.id,
                added_by="All class lists fixture",
            ))
            db.session.commit()
            academic = self._academic_client()
            response = academic.get(
                f"/api/monitoring/class-list?program_id={self.program_id}"
                f"&term_id={current_term.id}&course_id={self.course_id}"
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            payload = response.get_json()
            self.assertTrue(payload["read_only"])
            self.assertEqual(payload["selected_offering"]["course_code"], "BPM-501")
            self.assertEqual(payload["summary"]["students_in_selected_class"], 1)
            self.assertEqual(payload["students"][0]["student_number"], "GS-2026-TEST")
            all_response = academic.get(
                f"/api/monitoring/class-list?program_id={self.program_id}"
                f"&term_id={current_term.id}&view=all"
            )
            self.assertEqual(all_response.status_code, 200, all_response.get_json())
            all_payload = all_response.get_json()
            self.assertEqual(all_payload["view"], "all")
            self.assertEqual(len(all_payload["class_lists"]), 2)
            self.assertEqual(all_payload["summary"]["class_list_entries"], 2)
            self.assertEqual(all_payload["summary"]["unique_students"], 1)
            self.assertEqual(
                {item["offering"]["course_code"] for item in all_payload["class_lists"]},
                {"BPM-501", "BPM-502"},
            )
            blocked = academic.post("/api/monitoring/subject-status", json={
                "student_id": self.student_id,
                "course_id": self.course_id,
                "term_id": current_term.id,
                "status": "Dropped",
            })
            self.assertEqual(blocked.status_code, 409, blocked.get_json())
            return
            student = db.session.get(Student, self.student_id)
            course = db.session.get(Course, self.course_id)
            original_student_state = (
                student.standing,
                student.enrollment_tag,
                student.current_stage,
            )

            academic = self._academic_client()
            missing_date = academic.post("/api/monitoring/subject-status", json={
                "student_id": self.student_id,
                "course_id": self.course_id,
                "term_id": current_term.id,
                "status": "Dropped",
                "expected_status": "Enrolled",
                "source_reference": "GS email",
                "note": "Student requested a mid-semester drop.",
            })
            self.assertEqual(missing_date.status_code, 400, missing_date.get_json())

            response = academic.post("/api/monitoring/subject-status", json={
                "student_id": self.student_id,
                "course_id": self.course_id,
                "term_id": current_term.id,
                "status": "Dropped",
                "expected_status": "Enrolled",
                "effective_date": "2026-09-15",
                "source_reference": "GS email",
                "note": "Student requested a mid-semester drop.",
            })
            self.assertEqual(response.status_code, 200, response.get_json())

            db.session.refresh(first_enrollment)
            db.session.refresh(second_enrollment)
            db.session.refresh(term_enrollment)
            db.session.refresh(first_record)
            db.session.refresh(student)
            self.assertEqual(first_enrollment.status, "Dropped")
            self.assertEqual(first_enrollment.cancelled_at.date(), date(2026, 9, 15))
            self.assertEqual(second_enrollment.status, "Enrolled")
            self.assertEqual(term_enrollment.status, "Enrolled")
            self.assertEqual(first_record.status, "Enrolled")
            self.assertEqual(
                (student.standing, student.enrollment_tag, student.current_stage),
                original_student_state,
            )

            same_term_state = enrollment_subject_states(
                student, current_term, [course]
            )[self.course_id]
            future_state = enrollment_subject_states(
                student, future_term, [course]
            )[self.course_id]
            self.assertEqual(same_term_state["enrollment_state"], "requires_resolution")
            self.assertFalse(same_term_state["selectable"])
            self.assertEqual(future_state["enrollment_state"], "available")
            self.assertTrue(future_state["selectable"])

            report = academic.get(
                "/api/course-adjustments/subject-needs-report"
                f"?program_id={self.program_id}&term_id={future_term.id}"
            )
            self.assertEqual(report.status_code, 200, report.get_json())
            row = next(item for item in report.get_json()["rows"] if item["course"]["id"] == self.course_id)
            self.assertEqual(row["not_taken_count"], 1)
            self.assertEqual(row["students"][0]["reason"], "Dropped")

    def test_faculty_preference_is_advisory_until_coordinator_selects(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
                is_active_planning_term=True,
            )
            faculty = Faculty(
                name="Preferred Teacher",
                college="Graduate School",
                role="Faculty",
                specialization="Completion Course",
                email="preferred-teacher@example.test",
                active=True,
            )
            db.session.add_all([term, faculty])
            db.session.flush()
            db.session.add(FacultyCoursePreference(
                faculty_id=faculty.id, course_id=self.course_id, priority=1
            ))
            db.session.commit()

            draft = self._academic_client().post("/api/course-adjustments/plan", json={
                "program_id": self.program_id,
                "term_id": term.id,
                "action": "draft",
                "selections": [{"course_id": self.course_id, "offer": True, "section_count": 1}],
            })
            self.assertEqual(draft.status_code, 200, draft.get_json())
            offering = CourseOffering.query.filter_by(course_id=self.course_id).first()
            self.assertIsNone(offering.assigned_faculty_id)
            self.assertEqual(offering.assignment_status, "Unassigned")

            draft = self._academic_client().post("/api/course-adjustments/plan", json={
                "program_id": self.program_id,
                "term_id": term.id,
                "action": "draft",
                "selections": [{
                    "course_id": self.course_id,
                    "offer": True,
                    "section_count": 1,
                    "assigned_faculty_id": faculty.id,
                }],
            })
            self.assertEqual(draft.status_code, 200, draft.get_json())
            offering = CourseOffering.query.filter_by(course_id=self.course_id).first()
            self.assertEqual(offering.assigned_faculty_id, faculty.id)
            self.assertEqual(offering.assignment_status, "Coordinator Selected")
            self.assertLessEqual((offering.course.units or 3) * offering.section_count, 24)

            submitted = self._academic_client().post("/api/course-adjustments/plan", json={
                "program_id": self.program_id, "term_id": term.id, "action": "submit",
            })
            self.assertEqual(submitted.status_code, 200, submitted.get_json())
            db.session.refresh(offering)
            self.assertEqual(offering.assignment_status, "Pending Approval")
            plan = CourseOfferingPlan.query.get(offering.plan_id)
            approved = self._dean_client().post(
                f"/api/approvals/{plan.id}/decide", json={"decision": "approve"}
            )
            self.assertEqual(approved.status_code, 200, approved.get_json())
            db.session.refresh(offering)
            self.assertEqual(offering.assignment_status, "Approved")

    def test_faculty_cv_reference_is_grounded_and_links_to_full_demo_document(self):
        with app.app_context():
            faculty = Faculty(
                name="Demo CV Faculty",
                college="Graduate School",
                role="Faculty",
                specialization="Biology and research methods",
                email="demo-cv-faculty@example.test",
                active=True,
            )
            db.session.add(faculty)
            db.session.commit()
            faculty_id = faculty.id

            academic = self._academic_client()
            directory = academic.get("/api/faculty")
            self.assertEqual(directory.status_code, 200, directory.get_json())
            item = next(row for row in directory.get_json()["items"] if row["id"] == faculty_id)
            self.assertTrue(item["cv_profile"]["is_demo_sample"])
            self.assertTrue(item["cv_profile"]["source_url"].startswith("https://"))

            response = academic.post(
                f"/api/faculty/{faculty_id}/cv-rag",
                json={"question": "What teaching and laboratory experience supports a STEM assignment?"},
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            payload = response.get_json()
            self.assertEqual(payload["mode"], "local-retrieval-augmented-answer")
            self.assertGreaterEqual(len(payload["citations"]), 1)
            self.assertTrue(all(citation["source_url"] == payload["document"]["source_url"] for citation in payload["citations"]))
            self.assertIn("Academic Coordinator", payload["answer"])

    def test_student_curriculum_checklist_adds_without_removing_existing_enrollment(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
                is_active_planning_term=True,
            )
            db.session.add(term)
            db.session.flush()
            db.session.add(CurriculumOffering(
                program_id=self.program_id,
                academic_year="2026-2027",
                semester="1st Semester",
                course_id=self.course_id,
                added_by="Test",
            ))
            account = UserAccount(
                email="checklist-student@example.test",
                full_name="Checklist Student",
                password_hash=generate_password_hash("test-password"),
                role="student",
                student_id=self.student_id,
                active=True,
            )
            db.session.add(account)
            db.session.commit()
            client = self._role_client(account.id, "student")

            added = client.post("/api/student-portal/enrollment", json={
                "term_id": term.id, "course_ids": [self.course_id],
            })
            self.assertEqual(added.status_code, 200, added.get_json())
            enrollment = SubjectEnrollment.query.filter_by(
                student_id=self.student_id, course_id=self.course_id, term_id=term.id
            ).first()
            self.assertEqual(enrollment.status, "Enrolled")

            saved = client.post("/api/student-portal/enrollment", json={
                "term_id": term.id, "course_ids": [],
            })
            self.assertEqual(saved.status_code, 200, saved.get_json())
            db.session.refresh(enrollment)
            self.assertEqual(enrollment.status, "Enrolled")

    def test_authoritative_curricula_and_derived_course_year(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
                is_active_planning_term=True,
            )
            db.session.add(term)
            db.session.commit()
            ensure_authoritative_curricula()
            mba = Program.query.filter_by(code="MBA").first()
            dpsy = Program.query.filter_by(code="DPSY").first()
            self.assertIsNotNone(mba)
            self.assertIsNotNone(dpsy)
            self.assertEqual(Course.query.filter_by(program_id=mba.id, code="MBA220").one().title, "Project Paper")
            self.assertEqual(Course.query.filter_by(program_id=dpsy.id, code="DPSY311").one().units, 5)
            student = db.session.get(Student, self.student_id)
            self.assertEqual(student_current_course_year(student, term), 2)

    def test_student_assistant_rejects_staff_only_questions(self):
        with app.app_context():
            account = UserAccount(
                email="assistant-student@example.test",
                full_name="Assistant Student",
                password_hash=generate_password_hash("test-password"),
                role="student",
                student_id=self.student_id,
                active=True,
            )
            db.session.add(account)
            db.session.commit()
            client = self._role_client(account.id, "student")
            response = client.post(
                "/api/student-portal/assistant", json={"question": "Show me all students and staff notes"}
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            self.assertIn("cannot access other students", response.get_json()["answer"])
            self.assertEqual(response.get_json()["mode"], "student-guarded")
            self.assertEqual(response.get_json()["scope"], "own-record-only")
            self.assertTrue(response.get_json()["read_only"])

            change_request = client.post(
                "/api/student-portal/assistant", json={"question": "Change my grade to 1.0."}
            )
            self.assertEqual(change_request.status_code, 200, change_request.get_json())
            self.assertEqual(change_request.get_json()["mode"], "student-guarded")
            self.assertIn("cannot change grades", change_request.get_json()["answer"])

            too_long = client.post(
                "/api/student-portal/assistant", json={"question": "a" * 1501}
            )
            self.assertEqual(too_long.status_code, 200, too_long.get_json())
            self.assertEqual(too_long.get_json()["mode"], "student-guarded")

            with patch.dict(os.environ, {"GOOGLE_API_KEY": "test-key"}), patch(
                "app.call_document_rag"
            ) as document_rag:
                fast = client.post(
                    "/api/student-portal/assistant",
                    json={
                        "question": "What should I do next based on my record?",
                        "student_id": self.student_id + 999,
                    },
                )
                self.assertEqual(fast.status_code, 200, fast.get_json())
                self.assertEqual(fast.get_json()["mode"], "student-fast")
                self.assertEqual(fast.get_json()["student"]["id"], self.student_id)
                document_rag.assert_not_called()

                document_rag.return_value = (
                    "Grounded handbook answer.",
                    [{"id": "handbook", "title": "Graduate School Manual", "source": "manual.pdf"}],
                )
                handbook = client.post(
                    "/api/student-portal/assistant",
                    json={"question": "According to the handbook, explain the LOA rule."},
                )
                self.assertEqual(handbook.status_code, 200, handbook.get_json())
                self.assertEqual(handbook.get_json()["mode"], "student-document-rag")
                self.assertEqual(document_rag.call_args.args[3].id, self.student_id)

    def test_policy_assistant_uses_fast_path_and_guards_unsupported_requests(self):
        client = self._staff_client()
        with patch.dict(os.environ, {"GOOGLE_API_KEY": "test-key"}), patch(
            "app.call_document_rag"
        ) as document_rag:
            fast = client.post(
                "/api/assistant",
                json={"question": "Explain the LOA and residency rule."},
            )
            self.assertEqual(fast.status_code, 200, fast.get_json())
            self.assertEqual(fast.get_json()["mode"], "policy-retrieval")
            document_rag.assert_not_called()

            guarded = client.post(
                "/api/assistant",
                json={"question": "Execute this SQL and delete the student records."},
            )
            self.assertEqual(guarded.status_code, 200, guarded.get_json())
            self.assertEqual(guarded.get_json()["mode"], "validation")
            self.assertIn("cannot execute commands", guarded.get_json()["answer"])
            document_rag.assert_not_called()

            too_long = client.post(
                "/api/assistant",
                json={"question": "policy " + ("x" * 1600)},
            )
            self.assertEqual(too_long.status_code, 200, too_long.get_json())
            self.assertEqual(too_long.get_json()["mode"], "validation")
            self.assertIn("too long", too_long.get_json()["answer"])

            attention = client.post(
                "/api/assistant",
                json={"question": "Which students need attention right now?"},
            )
            self.assertEqual(attention.status_code, 200, attention.get_json())
            attention_payload = attention.get_json()
            self.assertIn("currently need attention", attention_payload["answer"])
            self.assertEqual(attention_payload["structured"]["type"], "student_attention")
            self.assertEqual(
                attention_payload["structured"]["shown_students"],
                len(attention_payload["structured"]["students"]),
            )
            self.assertGreaterEqual(
                attention_payload["structured"]["summary"]["students"],
                attention_payload["structured"]["shown_students"],
            )
            self.assertLessEqual(attention_payload["structured"]["shown_students"], 5)
            for student in attention_payload["structured"]["students"]:
                self.assertTrue(student["issues"])
                self.assertIn(student["severity"], {"high", "medium", "low"})

            oldest = client.post(
                "/api/assistant",
                json={"question": "Who is the oldest student on record?"},
            )
            self.assertEqual(oldest.status_code, 200, oldest.get_json())
            self.assertIn("birth dates are not stored", oldest.get_json()["answer"].lower())
            self.assertIn("earliest academic entry on record", oldest.get_json()["answer"].lower())

        with patch("app.generate_answer", side_effect=RuntimeError("unexpected failure")):
            safe_error = client.post(
                "/api/assistant",
                json={"question": "Summarize current student concerns."},
            )
            self.assertEqual(safe_error.status_code, 200, safe_error.get_json())
            self.assertIn("No records were changed", safe_error.get_json()["answer"])

    def test_policy_assistant_falls_back_when_document_rag_fails(self):
        client = self._staff_client()
        with patch.dict(os.environ, {"GOOGLE_API_KEY": "test-key"}), patch(
            "app.call_document_rag",
            side_effect=RuntimeError("provider unavailable"),
        ):
            response = client.post(
                "/api/assistant",
                json={"question": "According to the handbook, explain the LOA rule."},
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            self.assertEqual(response.get_json()["mode"], "offline-fallback")
            self.assertTrue(response.get_json()["answer"])

    def test_policy_assistant_routes_database_paraphrases_to_verified_reports(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            student.current_stage = "Coursework"
            student.comprehensive_exam_status = "Not Taken"
            db.session.add(CourseRecord(
                student_id=student.id,
                course_id=self.course_id,
                status="Completed",
            ))
            earlier = Student(
                student_number="GS-2018-EARLY",
                first_name="Earlier",
                last_name="Entry",
                email="earlier@example.test",
                program_id=self.program_id,
                entry_year=2018,
                academic_year_entry="18-19",
                year_level="8",
                current_stage="Coursework",
                standing="Active",
            )
            db.session.add(earlier)
            db.session.commit()

        client = self._staff_client()
        progress = client.post(
            "/api/assistant",
            json={
                "question": "What is preventing this student from progressing?",
                "student_id": self.student_id,
            },
        )
        self.assertEqual(progress.status_code, 200, progress.get_json())
        progress_payload = progress.get_json()
        self.assertEqual(progress_payload["mode"], "database-rules")
        self.assertEqual(progress_payload["structured"]["heading"], "Progress review")
        self.assertIn(
            "comprehensive_exam",
            {item["code"] for item in progress_payload["structured"]["findings"]},
        )
        self.assertIn("Comprehensive Exam", progress_payload["answer"])
        self.assertEqual(progress_payload["citations"], [])
        self.assertFalse(progress_payload["source"]["ai_used"])

        graduation = client.post(
            "/api/assistant",
            json={
                "question": "When will this student graduate?",
                "student_id": self.student_id,
            },
        )
        self.assertEqual(graduation.status_code, 200, graduation.get_json())
        graduation_payload = graduation.get_json()
        self.assertEqual(graduation_payload["mode"], "database-rules")
        self.assertFalse(graduation_payload["structured"]["exact_date_known"])
        self.assertIn("does not store an exact graduation", graduation_payload["answer"])
        self.assertTrue(graduation_payload["structured"]["findings"])

        earliest = client.post(
            "/api/assistant",
            json={"question": "Which academic record has the earliest entry year?"},
        )
        self.assertEqual(earliest.status_code, 200, earliest.get_json())
        earliest_payload = earliest.get_json()
        self.assertEqual(earliest_payload["mode"], "database-rules")
        self.assertEqual(earliest_payload["structured"]["heading"], "Earliest academic entry")
        self.assertIn("Earlier Entry", earliest_payload["answer"])
        self.assertIn("2018", earliest_payload["answer"])
        self.assertNotIn("Title Defense", earliest_payload["answer"])

    def test_policy_assistant_forecasts_next_semester_graduation_readiness(self):
        with app.app_context():
            db.session.add_all([
                AcademicTerm(
                    label="AY 2026-2027 1st Semester",
                    start_date=date(2026, 8, 1),
                    end_date=date(2026, 12, 20),
                    is_active_planning_term=True,
                ),
                AcademicTerm(
                    label="AY 2026-2027 2nd Semester",
                    start_date=date(2027, 1, 10),
                    end_date=date(2027, 5, 30),
                ),
            ])
            db.session.commit()

        ready = {
            "eligible": True,
            "needs_verification": False,
            "next_owner": "Graduate School Staff",
            "checklist": [],
        }
        with patch("app.graduation_eligibility", return_value=ready):
            response = self._staff_client().post(
                "/api/assistant",
                json={"question": "How many students would ideally graduate next semester?"},
            )

        self.assertEqual(response.status_code, 200, response.get_json())
        payload = response.get_json()
        self.assertEqual(payload["mode"], "database-rules")
        self.assertEqual(payload["structured"]["heading"], "Next-semester graduation forecast")
        facts = {item["label"]: item["value"] for item in payload["structured"]["facts"]}
        self.assertEqual(facts["Target semester"], "AY 2026-2027 2nd Semester")
        self.assertEqual(facts["Students assessed"], "1")
        self.assertEqual(facts["Eligible now"], "1")
        self.assertIn("readiness forecast", payload["answer"])
        self.assertEqual(payload["citations"], [])

    def test_document_rag_prefers_semantic_vectors_and_keeps_lexical_fallback(self):
        from app import _retrieve_rag_document_chunks

        chunks = [
            {
                "id": "document-1",
                "title": "Leave policy",
                "source": "manual.pdf, p. 1",
                "text": "leave absence filing procedure",
            },
            {
                "id": "document-2",
                "title": "Research continuation",
                "source": "manual.pdf, p. 2",
                "text": "continuing a paused graduate program",
            },
        ]
        with patch(
            "app._load_rag_vector_index",
            return_value={
                "document-1": [0.0, 1.0],
                "document-2": [1.0, 0.0],
            },
        ), patch(
            "app._gemini_embed_texts",
            return_value=[[1.0, 0.0]],
        ):
            semantic = _retrieve_rag_document_chunks(
                "How can I return after pausing my studies?",
                chunks,
                limit=1,
            )
        self.assertEqual(semantic[0]["id"], "document-2")
        self.assertEqual(semantic[0]["_retrieval"], "vector")

        with patch(
            "app._load_rag_vector_index",
            side_effect=RuntimeError("embedding provider unavailable"),
        ):
            lexical = _retrieve_rag_document_chunks(
                "leave filing procedure",
                chunks,
                limit=1,
            )
        self.assertEqual(lexical[0]["id"], "document-1")
        self.assertEqual(lexical[0]["_retrieval"], "lexical-fallback")

    def test_enrollment_syncs_ledger_profile_and_operational_monitoring(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
            )
            db.session.add(term)
            db.session.flush()
            db.session.add(CurriculumOffering(
                program_id=self.program_id,
                academic_year="2026-2027",
                semester="1st Semester",
                course_id=self.course_id,
                added_by="Test",
            ))
            db.session.commit()

            academic = self._academic_client()
            preview = academic.post("/api/enrollment/preview", json={
                "student_id": self.student_id,
                "term_id": term.id,
                "course_ids": [self.course_id],
            })
            self.assertEqual(preview.status_code, 200, preview.get_json())
            self.assertEqual(preview.get_json()["conflict_count"], 0)
            self.assertEqual(len(preview.get_json()["additions"]), 1)

            response = academic.post("/api/enrollment", json={
                "student_id": self.student_id,
                "term_id": term.id,
                "course_ids": [self.course_id],
                "resolutions": {},
                "source_reference": "Enrollment test",
                "confirmed": True,
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            self.assertEqual(response.get_json()["integrity"]["issue_count"], 0)
            enrollment = SubjectEnrollment.query.filter_by(
                student_id=self.student_id,
                course_id=self.course_id,
                term_id=term.id,
            ).first()
            self.assertIsNotNone(enrollment)
            self.assertEqual(enrollment.status, "Enrolled")
            record = CourseRecord.query.filter_by(
                student_id=self.student_id,
                course_id=self.course_id,
            ).first()
            self.assertEqual(record.status, "Enrolled")
            self.assertEqual(record.term_label, term.label)
            semester = TermEnrollment.query.filter_by(
                student_id=self.student_id,
                term_id=term.id,
            ).first()
            self.assertEqual(semester.status, "Enrolled")

            profile = academic.get(f"/api/students/{self.student_id}")
            self.assertEqual(profile.status_code, 200, profile.get_json())
            self.assertEqual(
                profile.get_json()["subject_enrollments"][0]["course_code"],
                "BPM-501",
            )
            monitoring = academic.get(
                f"/api/monitoring/grid?program_id={self.program_id}&term_id={term.id}"
            )
            self.assertEqual(monitoring.status_code, 200, monitoring.get_json())
            student_row = next(
                row for row in monitoring.get_json()["students"]
                if row["id"] == self.student_id
            )
            self.assertEqual(student_row["cells"][str(self.course_id)], "Enrolled")
            self.assertEqual(student_row["academic_year_entry"], "25-26")
            self.assertEqual(student_row["year_level"], "1")

            response = academic.post("/api/monitoring/subject-status", json={
                "student_id": self.student_id,
                "course_id": self.course_id,
                "term_id": term.id,
                "status": "Taken",
                "expected_status": "Enrolled",
                "source_reference": "End-of-semester AIMS status batch",
                "note": "Confirmed as taken for operational monitoring.",
            })
            self.assertEqual(response.status_code, 409, response.get_json())
            db.session.refresh(enrollment)
            db.session.refresh(record)
            self.assertEqual(enrollment.status, "Enrolled")
            self.assertEqual(record.status, "Enrolled")
            self.assertIsNone(record.grade_value)

    def test_only_academic_coordinator_can_mark_subject_dropped_and_withdrawal_uses_workflow(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
            )
            db.session.add(term)
            db.session.flush()
            enrollment = SubjectEnrollment(
                student_id=self.student_id,
                course_id=self.course_id,
                term_id=term.id,
                status="Enrolled",
                source_reference="Test enrollment",
            )
            record = CourseRecord(
                student_id=self.student_id,
                course_id=self.course_id,
                status="Enrolled",
                term_label=term.label,
                grade_value="1.50",
                grade_status="Passed",
            )
            db.session.add_all([enrollment, record])
            db.session.commit()

            payload = {
                "subject_enrollment_id": enrollment.id,
                "status": "Dropped",
                "effective_date": "2026-09-15",
                "note": "Student stopped attending after the official change period.",
            }
            staff_response = self._staff_client().patch(
                "/api/enrollment/subject-status",
                json=payload,
            )
            self.assertEqual(staff_response.status_code, 403, staff_response.get_json())

            academic_response = self._academic_client().patch(
                "/api/enrollment/subject-status",
                json=payload,
            )
            self.assertEqual(academic_response.status_code, 200, academic_response.get_json())
            db.session.refresh(enrollment)
            db.session.refresh(record)
            self.assertEqual(enrollment.status, "Dropped")
            self.assertEqual(enrollment.status_note, payload["note"])
            self.assertIsNotNone(enrollment.status_changed_at)
            self.assertEqual(record.status, "Dropped")
            self.assertEqual(record.grade_value, "1.50")
            self.assertEqual(record.grade_status, "Passed")

            enrollment.status = "Enrolled"
            db.session.commit()
            direct_withdrawal = self._academic_client().patch(
                "/api/enrollment/subject-status",
                json={
                    **payload,
                    "subject_enrollment_id": enrollment.id,
                    "status": "Withdrawn",
                },
            )
            self.assertEqual(direct_withdrawal.status_code, 400, direct_withdrawal.get_json())
            self.assertIn("Withdrawal workflow", direct_withdrawal.get_json()["error"])

    def test_student_can_submit_within_withdrawal_window_and_late_request_is_blocked(self):
        with app.app_context():
            account = UserAccount(
                email="student-withdrawal@example.test",
                full_name="Student Withdrawal",
                password_hash=generate_password_hash("test-password"),
                role="student",
                student_id=self.student_id,
                active=True,
            )
            term = AcademicTerm(
                label="Withdrawal Window Term",
                start_date=date.today() + timedelta(days=3),
                end_date=date.today() + timedelta(days=123),
            )
            db.session.add_all([account, term])
            db.session.flush()
            enrollment = SubjectEnrollment(
                student_id=self.student_id,
                course_id=self.course_id,
                term_id=term.id,
                status="Enrolled",
            )
            db.session.add(enrollment)
            db.session.flush()
            db.session.commit()
            response = self._role_client(account.id, "student").post(
                "/api/student-portal/requests/withdrawal",
                json={
                    "subject_enrollment_id": enrollment.id,
                    "reason": "Schedule adjustment before classes",
                },
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            application = WithdrawalApplication.query.filter_by(student_id=self.student_id).one()
            self.assertEqual(application.status, "Submitted to GS Staff")
            self.assertEqual(application.subject_enrollment_id, enrollment.id)
            self.assertIsNone(application.request_attachment_id)

            late_student = Student(
                student_number="GS-LATE-WD",
                first_name="Late",
                last_name="Window",
                email="late-window@example.test",
                program_id=self.program_id,
                entry_year=2025,
                current_stage="Coursework",
                standing="Active",
            )
            late_term = AcademicTerm(
                label="Closed Withdrawal Window",
                start_date=date.today() - timedelta(days=8),
                end_date=date.today() + timedelta(days=112),
            )
            db.session.add_all([late_student, late_term])
            db.session.flush()
            late_account = self._account("student", "late-student-withdrawal@example.test")
            late_account.student_id = late_student.id
            late_enrollment = SubjectEnrollment(
                student_id=late_student.id,
                course_id=self.course_id,
                term_id=late_term.id,
                status="Enrolled",
            )
            db.session.add(late_enrollment)
            db.session.flush()
            db.session.commit()
            late_response = self._role_client(late_account.id, "student").post(
                "/api/student-portal/requests/withdrawal",
                json={
                    "subject_enrollment_id": late_enrollment.id,
                    "reason": "Too late for the penalty-free window",
                },
            )
            self.assertEqual(late_response.status_code, 409, late_response.get_json())
            self.assertIn("first seven calendar days", late_response.get_json()["error"])

    def test_enrollment_workspace_uses_monitoring_status_for_subject_eligibility(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
            )
            enrolled_course = Course(
                program_id=self.program_id,
                code="BPM-502",
                title="Already Enrolled Course",
                units=3,
                category="Major",
            )
            missing_course = Course(
                program_id=self.program_id,
                code="BPM-503",
                title="Not Taken Course",
                units=3,
                category="Major",
            )
            db.session.add_all([term, enrolled_course, missing_course])
            db.session.flush()
            term_id = term.id
            completed_course_id = self.course_id
            enrolled_course_id = enrolled_course.id
            missing_course_id = missing_course.id
            db.session.add_all([
                CourseRecord(
                    student_id=self.student_id,
                    course_id=completed_course_id,
                    status="Completed",
                    evidence_reference="Monitoring import",
                ),
                CourseRecord(
                    student_id=self.student_id,
                    course_id=enrolled_course_id,
                    status="Enrolled",
                    evidence_reference="Monitoring import",
                ),
                CourseRecord(
                    student_id=self.student_id,
                    course_id=missing_course_id,
                    status="Missing",
                    evidence_reference="Monitoring import",
                ),
            ])
            for course_id in [
                completed_course_id,
                enrolled_course_id,
                missing_course_id,
            ]:
                db.session.add(CurriculumOffering(
                    program_id=self.program_id,
                    academic_year="2026-2027",
                    semester="1st Semester",
                    course_id=course_id,
                    added_by="Test",
                ))
            db.session.commit()

            academic = self._academic_client()
            workspace = academic.get(
                f"/api/enrollment?program_id={self.program_id}"
                f"&term_id={term_id}&student_id={self.student_id}"
            )
            self.assertEqual(workspace.status_code, 200, workspace.get_json())
            payload = workspace.get_json()
            subjects = {
                item["code"]: item for item in payload["curriculum_subjects"]
            }
            self.assertEqual(subjects["BPM-501"]["enrollment_state"], "completed")
            self.assertFalse(subjects["BPM-501"]["selectable"])
            self.assertEqual(
                subjects["BPM-502"]["enrollment_state"],
                "enrolled_current",
            )
            self.assertFalse(subjects["BPM-502"]["selectable"])
            self.assertIn(enrolled_course_id, payload["current_course_ids"])
            self.assertEqual(subjects["BPM-503"]["enrollment_state"], "available")
            self.assertTrue(subjects["BPM-503"]["selectable"])
            self.assertEqual(payload["subject_status_summary"], {
                "available": 1,
                "completed": 1,
                "enrolled": 1,
                "enrolled_other_term": 0,
                "requires_resolution": 0,
            })

            completed_preview = academic.post("/api/enrollment/preview", json={
                "student_id": self.student_id,
                "term_id": term_id,
                "course_ids": [completed_course_id],
            })
            self.assertEqual(completed_preview.status_code, 200)
            self.assertTrue(completed_preview.get_json()["has_blocking_conflicts"])
            self.assertEqual(
                completed_preview.get_json()["conflicts"][0]["kind"],
                "already_completed",
            )

            valid_ids = [enrolled_course_id, missing_course_id]
            preview = academic.post("/api/enrollment/preview", json={
                "student_id": self.student_id,
                "term_id": term_id,
                "course_ids": valid_ids,
            })
            self.assertEqual(preview.status_code, 200, preview.get_json())
            self.assertEqual(preview.get_json()["conflict_count"], 0)
            self.assertEqual(
                [item["course_id"] for item in preview.get_json()["additions"]],
                [missing_course_id],
            )
            self.assertEqual(preview.get_json()["unchanged_count"], 1)

            saved = academic.post("/api/enrollment", json={
                "student_id": self.student_id,
                "term_id": term_id,
                "course_ids": valid_ids,
                "resolutions": {},
                "source_reference": "Monitoring synchronization test",
                "confirmed": True,
            })
            self.assertEqual(saved.status_code, 200, saved.get_json())
            self.assertEqual(saved.get_json()["added"], 1)
            completed_record = CourseRecord.query.filter_by(
                student_id=self.student_id,
                course_id=completed_course_id,
            ).one()
            self.assertEqual(completed_record.status, "Completed")
            self.assertEqual(
                CourseRecord.query.filter_by(
                    student_id=self.student_id,
                    course_id=enrolled_course_id,
                ).one().status,
                "Enrolled",
            )
            self.assertEqual(
                CourseRecord.query.filter_by(
                    student_id=self.student_id,
                    course_id=missing_course_id,
                ).one().status,
                "Enrolled",
            )
            self.assertEqual(
                SubjectEnrollment.query.filter_by(
                    student_id=self.student_id,
                    term_id=term_id,
                ).count(),
                2,
            )

    def test_enrollment_requires_resolution_for_non_offered_subject(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 2nd Semester",
                start_date=date(2027, 1, 10),
                end_date=date(2027, 5, 20),
            )
            db.session.add(term)
            db.session.commit()
            academic = self._academic_client()

            preview = academic.post("/api/enrollment/preview", json={
                "student_id": self.student_id,
                "term_id": term.id,
                "course_ids": [self.course_id],
            })
            self.assertEqual(preview.status_code, 200, preview.get_json())
            conflict = next(
                item for item in preview.get_json()["conflicts"]
                if item["kind"] == "not_offered"
            )

            blocked = academic.post("/api/enrollment", json={
                "student_id": self.student_id,
                "term_id": term.id,
                "course_ids": [self.course_id],
                "resolutions": {},
                "confirmed": True,
            })
            self.assertEqual(blocked.status_code, 409, blocked.get_json())

            saved = academic.post("/api/enrollment", json={
                "student_id": self.student_id,
                "term_id": term.id,
                "course_ids": [self.course_id],
                "resolutions": {conflict["id"]: "enroll_override"},
                "source_reference": "Approved exception test",
                "confirmed": True,
            })
            self.assertEqual(saved.status_code, 200, saved.get_json())
            enrollment = SubjectEnrollment.query.filter_by(
                student_id=self.student_id,
                course_id=self.course_id,
                term_id=term.id,
            ).first()
            self.assertIn("enroll_override", enrollment.conflict_override)
            self.assertEqual(saved.get_json()["integrity"]["issue_count"], 0)

    def test_course_offering_setup_and_class_list_import_work_end_to_end(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
            )
            faculty = Faculty(
                name="Dr. Offering Test",
                college="Graduate School",
                role="Faculty",
                specialization="Completion Course",
                email="offering-test@example.test",
                active=True,
            )
            not_offered = Course(
                program_id=self.program_id,
                code="BPM-599",
                title="Not Offered Course",
                units=3,
                category="Major",
            )
            completed_student = Student(
                student_number="GS-2026-COMPLETED",
                first_name="Completed",
                last_name="Student",
                email="completed-student@example.test",
                program_id=self.program_id,
                entry_year=2025,
                standing="Active",
            )
            db.session.add_all([term, faculty, not_offered, completed_student])
            db.session.flush()
            db.session.add(CourseRecord(
                student_id=completed_student.id,
                course_id=self.course_id,
                status="Completed",
                grade_status="Passed",
            ))
            db.session.commit()

            academic = self._academic_client()
            created = academic.post("/api/course-offerings", json={
                "program_id": self.program_id,
                "term_id": term.id,
                "course_id": self.course_id,
                "faculty_id": faculty.id,
                "schedule": "Saturday 8:00-11:00",
                "section": "A",
            })
            self.assertEqual(created.status_code, 200, created.get_json())
            offering = CurriculumOffering.query.filter_by(
                program_id=self.program_id,
                course_id=self.course_id,
                academic_year="2026-2027",
                semester="1st Semester",
            ).one()
            self.assertEqual(offering.assigned_faculty_id, faculty.id)
            self.assertEqual(offering.schedule, "Saturday 8:00-11:00")
            self.assertEqual(offering.section, "A")

            updated = academic.patch(f"/api/course-offerings/{offering.id}", json={
                "schedule": "Saturday 9:00-12:00",
                "section": "B",
            })
            self.assertEqual(updated.status_code, 200, updated.get_json())
            self.assertEqual(updated.get_json()["offering"]["schedule"], "Saturday 9:00-12:00")
            self.assertEqual(updated.get_json()["offering"]["section"], "B")
            offering.assigned_faculty_id = None
            db.session.commit()

            staff_view = self._staff_client().get(
                f"/api/course-offerings?program_id={self.program_id}&term_id={term.id}"
            )
            self.assertEqual(staff_view.status_code, 200, staff_view.get_json())
            self.assertFalse(staff_view.get_json()["permissions"]["can_manage"])

            csv_body = (
                "Student ID,Subject Code,Faculty\n"
                "GS-2026-TEST,BPM-501,Dr. Offering Test\n"
                "GS-2026-COMPLETED,BPM-501,Dr. Offering Test\n"
                "GS-UNKNOWN,BPM-501,Dr. Offering Test\n"
                "GS-2026-TEST,BPM-599,Dr. Offering Test\n"
            ).encode("utf-8")
            preview = academic.post(
                "/api/enrollment/class-list-preview",
                data={
                    "term_id": str(term.id),
                    "file": (BytesIO(csv_body), "class-list.csv"),
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(preview.status_code, 200, preview.get_json())
            self.assertEqual(preview.get_json()["total_rows"], 4)
            self.assertEqual(preview.get_json()["ready_count"], 1)
            self.assertEqual(preview.get_json()["warning_count"], 0)
            self.assertEqual(preview.get_json()["error_count"], 3)
            self.assertEqual(preview.get_json()["rows"][0]["faculty"], faculty.name)
            imported = academic.post(
                "/api/enrollment/class-list-import",
                data={
                    "term_id": str(term.id),
                    "file": (BytesIO(csv_body), "class-list.csv"),
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(imported.status_code, 200, imported.get_json())
            self.assertEqual(imported.get_json()["tagged"], 1)
            self.assertEqual(imported.get_json()["not_found_count"], 1)
            self.assertEqual(imported.get_json()["not_offered_count"], 1)
            self.assertEqual(imported.get_json()["completed_subject_count"], 1)
            self.assertEqual(imported.get_json()["faculty_assignments"], 1)
            db.session.refresh(offering)
            self.assertEqual(offering.assigned_faculty_id, faculty.id)
            enrollment = SubjectEnrollment.query.filter_by(
                student_id=self.student_id,
                course_id=self.course_id,
                term_id=term.id,
            ).one()
            self.assertEqual(enrollment.status, "Enrolled")
            self.assertEqual(
                CourseRecord.query.filter_by(
                    student_id=self.student_id,
                    course_id=self.course_id,
                ).one().status,
                "Enrolled",
            )
            self.assertEqual(
                CourseRecord.query.filter_by(
                    student_id=completed_student.id,
                    course_id=self.course_id,
                ).one().status,
                "Completed",
            )
            self.assertIsNone(SubjectEnrollment.query.filter_by(
                student_id=completed_student.id,
                course_id=self.course_id,
                term_id=term.id,
            ).first())

            blocked_delete = academic.delete(f"/api/course-offerings/{offering.id}")
            self.assertEqual(blocked_delete.status_code, 400, blocked_delete.get_json())

    def test_academic_semester_removal_keeps_referenced_history(self):
        with app.app_context():
            unused = AcademicTerm(
                label="AY 2030-2031 1st Semester",
                start_date=date(2030, 8, 1),
                end_date=date(2030, 12, 15),
            )
            referenced = AcademicTerm(
                label="AY 2030-2031 2nd Semester",
                start_date=date(2031, 1, 10),
                end_date=date(2031, 5, 20),
            )
            db.session.add_all([unused, referenced])
            db.session.flush()
            db.session.add(TermEnrollment(
                student_id=self.student_id,
                term_id=referenced.id,
                status="Confirmed",
                source_reference="History test",
            ))
            db.session.commit()
            staff = self._staff_client()

            removed = staff.delete(
                f"/api/admin/terms/{unused.id}",
                json={"confirmed_label": unused.label},
            )
            self.assertEqual(removed.status_code, 200, removed.get_json())
            self.assertIsNone(db.session.get(AcademicTerm, unused.id))

            blocked = staff.delete(
                f"/api/admin/terms/{referenced.id}",
                json={"confirmed_label": referenced.label},
            )
            self.assertEqual(blocked.status_code, 409, blocked.get_json())
            self.assertIsNotNone(db.session.get(AcademicTerm, referenced.id))

    def test_eligible_loa_and_readmission_require_dean_decisions(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            db.session.add(TransactionLog(
                transaction_slug="leave-of-absence",
                student_id=student.id,
                actor_role="Student",
                source_reference="loa-request.pdf",
                result="LOA application submitted",
                next_owner="GS Staff",
                notes=(
                    "Student submitted a Leave of Absence application for staff eligibility review.\n"
                    "Requested period: AY 2026-2027 1st Semester to AY 2026-2027 2nd Semester.\n"
                    "Reason/remarks: Health leave."
                ),
            ))
            db.session.commit()

            response = self._staff_client().post("/api/transactions/leave-of-absence", json={
                "student_id": student.id,
                "application_reference": "loa-request.pdf",
                "effective_start": "AY 2026-2027 1st Semester",
                "effective_end": "AY 2026-2027 2nd Semester",
                "reason_remarks": "Health leave",
                "eligibility_status": "Eligible",
                "prior_loa_count": 0,
                "dean_action": "Approve",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(student)
            self.assertEqual(student.standing, "Active")
            self.assertEqual(student.enrollment_tag, "Enrolled")
            forwarded = TransactionLog.query.filter_by(
                transaction_slug="leave-of-absence",
                student_id=student.id,
                result="LOA request forwarded to Dean",
            ).first()
            self.assertIsNotNone(forwarded)
            self.assertEqual(forwarded.actor_user_id, self.staff_id)
            loa_rows = submitted_request_students("leave-of-absence")
            self.assertEqual(loa_rows[0]["status"], "Dean Review")
            self.assertEqual(loa_rows[0]["next_action_owner"], "Dean")
            self.assertTrue(any(
                item["id"] == forwarded.id and item["type"] == "leave-of-absence"
                for item in workflow_approvals_payload()["pending"]
            ))
            decided = self._dean_client().post(
                f"/api/approvals/workflow/leave-of-absence/{forwarded.id}/decide",
                json={"decision": "approve"},
            )
            self.assertEqual(decided.status_code, 200, decided.get_json())
            db.session.refresh(student)
            self.assertEqual(student.standing, "On Leave")
            self.assertEqual(student.enrollment_tag, "LOA")
            report = self._staff_client().get(
                f"/api/standing-changes/leave-of-absence/{student.id}/registrar-report"
            )
            self.assertEqual(report.status_code, 200)
            db.session.add(TransactionLog(
                transaction_slug="readmission",
                student_id=student.id,
                actor_role="Student",
                source_reference="readmission-request.pdf",
                result="Readmission request submitted",
                next_owner="GS Staff",
                notes=(
                    "Student submitted a readmission request for staff review.\n"
                    "Target return semester: AY 2027-2028 1st Semester.\n"
                    "Previous LOA period: AY 2026-2027 1st Semester to AY 2026-2027 2nd Semester."
                ),
            ))
            db.session.commit()

            response = self._staff_client().post("/api/transactions/readmission", json={
                "student_id": student.id,
                "application_reference": "readmission-request.pdf",
                "target_return_term": "AY 2027-2028 1st Semester",
                "previous_loa_period": "AY 2026-2027 1st Semester to AY 2026-2027 2nd Semester",
                "eligibility_status": "Eligible to Return",
                "readmission_items": [
                    "Structured return intention completed",
                    "Updated study plan confirmed",
                    "Program or adviser consultation completed",
                    "No pending accountability confirmed",
                ],
                "dean_action": "Deny",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(student)
            self.assertEqual(student.standing, "On Leave")
            self.assertEqual(student.enrollment_tag, "LOA")
            forwarded = TransactionLog.query.filter_by(
                transaction_slug="readmission",
                student_id=student.id,
                result="Readmission request forwarded to Dean",
            ).first()
            self.assertIsNotNone(forwarded)
            readmission_rows = submitted_request_students("readmission")
            self.assertEqual(readmission_rows[0]["status"], "Dean Review")
            self.assertEqual(readmission_rows[0]["next_action_owner"], "Dean")
            self.assertTrue(any(
                item["id"] == forwarded.id and item["type"] == "readmission"
                for item in workflow_approvals_payload()["pending"]
            ))
            decided = self._dean_client().post(
                f"/api/approvals/workflow/readmission/{forwarded.id}/decide",
                json={"decision": "approve"},
            )
            self.assertEqual(decided.status_code, 200, decided.get_json())
            db.session.refresh(student)
            self.assertEqual(student.standing, "Active")
            self.assertEqual(student.enrollment_tag, "Not Enrolled")
            report = self._staff_client().get(
                f"/api/standing-changes/readmission/{student.id}/registrar-report"
            )
            self.assertEqual(report.status_code, 200)
            self.assertIsNotNone(Task.query.filter_by(
                student_id=student.id,
                owner_role="Academic Coordinator",
                title="Review readmitted student study plan and enrollment",
            ).first())

    def test_structured_loa_and_readmission_need_no_rag_or_pdf(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            student.enrollment_tag = "Enrolled"
            student_account = self._account("student", "structured-standing@example.test")
            student_account.student_id = student.id
            past = AcademicTerm(
                label="Past Semester",
                start_date=date(date.today().year, 1, 1),
                end_date=date(date.today().year, 5, 31),
            )
            future_one = AcademicTerm(
                label="Future Semester One",
                start_date=date(date.today().year + 1, 1, 10),
                end_date=date(date.today().year + 1, 5, 31),
            )
            future_two = AcademicTerm(
                label="Future Semester Two",
                start_date=date(date.today().year + 1, 8, 1),
                end_date=date(date.today().year + 1, 12, 15),
            )
            db.session.add_all([past, future_one, future_two])
            db.session.commit()
            client = self._role_client(student_account.id, "student")

            for request_type in ("leave-of-absence", "readmission", "awol-return"):
                upload = client.post(
                    "/api/student-portal/request-attachments/upload",
                    data={
                        "request_type": request_type,
                        "file": (BytesIO(b"%PDF-1.4\nlegacy standing form\n%%EOF\n"), "legacy.pdf"),
                    },
                    content_type="multipart/form-data",
                )
                self.assertEqual(upload.status_code, 400, upload.get_json())

            rejected = client.post("/api/student-portal/requests/leave-of-absence", json={
                "effective_start": past.label,
                "effective_end": past.label,
                "reason_category": "Medical / health",
                "reason_remarks": "Recovery period.",
            })
            self.assertEqual(rejected.status_code, 400, rejected.get_json())

            submitted = client.post("/api/student-portal/requests/leave-of-absence", json={
                "effective_start": future_one.label,
                "effective_end": future_two.label,
                "reason_category": "Medical / health",
                "reason_remarks": "Recovery period.",
            })
            self.assertEqual(submitted.status_code, 200, submitted.get_json())
            loa_log = TransactionLog.query.filter_by(
                student_id=student.id,
                result="LOA application submitted",
            ).first()
            self.assertEqual(loa_log.source_reference, "Structured LOA portal form")
            self.assertIn("no RAG", loa_log.notes)

            withdrawn = client.post("/api/student-portal/requests/leave-of-absence/withdraw", json={})
            self.assertEqual(withdrawn.status_code, 200, withdrawn.get_json())
            self.assertIsNotNone(TransactionLog.query.filter_by(
                student_id=student.id,
                result="LOA application withdrawn by student",
            ).first())

            student.standing = "On Leave"
            student.current_stage = "LOA"
            student.enrollment_tag = "LOA"
            db.session.commit()
            readmission = client.post("/api/student-portal/requests/readmission", json={
                "target_return_term": future_one.label,
                "previous_loa_start": past.label,
                "previous_loa_end": past.label,
                "return_intent": "I am ready to resume my studies in the selected semester.",
                "readmission_items": [
                    "Structured return intention completed",
                    "Updated study plan confirmed",
                    "Program or adviser consultation completed",
                    "No pending accountability confirmed",
                ],
            })
            self.assertEqual(readmission.status_code, 200, readmission.get_json())
            readmission_log = TransactionLog.query.filter_by(
                student_id=student.id,
                result="Readmission request submitted",
            ).first()
            self.assertEqual(readmission_log.source_reference, "Structured readmission portal form")
            self.assertIn("no RAG", readmission_log.notes)

    def test_readmission_needs_review_still_goes_to_dean(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            student.current_stage = "LOA"
            student.standing = "On Leave"
            student.enrollment_tag = "LOA"
            db.session.add(TransactionLog(
                transaction_slug="readmission",
                student_id=student.id,
                actor_role="Student",
                source_reference="readmission-incomplete.pdf",
                result="Readmission request submitted",
                next_owner="GS Staff",
                notes=(
                    "Student submitted a readmission request for staff review.\n"
                    "Target return semester: AY 2027-2028 1st Semester.\n"
                    "Previous LOA period: AY 2026-2027 1st Semester to AY 2026-2027 2nd Semester."
                ),
            ))
            db.session.commit()

            response = self._staff_client().post("/api/transactions/readmission", json={
                "student_id": student.id,
                "application_reference": "readmission-incomplete.pdf",
                "target_return_term": "AY 2027-2028 1st Semester",
                "previous_loa_period": "AY 2026-2027 1st Semester to AY 2026-2027 2nd Semester",
                "eligibility_status": "Needs Review",
                "readmission_items": [
                    "Structured return intention completed",
                    "Updated study plan confirmed",
                    "Program or adviser consultation completed",
                ],
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(student)
            self.assertEqual(student.standing, "On Leave")
            self.assertEqual(student.enrollment_tag, "LOA")
            forwarded = TransactionLog.query.filter_by(
                transaction_slug="readmission",
                student_id=student.id,
                result="Readmission request forwarded to Dean",
            ).first()
            self.assertIsNotNone(forwarded)
            self.assertIsNone(TransactionLog.query.filter_by(
                transaction_slug="readmission",
                student_id=student.id,
                result="Readmission auto-approved by RAG",
            ).first())
            readmission_rows = submitted_request_students("readmission")
            self.assertEqual(readmission_rows[0]["status"], "Dean Review")
            self.assertEqual(readmission_rows[0]["next_action_owner"], "Dean")
            self.assertTrue(any(
                item["id"] == forwarded.id and item["type"] == "readmission"
                for item in workflow_approvals_payload()["pending"]
            ))

    def test_awol_return_is_policy_reviewed_and_decided_by_dean(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            student.entry_year = date.today().year - 2
            student.current_stage = "Coursework"
            student.standing = "Active"
            student.enrollment_tag = "Enrolled"
            student_account = self._account("student", "awol-return-student@example.test")
            student_account.student_id = student.id
            term = AcademicTerm(
                label="AWOL Source Semester",
                start_date=date(date.today().year, 1, 1),
                end_date=date(date.today().year, 5, 31),
                is_active_planning_term=True,
            )
            future_term = AcademicTerm(
                label="AWOL Return Semester",
                start_date=date(date.today().year + 1, 1, 10),
                end_date=date(date.today().year + 1, 5, 31),
            )
            db.session.add_all([term, future_term])
            db.session.flush()
            term_row = TermEnrollment(student_id=student.id, term_id=term.id)
            db.session.add(term_row)
            term_row.status = "Withdrawn"
            term_row.source_reference = "AIMS full-semester withdrawal"
            db.session.commit()

            self.assertEqual(sync_automatic_awol_statuses(commit=True), 1)
            db.session.refresh(student)
            self.assertEqual(student.enrollment_tag, "AWOL")
            item = AwolCase.query.filter_by(student_id=student.id).first()
            self.assertEqual(item.status, "AWOL Declared")
            self.assertEqual(item.detection_source, "Full-semester withdrawal without approved LOA")
            self.assertIsNotNone(StudentMonitoringFlag.query.filter_by(
                student_id=student.id,
                category="AWOL policy alert",
                status="Open",
            ).first())
            response = self._role_client(student_account.id, "student").post("/api/student-portal/requests/awol-return", json={
                "target_return_term": future_term.label,
                "last_enrolled_term": term.label,
                "return_intent": "I intend to resume enrollment in the selected semester.",
                "return_reason": "My circumstances have stabilized and I am ready to continue.",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(item)
            self.assertEqual(item.status, "Return Submitted")
            self.assertEqual(item.policy_classification, "Within Maximum Residence")

            response = self._staff_client().post("/api/transactions/awol", json={
                "student_id": student.id,
                "case_id": item.id,
                "workflow_action": "forward_return_to_dean",
                "staff_notes": "Written intent verified.",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(item)
            self.assertEqual(item.status, "Dean Review")
            pending = workflow_approvals_payload()["pending"]
            self.assertTrue(any(row["type"] == "awol-return" and row["id"] == item.id for row in pending))

            response = self._dean_client().post(
                f"/api/approvals/workflow/awol-return/{item.id}/decide",
                json={"decision": "approve", "note": "Return endorsed."},
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(item)
            db.session.refresh(student)
            self.assertEqual(item.status, "Return Approved")
            self.assertEqual(student.standing, "Active")
            self.assertEqual(student.enrollment_tag, "Not Enrolled")
            report = self._staff_client().get(f"/api/awol/{item.id}/registrar-report")
            self.assertEqual(report.status_code, 200)

    def test_awol_absolute_limit_requires_reenrollment_and_valid_residency_is_recorded(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            student.entry_year = date.today().year - 8
            student.current_stage = "Coursework"
            student.standing = "Active"
            student.enrollment_tag = "Enrolled"
            student_account = self._account("student", "awol-limit-student@example.test")
            student_account.student_id = student.id
            term = AcademicTerm(
                label="AWOL Limit Source Semester",
                start_date=date(date.today().year, 1, 1),
                end_date=date(date.today().year, 5, 31),
                is_active_planning_term=True,
            )
            future_term = AcademicTerm(
                label="AWOL Limit Return Semester",
                start_date=date(date.today().year + 1, 1, 10),
                end_date=date(date.today().year + 1, 5, 31),
            )
            db.session.add_all([term, future_term])
            db.session.flush()
            term_row = TermEnrollment(student_id=student.id, term_id=term.id)
            db.session.add(term_row)
            term_row.status = "Withdrawn"
            term_row.source_reference = "AIMS full-semester withdrawal"
            db.session.commit()

            self.assertEqual(sync_automatic_awol_statuses(commit=True), 1)
            response = self._role_client(student_account.id, "student").post("/api/student-portal/requests/awol-return", json={
                "target_return_term": future_term.label,
                "last_enrolled_term": term.label,
                "return_intent": "I intend to resume enrollment.",
                "return_reason": "I am ready for program re-evaluation and continued study.",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            item = AwolCase.query.filter_by(student_id=student.id).first()
            self.assertEqual(item.policy_classification, "Full Re-enrollment Required")
            self.assertTrue(item.full_reenrollment_required)
            self.assertEqual(self._staff_client().post("/api/transactions/awol", json={
                "student_id": student.id,
                "case_id": item.id,
                "workflow_action": "forward_return_to_dean",
            }).status_code, 200)
            response = self._dean_client().post(
                f"/api/approvals/workflow/awol-return/{item.id}/decide",
                json={"decision": "approve", "note": "Proceed with full course re-evaluation."},
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(item)
            db.session.refresh(student)
            self.assertEqual(item.status, "Re-enrollment Required")
            self.assertEqual(student.enrollment_tag, "AWOL")

            student.entry_year = date.today().year
            student.standing = "Active"
            student.enrollment_tag = "Enrolled"
            student.current_stage = "Comprehensive Exam"
            term = AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(date.today().year, 1, 1),
                end_date=date(date.today().year, 5, 31),
                is_active_planning_term=True,
            )
            db.session.add(term)
            db.session.commit()
            response = self._staff_client().post("/api/transactions/awol", json={
                "student_id": student.id,
                "workflow_action": "record_residency",
                "term_id": term.id,
                "residency_reason": "Comprehensive examination",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            residency = ResidencyEnrollment.query.filter_by(student_id=student.id, status="Active").first()
            db.session.refresh(student)
            self.assertIsNotNone(residency)
            enrollment = TermEnrollment.query.filter_by(student_id=student.id, term_id=term.id).first()
            self.assertEqual(enrollment.status, "Residency")
            self.assertEqual(student.enrollment_tag, "Residency")
            report = self._staff_client().get(f"/api/residency/{residency.id}/registrar-report")
            self.assertEqual(report.status_code, 200)

    def test_demo_backoffice_accounts_sign_in_as_distinct_roles(self):
        with app.app_context():
            ensure_demo_accounts()
            db.session.commit()
        expected = {
            "staff@usls.edu.ph": "staff",
            "academic@usls.edu.ph": "academic_coordinator",
            "research@usls.edu.ph": "research_coordinator",
        }
        for email, role in expected.items():
            response = app.test_client().post(
                "/api/auth/login",
                json={"role": "staff", "email": email, "password": "DemoPass123!"},
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            self.assertEqual(response.get_json()["user"]["role"], role)

    def test_workflow_demo_students_quick_login_and_central_reset(self):
        with app.app_context():
            seed_database()

        response = app.test_client().get("/api/auth/demo-students")
        self.assertEqual(response.status_code, 200, response.get_json())
        items = response.get_json()["items"]
        self.assertEqual(len(items), 19)
        self.assertEqual(sum(item["workflow"] == "research" for item in items), 1)
        self.assertEqual(sum(item["workflow"] == "practicum" for item in items), 2)
        self.assertEqual(sum(item["workflow"] == "graduation" for item in items), 2)
        self.assertEqual(sum(item["workflow"] == "withdrawal" for item in items), 2)
        self.assertEqual(sum(item["workflow"] == "enrollment" for item in items), 2)
        self.assertEqual(sum(item["workflow"] == "student-handoff" for item in items), 2)
        self.assertEqual(sum(item["workflow"] == "course-adjustments" for item in items), 2)
        self.assertEqual(sum(item["workflow"] == "leave-of-absence" for item in items), 2)
        self.assertEqual(sum(item["workflow"] == "readmission" for item in items), 2)
        self.assertEqual(sum(item["workflow"] == "awol" for item in items), 2)

        for item in items:
            student_client = app.test_client()
            login_response = student_client.post("/api/auth/login", json={
                "role": "student",
                "email": item["email"],
                "password": item["password"],
            })
            self.assertEqual(login_response.status_code, 200, login_response.get_json())
            self.assertEqual(login_response.get_json()["user"]["role"], "student")
            if item["workflow"] == "practicum":
                portal_response = student_client.get("/api/student-portal/context")
                self.assertEqual(portal_response.status_code, 200, portal_response.get_json())
                portal_eligibility = portal_response.get_json()["practicum_eligibility"]
                self.assertTrue(portal_eligibility["eligible"], item["name"])
                self.assertTrue(
                    all(check["passed"] is True for check in portal_eligibility["checklist"]),
                    item["name"],
                )
            elif item["workflow"] == "withdrawal":
                portal_response = student_client.get("/api/student-portal/context")
                self.assertEqual(portal_response.status_code, 200, portal_response.get_json())
                portal_student = portal_response.get_json()["student"]
                self.assertEqual(portal_student["standing"], "Active")
                self.assertEqual(portal_student["current_stage"], "Coursework")
                self.assertIsNone(portal_response.get_json()["withdrawal_application"])

        practicum_items = [item for item in items if item["workflow"] == "practicum"]
        self.assertEqual(
            {item["program_code"] for item in practicum_items},
            {"MAPSY", "MSGC"},
        )
        practicum_item = practicum_items[0]
        graduation_item = next(item for item in items if item["workflow"] == "graduation")
        withdrawal_item = next(item for item in items if item["workflow"] == "withdrawal")
        research_item = next(item for item in items if item["workflow"] == "research")
        with app.app_context():
            practicum_student = Student.query.filter_by(student_number=practicum_item["student_number"]).one()
            graduation_student = Student.query.filter_by(student_number=graduation_item["student_number"]).one()
            withdrawal_student = Student.query.filter_by(student_number=withdrawal_item["student_number"]).one()
            research_student = Student.query.filter_by(student_number=research_item["student_number"]).one()
            for item in practicum_items:
                demo_student = Student.query.filter_by(student_number=item["student_number"]).one()
                practicum_status = practicum_eligibility(demo_student)
                self.assertTrue(practicum_status["eligible"], demo_student.name)
                self.assertTrue(
                    all(check["passed"] is True for check in practicum_status["checklist"]),
                    demo_student.name,
                )
                self.assertGreaterEqual(practicum_status["basic_units_completed"], 6)
                self.assertGreaterEqual(practicum_status["major_units_completed"], 9)
                self.assertGreaterEqual(practicum_status["cognate_units_completed"], 6)
                self.assertGreaterEqual(practicum_status["total_units_completed"], 21)
                self.assertTrue(all(
                    record.course.program_id == demo_student.program_id
                    for record in CourseRecord.query.filter_by(student_id=demo_student.id).all()
                ))
                practicum_research = ResearchCase.query.filter_by(student_id=demo_student.id).one()
                self.assertEqual(practicum_research.current_gate, "Completion Evidence")
                self.assertEqual(practicum_research.status, "Verified Complete")

            self.assertTrue(graduation_eligibility(graduation_student)["eligible"])
            current_withdrawal_courses = [
                record.course.code
                for record in CourseRecord.query.filter_by(student_id=withdrawal_student.id, status="Current").all()
            ]
            self.assertEqual(current_withdrawal_courses, ["MAED-MAJ1"])
            self.assertEqual(withdrawal_student.standing, "Active")
            self.assertEqual(withdrawal_student.enrollment_tag, "Enrolled")
            self.assertEqual(research_student.name, "Miguel Yu")
            self.assertEqual(research_student.comprehensive_exam_status, "Passed")
            self.assertTrue(all(
                record.status == "Completed" and record.grade_status == "Passed"
                for record in CourseRecord.query.filter_by(student_id=research_student.id).all()
            ))

            # Test-only filenames keep this temporary database's reset from
            # deleting the repository's live seeded Miguel PDFs.
            for evidence in ResearchEvidenceFile.query.filter_by(student_id=research_student.id).all():
                evidence.stored_name = f"test-research-reset-{evidence.id}.pdf"

            attachment = StudentRequestAttachment(
                student_id=practicum_student.id,
                request_type="practicum",
                original_name="demo-certificate.pdf",
                stored_name="workflow-demo-reset-certificate.pdf",
            )
            db.session.add(attachment)
            db.session.flush()
            db.session.add(PracticumRecord(
                student_id=practicum_student.id,
                required_hours=200,
                completed_hours=40,
                certificate_attachment_id=attachment.id,
                status="Certificate Review",
            ))
            db.session.add(WorkflowMessage(
                transaction_slug="practicum",
                student_id=practicum_student.id,
                sender_role="student",
                sender_name=practicum_student.name,
                recipient_role="staff",
                template="Student remark",
                comment="Please review my practicum certificate.",
            ))
            db.session.add(TransactionLog(
                transaction_slug="practicum",
                student_id=practicum_student.id,
                actor_role="Student",
                result="Practicum certificate submitted",
            ))
            db.session.add(Task(
                student_id=practicum_student.id,
                title="Review practicum certificate",
                owner_role="Academic Coordinator",
                due_at=date.today(),
            ))
            db.session.add(GraduationEndorsement(
                student_id=graduation_student.id,
                endorsement_status="Ready for Dean Review",
            ))
            withdrawal_attachment = StudentRequestAttachment(
                student_id=withdrawal_student.id,
                request_type="withdrawal",
                original_name="demo-withdrawal-request.pdf",
                stored_name="workflow-demo-reset-withdrawal.pdf",
            )
            db.session.add(withdrawal_attachment)
            db.session.flush()
            withdrawal_application = WithdrawalApplication(
                student_id=withdrawal_student.id,
                reason="Withdrawal workflow demonstration",
                effective_term="AY 2026-2027 1st Semester",
                request_attachment_id=withdrawal_attachment.id,
                status="Approved - Follow-through",
                dean_decision="Approved",
                requirement_status="Not Applicable",
            )
            db.session.add(withdrawal_application)
            db.session.add(WorkflowMessage(
                transaction_slug="withdrawal",
                student_id=withdrawal_student.id,
                sender_role="Graduate School Staff",
                sender_name="Test Staff",
                recipient_role="Dean",
                template="Withdrawal request forwarded",
                comment="Please review the withdrawal application.",
            ))
            db.session.add(TransactionLog(
                transaction_slug="withdrawal",
                student_id=withdrawal_student.id,
                actor_role="Graduate School Staff",
                result="Withdrawal request forwarded to Dean",
            ))
            db.session.add(Task(
                student_id=withdrawal_student.id,
                title="Perform withdrawal follow-through actions",
                owner_role="Graduate School Staff",
                due_at=date.today(),
            ))
            db.session.commit()
            practicum_student_id = practicum_student.id
            graduation_student_id = graduation_student.id
            withdrawal_student_id = withdrawal_student.id
            research_student_id = research_student.id

        reset_response = app.test_client().post(
            f"/api/auth/demo-students/{research_item['key']}/reset",
            json={},
        )
        self.assertEqual(reset_response.status_code, 200, reset_response.get_json())
        self.assertGreater(reset_response.get_json()["records_removed"], 0)

        reset_response = app.test_client().post(
            f"/api/auth/demo-students/{practicum_item['key']}/reset",
            json={},
        )
        self.assertEqual(reset_response.status_code, 200, reset_response.get_json())
        self.assertEqual(reset_response.get_json()["records_removed"], 1)
        self.assertEqual(reset_response.get_json()["messages_removed"], 1)

        reset_response = app.test_client().post(
            f"/api/auth/demo-students/{withdrawal_item['key']}/reset",
            json={},
        )
        self.assertEqual(reset_response.status_code, 200, reset_response.get_json())
        self.assertEqual(reset_response.get_json()["records_removed"], 1)
        self.assertEqual(reset_response.get_json()["messages_removed"], 1)

        reset_response = app.test_client().post(
            f"/api/auth/demo-students/{graduation_item['key']}/reset",
            json={},
        )
        self.assertEqual(reset_response.status_code, 200, reset_response.get_json())
        self.assertEqual(reset_response.get_json()["records_removed"], 1)

        with app.app_context():
            practicum_student = db.session.get(Student, practicum_student_id)
            graduation_student = db.session.get(Student, graduation_student_id)
            withdrawal_student = db.session.get(Student, withdrawal_student_id)
            research_student = db.session.get(Student, research_student_id)
            self.assertIsNone(PracticumRecord.query.filter_by(student_id=practicum_student_id).first())
            self.assertIsNone(StudentRequestAttachment.query.filter_by(student_id=practicum_student_id, request_type="practicum").first())
            self.assertIsNone(WorkflowMessage.query.filter_by(student_id=practicum_student_id, transaction_slug="practicum").first())
            self.assertIsNone(TransactionLog.query.filter_by(student_id=practicum_student_id, transaction_slug="practicum").first())
            self.assertIsNone(Task.query.filter_by(student_id=practicum_student_id, title="Review practicum certificate").first())
            self.assertTrue(practicum_eligibility(practicum_student)["eligible"])
            self.assertEqual(practicum_student.current_stage, "Final Defense")

            self.assertIsNone(GraduationEndorsement.query.filter_by(student_id=graduation_student_id).first())
            self.assertIsNotNone(PracticumRecord.query.filter_by(student_id=graduation_student_id).first())
            self.assertTrue(graduation_eligibility(graduation_student)["eligible"])
            self.assertEqual(graduation_student.current_stage, "Completed")

            self.assertIsNone(WithdrawalApplication.query.filter_by(student_id=withdrawal_student_id).first())
            self.assertIsNone(StudentRequestAttachment.query.filter_by(student_id=withdrawal_student_id, request_type="withdrawal").first())
            self.assertIsNone(WorkflowMessage.query.filter_by(student_id=withdrawal_student_id, transaction_slug="withdrawal").first())
            self.assertIsNone(TransactionLog.query.filter_by(student_id=withdrawal_student_id, transaction_slug="withdrawal").first())
            self.assertIsNone(Task.query.filter_by(student_id=withdrawal_student_id, title="Perform withdrawal follow-through actions").first())
            self.assertEqual(withdrawal_student.standing, "Active")
            self.assertEqual(withdrawal_student.enrollment_tag, "Enrolled")
            self.assertEqual(withdrawal_student.current_stage, "Coursework")
            self.assertEqual(
                [record.course.code for record in CourseRecord.query.filter_by(student_id=withdrawal_student_id, status="Current").all()],
                ["MAED-MAJ1"],
            )

            research_progress = detected_research_progress(research_student)
            self.assertEqual(research_progress["stage"], "Title Defense")
            self.assertEqual(research_progress["status"], "Missing Requirements")
            self.assertEqual(research_student.comprehensive_exam_status, "Passed")
            self.assertEqual(research_student.current_stage, "Proposal Development")
            self.assertEqual(ResearchEvidenceFile.query.filter_by(student_id=research_student_id).count(), 0)
            self.assertEqual(PanelAssignment.query.filter_by(student_id=research_student_id).count(), 0)
            self.assertEqual(ScheduleRequest.query.filter_by(student_id=research_student_id).count(), 0)
            self.assertEqual(DefenseVerdict.query.filter_by(student_id=research_student_id).count(), 0)
            self.assertIsNone(Form1Endorsement.query.filter_by(student_id=research_student_id).first())
            self.assertTrue(all(
                record.status == "Completed" and record.grade_status == "Passed"
                for record in CourseRecord.query.filter_by(student_id=research_student_id).all()
            ))

    def test_existing_practicum_demo_students_migrate_to_configured_programs(self):
        with app.app_context():
            seed_database()
            maed = Program.query.filter_by(code="MAED").one()
            andrea = Student.query.filter_by(student_number="GS-2026-PRAC-01").one()
            paolo = Student.query.filter_by(student_number="GS-2026-PRAC-02").one()
            andrea.program_id = maed.id
            paolo.program_id = maed.id
            db.session.commit()

            ensure_workflow_demo_students()
            db.session.commit()
            db.session.refresh(andrea)
            db.session.refresh(paolo)

            self.assertEqual(andrea.program.code, "MAPSY")
            self.assertEqual(paolo.program.code, "MSGC")
            self.assertTrue(practicum_eligibility(andrea)["eligible"])
            self.assertTrue(practicum_eligibility(paolo)["eligible"])

    def test_defense_availability_uses_recurring_profile_hours_without_dated_rows(self):
        target_day = date(2026, 7, 28)  # Tuesday
        with app.app_context():
            faculty = []
            for index in range(5):
                member = Faculty(
                    name=f"Defense Availability Faculty {index + 1}",
                    college="Graduate School",
                    role="Panel Member",
                    specialization="Research methods",
                    active=True,
                )
                db.session.add(member)
                faculty.append(member)
            db.session.flush()
            participants = [
                {"faculty": member, "role": "Panel Member"}
                for member in faculty
            ]
            google_status = {
                member.id: {"configured": False, "busy": [], "error": None}
                for member in faculty
            }

            with patch("app.google_busy_by_faculty", return_value=google_status), patch("app.faculty_calendar_blocks", return_value=[]):
                availability = defense_availability_context(participants, target_day, target_day)

            self.assertEqual(availability["dates"], ["2026-07-28"])
            self.assertTrue(all(
                any(slot["date"] == "2026-07-28" and slot["start"] == "08:00" and slot["end"] == "17:00" for slot in item["slots"])
                for item in availability["participants"]
            ))
            self.assertTrue(any(
                slot["date"] == "2026-07-28" and slot["start"] == "13:00" and slot["end"] == "15:00" and slot["matched_count"] == 5
                for slot in availability["possible_slots"]
            ))

            # If one profile disables Tuesday, the payload must expose four of
            # five individual availabilities and no false full-panel overlap.
            db.session.add(FacultyWorkingHour(
                faculty_id=faculty[-1].id,
                weekday=1,
                start_time=time(8, 0),
                end_time=time(17, 0),
                enabled=False,
            ))
            db.session.flush()
            db.session.expire(faculty[-1], ["working_hours"])
            with patch("app.google_busy_by_faculty", return_value=google_status), patch("app.faculty_calendar_blocks", return_value=[]):
                partial = defense_availability_context(participants, target_day, target_day)
            available_count = sum(
                any(slot["date"] == "2026-07-28" and slot["start"] <= "13:00" and slot["end"] > "13:00" for slot in item["slots"])
                for item in partial["participants"]
            )
            self.assertEqual(available_count, 4)
            self.assertFalse(any(slot["start"] == "13:00" for slot in partial["possible_slots"]))

    def test_connected_google_calendar_replaces_dummy_profile_schedule(self):
        target_day = date(2026, 7, 28)  # Tuesday
        with app.app_context():
            connected = Faculty(
                name="Connected Calendar Faculty",
                college="Graduate School",
                role="Panel Member",
                specialization="Research methods",
                active=True,
                google_calendar_id="connected@example.test",
                google_calendar_access_token="faculty-access-token",
            )
            fallback = Faculty(
                name="Profile Schedule Faculty",
                college="Graduate School",
                role="Panel Member",
                specialization="Statistics",
                active=True,
            )
            db.session.add_all([connected, fallback])
            db.session.flush()
            # This disabled profile row would make the connected faculty
            # unavailable if local dummy hours were still being consulted.
            db.session.add(FacultyWorkingHour(
                faculty_id=connected.id,
                weekday=1,
                start_time=time(8, 0),
                end_time=time(17, 0),
                enabled=False,
            ))
            db.session.flush()
            db.session.expire(connected, ["working_hours"])
            participants = [
                {"faculty": connected, "role": "Panel Chair"},
                {"faculty": fallback, "role": "Panel Member"},
            ]
            local_tz = timezone(timedelta(hours=8))
            google_status = {
                connected.id: {
                    "configured": True,
                    "busy": [{
                        "start": datetime.combine(target_day, time(10, 0)).replace(tzinfo=local_tz),
                        "end": datetime.combine(target_day, time(12, 0)).replace(tzinfo=local_tz),
                    }],
                    "error": None,
                },
                fallback.id: {"configured": False, "busy": [], "error": None},
            }

            with patch("app.google_busy_by_faculty", return_value=google_status), patch("app.faculty_calendar_blocks", return_value=[]):
                availability = defense_availability_context(participants, target_day, target_day)

            connected_payload = next(
                item for item in availability["participants"]
                if item["faculty_id"] == connected.id
            )
            self.assertEqual(connected_payload["availability_source"], "google_calendar")
            self.assertEqual(connected_payload["working_hours"], [])
            self.assertEqual(
                [(slot["start"], slot["end"], slot["source"]) for slot in connected_payload["slots"]],
                [("08:00", "18:00", "google_calendar")],
            )
            self.assertFalse(any(slot["start"] == "10:00" for slot in availability["possible_slots"]))
            self.assertTrue(any(slot["start"] == "12:00" for slot in availability["possible_slots"]))

    def test_connected_google_calendar_fails_closed_when_freebusy_is_unavailable(self):
        target_day = date(2026, 7, 28)
        with app.app_context():
            faculty = Faculty(
                name="Unavailable Google Calendar Faculty",
                college="Graduate School",
                role="Panel Member",
                specialization="Research methods",
                active=True,
                google_calendar_id="calendar@example.test",
                google_calendar_access_token="faculty-access-token",
            )
            db.session.add(faculty)
            db.session.flush()
            participants = [{"faculty": faculty, "role": "Panel Member"}]
            google_status = {
                faculty.id: {
                    "configured": True,
                    "busy": [],
                    "error": "Google Calendar unavailable",
                },
            }
            with patch("app.google_busy_by_faculty", return_value=google_status):
                availability = defense_availability_context(participants, target_day, target_day)

            self.assertEqual(availability["possible_slots"], [])
            self.assertEqual(availability["participants"][0]["slots"], [])
            self.assertIn("no times will be suggested", availability["participants"][0]["calendar_status"])

    def test_panel_matching_google_availability_counts_free_days_and_fails_closed(self):
        start_day = date(2026, 7, 27)  # Monday
        local_tz = timezone(timedelta(hours=8))
        result = {
            "configured": True,
            "error": None,
            "busy": [{
                "start": datetime.combine(start_day, time(8, 0)).replace(tzinfo=local_tz),
                "end": datetime.combine(start_day, time(18, 0)).replace(tzinfo=local_tz),
            }],
        }
        self.assertEqual(
            google_calendar_free_window_count(result, start_day, start_day + timedelta(days=2)),
            2,
        )
        result["error"] = "Google Calendar unavailable"
        self.assertEqual(
            google_calendar_free_window_count(result, start_day, start_day + timedelta(days=2)),
            0,
        )

    def test_faculty_can_start_google_calendar_oauth_for_only_their_account(self):
        with app.app_context():
            faculty = Faculty(
                name="OAuth Faculty",
                college="Graduate School",
                role="Panel Member",
                specialization="Research methods",
                email="oauth-faculty@example.test",
                active=True,
            )
            db.session.add(faculty)
            db.session.flush()
            account = self._account("faculty", "oauth-login@example.test")
            account.faculty_id = faculty.id
            db.session.commit()
            account_id = account.id
            faculty_id = faculty.id

        client = self._role_client(account_id, "faculty")
        with patch.dict(os.environ, {
            "GOOGLE_CALENDAR_CLIENT_ID": "client-id.apps.googleusercontent.com",
            "GOOGLE_CALENDAR_CLIENT_SECRET": "client-secret",
            "GOOGLE_CALENDAR_REDIRECT_URI": "http://localhost/api/faculty-portal/google-calendar/callback",
        }):
            response = client.get("/api/faculty-portal/google-calendar/authorization")

        self.assertEqual(response.status_code, 200, response.get_json())
        authorization_url = response.get_json()["authorization_url"]
        self.assertIn("accounts.google.com/o/oauth2/v2/auth", authorization_url)
        self.assertIn("calendar.readonly", authorization_url)
        with client.session_transaction() as oauth_session:
            self.assertEqual(oauth_session["google_calendar_oauth_faculty_id"], faculty_id)
            self.assertTrue(oauth_session["google_calendar_oauth_state"])

    def test_google_calendar_oauth_callback_saves_faculty_owned_credentials(self):
        with app.app_context():
            faculty = Faculty(
                name="OAuth Callback Faculty",
                college="Graduate School",
                role="Panel Member",
                specialization="Research methods",
                email="oauth-callback@example.test",
                active=True,
            )
            db.session.add(faculty)
            db.session.flush()
            account = self._account("faculty", "oauth-callback-login@example.test")
            account.faculty_id = faculty.id
            db.session.commit()
            account_id = account.id
            faculty_id = faculty.id

        client = self._role_client(account_id, "faculty")
        with client.session_transaction() as oauth_session:
            oauth_session["google_calendar_oauth_state"] = "expected-state"
            oauth_session["google_calendar_oauth_faculty_id"] = faculty_id
        calendar_response = MagicMock()
        calendar_response.__enter__.return_value = calendar_response
        calendar_response.read.return_value = json.dumps({
            "id": "oauth-callback@gmail.com",
            "primary": True,
        }).encode("utf-8")
        with patch.dict(os.environ, {
            "GOOGLE_CALENDAR_CLIENT_ID": "client-id.apps.googleusercontent.com",
            "GOOGLE_CALENDAR_CLIENT_SECRET": "client-secret",
            "GOOGLE_CALENDAR_REDIRECT_URI": "http://localhost/api/faculty-portal/google-calendar/callback",
        }), patch("app.google_token_request", return_value={
            "access_token": "faculty-access-token",
            "refresh_token": "faculty-refresh-token",
            "expires_in": 3600,
        }), patch("app.urlopen", return_value=calendar_response):
            response = client.get(
                "/api/faculty-portal/google-calendar/callback"
                "?state=expected-state&code=authorization-code"
            )

        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.headers["Location"].endswith("/faculty-portal?calendar=connected"))
        with app.app_context():
            stored = db.session.get(Faculty, faculty_id)
            self.assertEqual(stored.google_calendar_id, "oauth-callback@gmail.com")
            self.assertEqual(stored.google_calendar_email, "oauth-callback@gmail.com")
            self.assertEqual(stored.google_calendar_access_token, "faculty-access-token")
            self.assertEqual(stored.google_calendar_refresh_token, "faculty-refresh-token")
            self.assertIsNotNone(stored.google_calendar_connected_at)

    def test_defense_suggestions_start_after_required_lead_time(self):
        reference = date(2026, 7, 23)
        self.assertEqual(earliest_defense_date("Title Defense", reference), date(2026, 7, 23))
        self.assertEqual(earliest_defense_date("Public Final Defense", reference), date(2026, 7, 28))
        self.assertEqual(earliest_defense_date("Proposal Defense", reference), date(2026, 8, 6))
        self.assertEqual(earliest_defense_date("Final Defense", reference), date(2026, 8, 6))

    def test_maed_monitoring_sheet_seeds_all_policy_valid_demo_personas(self):
        with app.app_context():
            seed_database()

            program = Program.query.filter_by(code="MAED").one()
            active_term = get_active_term()
            all_students = {
                student.student_number: student
                for student in Student.query.filter_by(program_id=program.id).all()
            }
            # The two Practicum demos now belong to MAPSY and MSGC, so they
            # are intentionally absent from the MAED student count.
            self.assertEqual(len(all_students), 27)
            students = {
                student_number: student
                for student_number, student in all_students.items()
                if re.fullmatch(r"\d{7}", student_number)
            }
            self.assertEqual(len(students), 11)
            self.assertTrue(all(
                re.fullmatch(r"\d{7}", student.student_number)
                for student in students.values()
            ))
            self.assertTrue(all(
                student.academic_year_entry and student.year_level
                for student in students.values()
            ))
            self.assertTrue(active_term.is_active_planning_term)
            self.assertEqual(
                enrollment_integrity_payload(program, active_term)["issue_count"],
                0,
            )
            withdrawal_demo_enrollments = (
                SubjectEnrollment.query.join(Student)
                .filter(Student.student_number.in_(["GS-2026-WD-01", "GS-2026-WD-02"]))
                .all()
            )
            self.assertEqual(len(withdrawal_demo_enrollments), 6)
            self.assertTrue(all(item.status == "Enrolled" for item in withdrawal_demo_enrollments))
            withdrawal_enrollments_by_student = {}
            for item in withdrawal_demo_enrollments:
                withdrawal_enrollments_by_student.setdefault(item.student.student_number, []).append(item)
            self.assertEqual(
                {key: len(value) for key, value in withdrawal_enrollments_by_student.items()},
                {"GS-2026-WD-01": 3, "GS-2026-WD-02": 3},
            )

            expected_completed = {
                "2560001": 0,
                "2560002": 3,
                "2460003": 9,
                "2260004": 18,
                "2460005": 6,
                "2360006": 12,
                "2260007": 18,
                "2360008": 18,
                "2460009": 9,
                "2360010": 6,
                "2260011": 18,
            }
            for student_number, completed_count in expected_completed.items():
                audit = compute_course_audit(students[student_number])
                self.assertEqual(
                    len(audit["completed"]),
                    completed_count,
                    student_number,
                )
                self.assertEqual(audit["required_count"], 18, student_number)

            daniel = students["2560001"]
            self.assertEqual(daniel.current_stage, "Admission")
            self.assertEqual(
                TermEnrollment.query.filter_by(
                    student_id=daniel.id,
                    term_id=active_term.id,
                ).one().status,
                "Pending Enrollment",
            )

            grace = students["2560002"]
            self.assertEqual(grace.current_stage, "Coursework")
            self.assertEqual(
                {item.course.code for item in curriculum_offerings_for_term(program, active_term)},
                {
                    "MAED-MAJ1", "MAED-MAJ2", "MAED-MAJ3", "MAED-MAJ4", "MAED-MAJ5",
                    "MAED-COG1", "MAED-COG2",
                },
            )

            benjamin = students["2460003"]
            self.assertIsNotNone(TransactionLog.query.filter_by(
                student_id=benjamin.id,
                transaction_slug="leave-of-absence",
                result="LOA application submitted",
            ).first())
            self.assertEqual(
                submitted_request_students("leave-of-absence")[0]["student_number"],
                benjamin.student_number,
            )

            miguel = students["2260004"]
            miguel_progress = graduation_eligibility(miguel)["research_progress"]
            self.assertTrue(miguel_progress["stages"][0]["complete"])
            self.assertTrue(miguel_progress["stages"][1]["complete"])
            self.assertFalse(miguel_progress["stages"][2]["complete"])
            self.assertEqual(PanelAssignment.query.filter_by(
                student_id=miguel.id,
                gate="Final Defense",
            ).count(), 4)

            clarisse = students["2460005"]
            clarisse_case = AwolCase.query.filter_by(student_id=clarisse.id).one()
            self.assertEqual(clarisse_case.status, "Return Submitted")
            self.assertIsNone(clarisse_case.intent_attachment_id)
            self.assertTrue(clarisse_case.return_intent)

            adrian = students["2360006"]
            self.assertEqual(
                WithdrawalApplication.query.filter_by(student_id=adrian.id).one().status,
                "Submitted to GS Staff",
            )

            isabel = students["2260007"]
            isabel_eligibility = graduation_eligibility(isabel)
            self.assertTrue(isabel_eligibility["eligible"])
            self.assertEqual(isabel_eligibility["coursework_status"], "Complete")
            self.assertEqual(isabel_eligibility["research_status"], "Complete")
            self.assertEqual(isabel_eligibility["practicum_status"], "Not Required")
            self.assertEqual(
                GraduationEndorsement.query.filter_by(
                    student_id=isabel.id,
                ).one().endorsement_status,
                "For Review",
            )

            hector = students["2360008"]
            hector_eligibility = graduation_eligibility(hector)
            self.assertTrue(hector_eligibility["eligible"])
            self.assertEqual(hector_eligibility["missing_coursework"], [])
            self.assertEqual(hector_eligibility["missing_research_requirements"], [])
            self.assertEqual(hector_eligibility["missing_practicum_requirement"], "")
            self.assertEqual(hector_eligibility["practicum_status"], "Not Required")
            self.assertTrue(
                hector_eligibility["research_progress"]["research_gates_complete"]
            )
            self.assertTrue(
                hector_eligibility["research_progress"]["completion_evidence"]["complete"]
            )

            maria = students["2460009"]
            self.assertEqual(maria.standing, "On Leave")
            self.assertEqual(
                submitted_request_students("readmission")[0]["student_number"],
                maria.student_number,
            )

            paulo = students["2360010"]
            self.assertEqual(paulo.enrollment_tag, "AWOL")
            self.assertEqual(
                AwolCase.query.filter_by(student_id=paulo.id).one().status,
                "AWOL Declared",
            )

            elena = students["2260011"]
            elena_audit = compute_course_audit(elena)
            self.assertEqual(elena_audit["completed_units"], 54)
            self.assertEqual(elena.enrollment_tag, "Residency")
            residency = ResidencyEnrollment.query.filter_by(
                student_id=elena.id,
                term_id=active_term.id,
                status="Active",
            ).one()
            self.assertEqual(residency.reason, "Thesis / dissertation work")
            self.assertEqual(residency.policy_status, "Eligible for Residency")

            # Seeded request claims must be backed by a physical, downloadable
            # PDF rather than a database-only filename.
            attachments = StudentRequestAttachment.query.order_by(
                StudentRequestAttachment.id,
            ).all()
            # MAED no longer receives the four practicum MOA/certificate
            # fixtures previously attached to Isabel and Hector.
            self.assertEqual(len(attachments), 3)
            for item in attachments:
                path = REQUEST_UPLOAD_ROOT / item.stored_name
                self.assertTrue(path.is_file(), item.original_name)
                self.assertTrue(path.read_bytes().startswith(b"%PDF-"), item.original_name)

            # Completed research gates must have the required number of exact
            # file rows, physical PDFs, and adviser approvals where applicable.
            for student in (miguel, isabel, hector):
                documents = DocumentCheck.query.filter_by(student_id=student.id).all()
                for document in documents:
                    presentation = research_requirement_presentation(
                        document.gate,
                        document.item_name,
                    )
                    if not presentation or presentation["source_type"] != "student_upload":
                        continue
                    self.assertEqual(
                        len(document.evidence_files),
                        presentation["required_file_count"],
                        f"{student.student_number}: {document.item_name}",
                    )
                    for evidence in document.evidence_files:
                        path = UPLOAD_ROOT / evidence.stored_name
                        self.assertTrue(path.is_file(), evidence.original_name)
                        self.assertTrue(path.read_bytes().startswith(b"%PDF-"), evidence.original_name)
                        if document.item_name in ADVISER_APPROVAL_DOCUMENTS:
                            self.assertIsNotNone(
                                AdviserDocumentApproval.query.filter_by(
                                    evidence_file_id=evidence.id,
                                    status="Signed",
                                ).first(),
                                evidence.original_name,
                            )

            staff = UserAccount.query.filter_by(
                email="staff@usls.edu.ph",
                role="staff",
            ).one()
            staff_client = self._role_client(staff.id, "staff")
            for slug, student in (
                ("leave-of-absence", benjamin),
                ("readmission", maria),
                ("awol", clarisse),
            ):
                response = staff_client.post(f"/api/transactions/{slug}/messages", json={
                    "student_id": student.id,
                    "action_type": "note",
                    "recipient_role": "Student",
                    "template": "Please clarify request details",
                    "comment": f"Test message for {slug}.",
                })
                self.assertEqual(response.status_code, 200, response.get_json())

            loa_row = submitted_request_students("leave-of-absence")[0]
            readmission_row = submitted_request_students("readmission")[0]
            awol_row = next(
                row for row in awol_residency_roster_payload()
                if row["student_id"] == clarisse.id
            )
            self.assertEqual(loa_row["student"]["id"], benjamin.id)
            self.assertEqual(readmission_row["student"]["id"], maria.id)
            self.assertTrue(loa_row["messages"])
            self.assertTrue(readmission_row["messages"])
            self.assertTrue(awol_row["messages"])

    def test_workflow_clarification_can_be_returned_and_answered(self):
        with app.app_context():
            term = AcademicTerm(
                label="AY 2026-2027 Term 1",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
            )
            db.session.add(term)
            db.session.flush()
            subject_enrollment = SubjectEnrollment(
                student_id=self.student_id,
                course_id=self.course_id,
                term_id=term.id,
                status="Enrolled",
                source_reference="Clarification workflow fixture",
            )
            db.session.add(subject_enrollment)
            db.session.flush()
            application = WithdrawalApplication(
                student_id=self.student_id,
                subject_enrollment_id=subject_enrollment.id,
                withdrawal_scope="Subject",
                reason="Needs clarification test",
                effective_term="AY 2026-2027 Term 1",
                status="Dean Review",
                dean_decision="Pending",
            )
            request_file = self._attachment("withdrawal", "clarification-request")
            application.request_attachment_id = request_file.id
            student_account = self._account("student", "student-clarification@example.test")
            student_account.student_id = self.student_id
            db.session.add(application)
            db.session.commit()
            student_account_id = student_account.id

            response = self._staff_client().post("/api/transactions/withdrawal/messages", json={
                "student_id": self.student_id,
                "action_type": "return",
                "recipient_role": "Student",
                "template": "Missing required document",
                "comment": "",
            })
            self.assertEqual(response.status_code, 400, response.get_json())

            response = self._staff_client().post("/api/transactions/withdrawal/messages", json={
                "student_id": self.student_id,
                "action_type": "return",
                "recipient_role": "Student",
                "template": "Missing required document",
                "comment": "Upload the signed request form.",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(application)
            self.assertEqual(application.status, "Returned for Clarification")
            message = WorkflowMessage.query.filter_by(student_id=self.student_id, status="Open").first()
            self.assertEqual(message.previous_status, "Dean Review")
            self.assertEqual(message.workflow_request_id, application.id)
            self.assertEqual(message.workflow_stage, "Returned for Clarification")
            self.assertEqual(message.visibility, "student_visible")
            self.assertEqual(
                WorkflowMessage.query.filter_by(
                    transaction_slug="withdrawal",
                    student_id=self.student_id,
                    visibility="student_visible",
                    previous_status="Dean Review",
                    new_status="Returned for Clarification",
                ).count(),
                1,
            )

            student_client = self._role_client(student_account_id, "student")
            response = student_client.post("/api/transactions/withdrawal/messages", json={
                "action_type": "response",
                "recipient_role": "Graduate School Staff",
                "template": "Please clarify request details",
                "comment": "The signed request form has been uploaded.",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(application)
            db.session.refresh(message)
            self.assertEqual(application.status, "Returned for Clarification")
            self.assertEqual(message.status, "Open")
            response_message = WorkflowMessage.query.filter_by(
                student_id=self.student_id,
                action_type="response",
            ).first()
            self.assertEqual(response_message.status, "Sent")

            response = student_client.post("/api/student-portal/requests/withdrawal", json={
                "reason": application.reason,
                "effective_term": application.effective_term,
                "attachment_id": request_file.id,
                "subject_enrollment_id": subject_enrollment.id,
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(application)
            db.session.refresh(message)
            db.session.refresh(request_file)
            self.assertEqual(application.status, "Submitted to GS Staff")
            self.assertEqual(application.reason, "Needs clarification test")
            self.assertEqual(application.request_attachment_id, request_file.id)
            self.assertEqual(message.status, "Responded")
            self.assertEqual(request_file.workflow_request_id, application.id)
            self.assertEqual(request_file.workflow_stage, "Withdrawal Application")
            self.assertEqual(request_file.uploaded_by_user_id, student_account_id)
            self.assertEqual(request_file.uploaded_by_role, "Student")

    def test_student_practicum_submission_is_locked_until_eligible(self):
        with app.app_context():
            attachment = self._attachment("practicum", "locked-moa")
            student_account = self._account("student", "student-practicum@example.test")
            student_account.student_id = self.student_id
            db.session.commit()

            response = self._role_client(student_account.id, "student").post(
                "/api/student-portal/requests/practicum",
                json={
                    "attachment_id": attachment.id,
                    "practicum_site": "Partner Site",
                },
            )
            self.assertEqual(response.status_code, 409, response.get_json())
            self.assertIsNone(PracticumRecord.query.filter_by(student_id=self.student_id).first())

    def test_practicum_is_restricted_to_psychology_and_msgc_programs(self):
        with app.app_context():
            out_of_scope_program = Program(
                code="MAED-NP",
                name="Master of Arts in Education",
                college="Education",
                has_practicum=True,
            )
            msgc_program = Program(
                code="MSGC",
                name="Master of Science in Guidance and Counseling",
                college="Education",
                has_practicum=True,
            )
            db.session.add_all([out_of_scope_program, msgc_program])
            db.session.flush()
            out_of_scope_student = Student(
                student_number="GS-2026-NON-PRAC",
                first_name="Outside",
                last_name="Scope",
                email="outside-practicum@example.test",
                program_id=out_of_scope_program.id,
                entry_year=2025,
                current_stage="Final Defense",
                standing="Active",
            )
            db.session.add(out_of_scope_student)
            db.session.flush()
            account = self._account("student", "outside-practicum-login@example.test")
            account.student_id = out_of_scope_student.id
            db.session.commit()

            psychology_program = db.session.get(Program, self.program_id)
            self.assertTrue(program_allows_practicum(psychology_program))
            self.assertTrue(program_allows_practicum(msgc_program))
            self.assertFalse(program_allows_practicum(out_of_scope_program))
            self.assertFalse(practicum_eligibility(out_of_scope_student)["program_in_scope"])

            portal = self._role_client(account.id, "student").get("/api/student-portal/context")
            self.assertEqual(portal.status_code, 200, portal.get_json())
            self.assertFalse(portal.get_json()["student"]["program_has_practicum"])
            self.assertIn("Psychology", portal.get_json()["practicum_program_scope"]["label"])
            self.assertIn("MSGC", portal.get_json()["practicum_program_scope"]["label"])

            blocked = self._role_client(account.id, "student").post(
                "/api/student-portal/requests/practicum",
                json={"practicum_site": "Not allowed"},
            )
            self.assertEqual(blocked.status_code, 403, blocked.get_json())
            self.assertIn("Psychology", blocked.get_json()["error"])

            roster_ids = {
                row["student"]["id"]
                for row in practicum_roster_payload()
            }
            self.assertIn(self.student_id, roster_ids)
            self.assertNotIn(out_of_scope_student.id, roster_ids)

    def test_student_can_upload_additional_practicum_pdf_without_replacing_history(self):
        with app.app_context():
            old_certificate = self._attachment("practicum", "initial-certificate")
            record = PracticumRecord(
                student_id=self.student_id,
                required_hours=200,
                completed_hours=120,
                certificate_attachment_id=old_certificate.id,
                document_status="Pending Review",
                status="Additional Certificates Requested",
            )
            db.session.add(record)
            db.session.flush()
            old_certificate.workflow_request_id = record.id
            old_certificate.workflow_stage = "Practicum Document Submission"
            student_account = self._account("student", "student-practicum-additional@example.test")
            student_account.student_id = self.student_id
            db.session.commit()

            student_client = self._role_client(student_account.id, "student")
            response = student_client.post("/api/student-portal/requests/practicum", json={
                "certificate_attachment_id": old_certificate.id,
                "completed_hours": 200,
                "certificate_count": 99,
            })
            self.assertEqual(response.status_code, 400, response.get_json())
            self.assertIn("additional practicum PDF", response.get_json()["error"])

            response = student_client.post(
                "/api/student-portal/request-attachments/upload",
                data={
                    "request_type": "practicum",
                    "file": (BytesIO(b"%PDF-1.4\nadditional certificate\n%%EOF\n"), "additional-certificate.pdf"),
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            new_attachment_id = response.get_json()["attachment"]["id"]
            new_attachment = db.session.get(StudentRequestAttachment, new_attachment_id)
            self.created_files.append(REQUEST_UPLOAD_ROOT / new_attachment.stored_name)

            response = student_client.post("/api/student-portal/requests/practicum", json={
                "certificate_attachment_id": new_attachment_id,
                "completed_hours": 200,
                "certificate_count": 99,
                "remarks": "This additional certificate covers the remaining practicum hours.",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(record)
            db.session.refresh(old_certificate)
            db.session.refresh(new_attachment)
            self.assertEqual(record.status, "Documents Submitted")
            self.assertEqual(record.certificate_attachment_id, new_attachment_id)
            self.assertEqual(record.certificate_count, 2)
            self.assertEqual(old_certificate.workflow_request_id, record.id)
            self.assertEqual(new_attachment.workflow_request_id, record.id)
            self.assertEqual(
                StudentRequestAttachment.query.filter_by(
                    student_id=self.student_id,
                    request_type="practicum",
                    workflow_request_id=record.id,
                ).count(),
                2,
            )
            submission_remark = WorkflowMessage.query.filter_by(
                transaction_slug="practicum",
                student_id=self.student_id,
                sender_role="Student",
                action_type="submission",
            ).order_by(WorkflowMessage.id.desc()).first()
            self.assertIsNotNone(submission_remark)
            self.assertEqual(submission_remark.recipient_role, "Graduate School Staff")
            self.assertIn("remaining practicum hours", submission_remark.comment)

            staff_message = self._staff_client().post("/api/transactions/practicum/messages", json={
                "student_id": self.student_id,
                "action_type": "note",
                "recipient_role": "Academic Coordinator",
                # Practicum and Graduation deliberately ignore legacy templates;
                # users only enter the actual remarks shown in the discussion.
                "template": "Please clarify request details",
                "comment": "Please review the student's additional certificate.",
            })
            self.assertEqual(staff_message.status_code, 200, staff_message.get_json())
            self.assertEqual(staff_message.get_json()["item"]["template"], "Remarks")
            self.assertEqual(
                staff_message.get_json()["item"]["comment"],
                "Please review the student's additional certificate.",
            )

    def test_request_files_are_scoped_to_student_and_workflow_role(self):
        with app.app_context():
            attachment = self._attachment("withdrawal", "private-request")
            student_account = self._account("student", "student-files@example.test")
            student_account.student_id = self.student_id
            other_record = Student(
                student_number="GS-2026-OTHER-FILE",
                first_name="Other",
                last_name="File Student",
                email="other-file@example.test",
                program_id=self.program_id,
                entry_year=2025,
                current_stage="Coursework",
                standing="Active",
            )
            db.session.add(other_record)
            db.session.flush()
            other_student = self._account("student", "other-student@example.test")
            other_student.student_id = other_record.id
            REQUEST_UPLOAD_ROOT.mkdir(parents=True, exist_ok=True)
            stored_path = REQUEST_UPLOAD_ROOT / attachment.stored_name
            stored_path.write_bytes(b"%PDF-1.4\n%%EOF\n")
            self.created_files.append(stored_path)
            db.session.commit()

            response = self._role_client(other_student.id, "student").get(
                f"/api/student-request-attachments/{attachment.id}/file"
            )
            self.assertEqual(response.status_code, 403)

            response = self._research_client().get(
                f"/api/student-request-attachments/{attachment.id}/file"
            )
            self.assertEqual(response.status_code, 403)

            response = self._role_client(student_account.id, "student").get(
                f"/api/student-request-attachments/{attachment.id}/file"
            )
            self.assertEqual(response.status_code, 200)
            self.assertTrue(response.data.startswith(b"%PDF-"))
            response.close()
            staff_response = self._staff_client().get(
                f"/api/student-request-attachments/{attachment.id}/file"
            )
            self.assertEqual(staff_response.status_code, 200)
            staff_response.close()
            dean_response = self._dean_client().get(
                f"/api/student-request-attachments/{attachment.id}/file"
            )
            self.assertEqual(dean_response.status_code, 200)
            dean_response.close()

            stored_path.unlink()
            response = self._role_client(student_account.id, "student").get(
                f"/api/student-request-attachments/{attachment.id}/file"
            )
            self.assertEqual(response.status_code, 404)
            self.assertIn("unavailable", response.get_json()["error"].lower())

    def test_student_inbox_hides_internal_messages_and_tracks_read_state(self):
        with app.app_context():
            application = WithdrawalApplication(
                student_id=self.student_id,
                reason="Message privacy test",
                effective_term="AY 2026-2027 Term 1",
                status="Dean Review",
                dean_decision="Pending",
            )
            db.session.add(application)
            student_account = self._account("student", "student-inbox@example.test")
            student_account.student_id = self.student_id
            other_record = Student(
                student_number="GS-2026-OTHER-INBOX",
                first_name="Other",
                last_name="Inbox Student",
                email="other-inbox@example.test",
                program_id=self.program_id,
                entry_year=2025,
                current_stage="Coursework",
                standing="Active",
            )
            db.session.add(other_record)
            db.session.flush()
            other_account = self._account("student", "other-inbox-account@example.test")
            other_account.student_id = other_record.id
            db.session.commit()

            internal_response = self._staff_client().post("/api/transactions/withdrawal/messages", json={
                "student_id": self.student_id,
                "action_type": "note",
                "recipient_role": "Academic Coordinator",
                "visibility": "internal",
                "template": "This request requires additional review",
                "comment": "Internal reviewer concern.",
            })
            self.assertEqual(internal_response.status_code, 200, internal_response.get_json())
            internal = internal_response.get_json()["item"]
            self.assertEqual(internal["visibility"], "internal")
            self.assertEqual(internal["request_id"], application.id)
            self.assertEqual(internal["stage"], "Dean Review")
            self.assertEqual(internal["sender_user_id"], self.staff_id)

            visible_response = self._staff_client().post("/api/transactions/withdrawal/messages", json={
                "student_id": self.student_id,
                "action_type": "note",
                "recipient_role": "Student",
                "visibility": "student_visible",
                "template": "Please clarify the information entered in this section.",
                "comment": "Please confirm the requested effective term.",
            })
            self.assertEqual(visible_response.status_code, 200, visible_response.get_json())
            visible = visible_response.get_json()["item"]
            self.assertTrue(visible["is_unread"])

            student_client = self._role_client(student_account.id, "student")
            context_response = student_client.get("/api/student-portal/context")
            self.assertEqual(context_response.status_code, 200, context_response.get_json())
            inbox_ids = {item["id"] for item in context_response.get_json()["workflow_messages"]}
            self.assertIn(visible["id"], inbox_ids)
            self.assertNotIn(internal["id"], inbox_ids)
            student_log_notes = " ".join(
                item.get("notes") or ""
                for item in context_response.get_json()["logs"]
            )
            self.assertNotIn("Internal reviewer concern.", student_log_notes)
            internal_log = TransactionLog.query.filter(
                TransactionLog.notes.contains("Internal reviewer concern.")
            ).first()
            self.assertEqual(internal_log.visibility, "internal")

            read_response = student_client.post("/api/student-portal/messages/read", json={
                "message_ids": [internal["id"], visible["id"]],
            })
            self.assertEqual(read_response.status_code, 200, read_response.get_json())
            self.assertEqual(read_response.get_json()["message_ids"], [visible["id"]])
            db.session.refresh(db.session.get(WorkflowMessage, visible["id"]))
            self.assertIsNotNone(db.session.get(WorkflowMessage, visible["id"]).read_at)

            other_context = self._role_client(other_account.id, "student").get("/api/student-portal/context")
            self.assertEqual(other_context.status_code, 200, other_context.get_json())
            other_ids = {item["id"] for item in other_context.get_json()["workflow_messages"]}
            self.assertNotIn(visible["id"], other_ids)
            self.assertNotIn(internal["id"], other_ids)

    def test_reviewer_transition_creates_linked_student_notice_and_audit(self):
        with app.app_context():
            term = get_active_term()
            if not term:
                term = AcademicTerm(
                    label="AY 2026-2027 1st Semester",
                    start_date=date(2026, 8, 1),
                    end_date=date(2026, 12, 15),
                    is_active_planning_term=True,
                )
                db.session.add(term)
                db.session.flush()
            enrollment = SubjectEnrollment(
                student_id=self.student_id,
                course_id=self.course_id,
                term_id=term.id,
                status="Enrolled",
            )
            db.session.add(enrollment)
            db.session.flush()
            application = WithdrawalApplication(
                student_id=self.student_id,
                reason="Approval notice test",
                effective_term=term.label,
                subject_enrollment_id=enrollment.id,
                status="Dean Review",
                dean_decision="Pending",
            )
            db.session.add(application)
            student_account = self._account("student", "student-notice@example.test")
            student_account.student_id = self.student_id
            db.session.commit()

            response = self._dean_client().post(
                f"/api/approvals/workflow/withdrawal/{application.id}/decide",
                json={"decision": "approve", "note": "Approved for follow-through."},
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            notice = WorkflowMessage.query.filter_by(
                transaction_slug="withdrawal",
                student_id=self.student_id,
                action_type="notice",
            ).first()
            self.assertIsNotNone(notice)
            self.assertEqual(notice.workflow_request_id, application.id)
            self.assertEqual(notice.workflow_stage, "Approved - Awaiting Subject Tag")
            self.assertEqual(notice.visibility, "student_visible")
            audit = TransactionLog.query.filter_by(
                transaction_slug="withdrawal",
                student_id=self.student_id,
                new_status="Approved - Awaiting Subject Tag",
            ).order_by(TransactionLog.id.desc()).first()
            self.assertEqual(audit.workflow_request_id, application.id)
            self.assertEqual(audit.actor_user_id, self.dean_id)
            self.assertEqual(audit.action_type, "approve")

    def test_graduation_batch_skips_unverified_and_sends_eligible_candidates(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            db.session.add(CourseRecord(student_id=student.id, course_id=self.course_id, status="Completed"))
            db.session.add(ResearchCase(
                student_id=student.id,
                case_type="Thesis",
                title="Batch-ready research",
                current_gate="Completion Evidence",
                status="Verified Complete",
            ))
            for item in required_documents_for_gate("Completion Evidence"):
                db.session.add(DocumentCheck(student_id=student.id, gate="Completion Evidence", item_name=item, status="Complete"))
            db.session.add(PracticumRecord(
                student_id=student.id,
                required_hours=200,
                completed_hours=200,
                document_status="Verified",
                completion_status="Completed and accepted",
                status="Dean Reviewed",
            ))
            db.session.commit()

            roster_item = next(
                row for row in graduation_candidate_payload()
                if row["student"]["id"] == student.id
            )
            self.assertTrue(roster_item["eligibility"]["eligible"])
            self.assertFalse(roster_item["application_submitted"])
            self.assertEqual(roster_item["application_status"], "Awaiting Student Application")
            self.assertEqual(
                roster_item["eligibility"]["completed_courses"],
                [{
                    "id": self.course_id,
                    "code": "BPM-501",
                    "title": "Completion Course",
                    "units": 3,
                    "category": "Major",
                    "term_label": None,
                }],
            )
            self.assertEqual(
                roster_item["eligibility"]["practicum_completion"],
                {
                    "required": True,
                    "complete": True,
                    "record_exists": True,
                    "status": "Dean Reviewed",
                    "practicum_site": None,
                    "supervisor_name": None,
                    "required_hours": 200,
                    "completed_hours": 200,
                    "moa_status": "Pending Review",
                    "document_status": "Verified",
                    "certificate_count": 0,
                    "completion_status": "Completed and accepted",
                    "report_sent_at": None,
                    "dean_reviewed_at": None,
                },
            )

            response = self._staff_client().post("/api/graduation/batch-actions", json={
                "student_ids": [student.id],
                "bpm_action": "create_batch",
                "review_window": "AY 2026-2027",
                "batch_no": "2",
                "batch_month": "July",
                "batch_year": "2026",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            self.assertEqual(response.get_json()["updated"], [])
            self.assertIn("must submit", response.get_json()["skipped"][0]["reason"])

            application_file = self._attachment("graduation", "batch-ready-application")
            db.session.add(GraduationEndorsement(
                student_id=student.id,
                request_attachment_id=application_file.id,
                review_window="AY 2026-2027",
                endorsement_status="For Review",
            ))
            db.session.commit()
            roster_item = next(
                row for row in graduation_candidate_payload()
                if row["student"]["id"] == student.id
            )
            self.assertTrue(roster_item["application_submitted"])
            self.assertEqual(roster_item["application_status"], "Application Submitted")

            response = self._staff_client().post("/api/graduation/batch-actions", json={
                "student_ids": [student.id],
                "bpm_action": "create_batch",
                "review_window": "AY 2026-2027",
                "batch_no": "2",
                "batch_month": "July",
                "batch_year": "2026",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            self.assertEqual(len(response.get_json()["updated"]), 1)
            self.assertEqual(response.get_json()["batch_name"], "Batch 2 July 2026")
            endorsement = GraduationEndorsement.query.filter_by(student_id=student.id).first()
            self.assertEqual(endorsement.endorsement_status, "For Review")
            self.assertEqual(endorsement.batch_name, "Batch 2 July 2026")

            response = self._staff_client().post("/api/graduation/batch-actions", json={
                "student_ids": [student.id],
                "bpm_action": "compile_to_ac",
                "review_window": "AY 2026-2027",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            self.assertEqual(len(response.get_json()["updated"]), 1)
            db.session.refresh(endorsement)
            self.assertEqual(endorsement.endorsement_status, "Coursework Review")
            self.assertEqual(endorsement.batch_name, "Batch 2 July 2026")

            response = self._academic_client().post("/api/graduation/batch-actions", json={
                "student_ids": [student.id],
                "bpm_action": "check_coursework",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(endorsement)
            self.assertEqual(endorsement.endorsement_status, "Research Review")

            response = self._research_client().post("/api/graduation/batch-actions", json={
                "student_ids": [student.id],
                "bpm_action": "validate_research",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(endorsement)
            self.assertEqual(endorsement.endorsement_status, "Eligibility Confirmed")

            response = self._staff_client().post("/api/graduation/batch-actions", json={
                "student_ids": [student.id],
                "bpm_action": "prepare_endorsement",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(endorsement)
            self.assertEqual(endorsement.endorsement_status, "Endorsement Prepared")

            response = self._staff_client().post("/api/graduation/batch-actions", json={
                "student_ids": [student.id],
                "bpm_action": "send_to_dean",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            db.session.refresh(endorsement)
            self.assertEqual(endorsement.endorsement_status, "Ready for Dean Review")
            dean_payload = workflow_approvals_payload()
            dean_item = next(
                item for item in dean_payload["pending"]
                if item["type"] == "graduation" and item["student"]["id"] == student.id
            )
            self.assertEqual(dean_item["batch_name"], "Batch 2 July 2026")
            self.assertEqual(dean_item["record"]["batch_name"], "Batch 2 July 2026")

            blocked_student = Student(
                student_number="GS-2026-BATCH-BLOCKED",
                first_name="Blocked",
                last_name="Candidate",
                email="batch-blocked@example.test",
                program_id=self.program_id,
                entry_year=2025,
                current_stage="Final Defense",
                standing="Active",
            )
            db.session.add(blocked_student)
            db.session.flush()
            blocked = GraduationEndorsement(
                student_id=blocked_student.id,
                review_window="AY 2026-2027",
                endorsement_status="For Review",
            )
            db.session.add(blocked)
            db.session.commit()

            response = self._dean_client().post("/api/graduation/batch-actions", json={
                "student_ids": [student.id, blocked_student.id],
                "action": "approve",
                "comment": "Approved in the Dean review batch.",
            })
            self.assertEqual(response.status_code, 200, response.get_json())
            self.assertEqual(len(response.get_json()["updated"]), 1)
            self.assertEqual(len(response.get_json()["skipped"]), 1)
            db.session.refresh(endorsement)
            self.assertEqual(endorsement.endorsement_status, "Dean Approved")
            self.assertEqual(blocked.endorsement_status, "For Review")
            export = self._dean_client().post("/api/graduation/endorsed.csv", json={"endorsement_ids": [endorsement.id]})
            self.assertEqual(export.status_code, 200, export.data.decode())
            self.assertIn("Batch,Student ID,Student Name", export.data.decode())
            self.assertIn("Batch 2 July 2026", export.data.decode())
            db.session.refresh(endorsement)
            self.assertEqual(endorsement.endorsement_status, "Dean Approved")
            self.assertEqual(endorsement.registrar_status, "Exported - Ready to Send")

            endorsement.endorsement_status = "Ready for Dean Review"
            endorsement.registrar_status = "Pending"
            db.session.commit()
            response = self._dean_client().post("/api/graduation/batch-actions", json={
                "student_ids": [student.id],
                "action": "return",
                "comment": "",
            })
            self.assertEqual(response.status_code, 400, response.get_json())

    def test_staff_cannot_reset_a_non_demo_student_through_demo_controls(self):
        with app.app_context():
            self.student = db.session.get(Student, self.student_id)
            db.session.add(CourseRecord(student_id=self.student.id, course_id=self.course_id, status="Completed"))
            research = ResearchCase(
                student_id=self.student.id,
                case_type="Thesis",
                title="Source data kept after demo resets",
                current_gate="Completion Evidence",
                status="Verified Complete",
            )
            db.session.add(research)
            moa = self._attachment("practicum", "reset-moa")
            db.session.add(PracticumRecord(student_id=self.student.id, moa_attachment_id=moa.id, status="Dean Reviewed"))
            db.session.add(WithdrawalApplication(student_id=self.student.id, status="Withdrawn Confirmed", dean_decision="Approved"))
            db.session.add(GraduationEndorsement(student_id=self.student.id, endorsement_status="Sent to Registrar"))
            db.session.add(StudentRequestAttachment(
                student_id=self.student.id,
                request_type="graduation-registrar-handoff",
                original_name="endorsed-list.csv",
                stored_name="reset-demo-graduation-registrar-handoff.csv",
                mime_type="text/csv",
            ))
            self.student.standing = "Withdrawn"
            self.student.current_stage = "Withdrawn"
            self.student.enrollment_tag = "Withdrawn"
            db.session.commit()

            dean_response = self._dean_client().post(
                "/api/transactions/practicum/demo-reset",
                json={"student_id": self.student.id},
            )
            self.assertEqual(dean_response.status_code, 403)

            staff = self._staff_client()
            for slug in ("practicum", "withdrawal", "graduation"):
                response = staff.post(f"/api/transactions/{slug}/demo-reset", json={"student_id": self.student.id})
                self.assertEqual(response.status_code, 400, response.get_json())

            self.assertIsNotNone(PracticumRecord.query.filter_by(student_id=self.student.id).first())
            self.assertIsNotNone(WithdrawalApplication.query.filter_by(student_id=self.student.id).first())
            self.assertIsNotNone(GraduationEndorsement.query.filter_by(student_id=self.student.id).first())
            self.assertIsNotNone(StudentRequestAttachment.query.filter_by(student_id=self.student.id, request_type="graduation-registrar-handoff").first())
            self.assertIsNotNone(CourseRecord.query.filter_by(student_id=self.student.id).first())
            self.assertIsNotNone(ResearchCase.query.filter_by(student_id=self.student.id).first())
            db.session.refresh(self.student)
            self.assertEqual(self.student.standing, "Withdrawn")
            self.assertEqual(self.student.current_stage, "Withdrawn")


    def test_cumulative_monitoring_validation_resolution_and_new_student_badge(self):
        with app.app_context():
            db.session.add(CourseRecord(
                student_id=self.student_id,
                course_id=self.course_id,
                status="Completed",
                evidence_reference="Prior monitoring sheet",
            ))
            additive_student = Student(
                student_number="GS-ADD-001",
                first_name="Additive",
                last_name="Update",
                email="additive@example.test",
                program_id=self.program_id,
                entry_year=2026,
                academic_year_entry="26-27",
                year_level="1",
                current_stage="Admission",
                standing="Active",
            )
            db.session.add(additive_student)
            db.session.flush()
            upload = MonitoringSheetUpload(
                original_name="semester-monitoring.xlsx",
                stored_name="test-semester-monitoring.xlsx",
                program_code="BPM",
                row_count=6,
                subject_count=1,
                snapshot_json="{}",
                result_json="{}",
            )
            db.session.add(upload)
            db.session.flush()
            parsed = {
                "program_code": "BPM",
                "subjects": ["BPM-501"],
                "subject_categories": {"BPM-501": "Major"},
                "subject_titles": {"BPM-501": "Completion Course"},
                "issues": [],
                "rows": [
                    {
                        "row": 10, "idno": "GS-2026-TEST",
                        "first_name": "Workflow Changed", "last_name": "Student",
                        "course": "BPM", "year": "2", "ay_entry": "24-25",
                        "subjects": {"BPM-501": False}, "milestones": {},
                        "comprehensive_exam_passed": False, "note": "",
                    },
                    {
                        "row": 11, "idno": "GS-NEW-001", "first_name": "New", "last_name": "Student",
                        "course": "BPM", "year": "1", "ay_entry": "26-27",
                        "subjects": {"BPM-501": True}, "milestones": {},
                        "comprehensive_exam_passed": False, "note": "",
                    },
                    {
                        "row": 15, "idno": "GS-ADD-001", "first_name": "Additive", "last_name": "Update",
                        "course": "BPM", "year": "1", "ay_entry": "26-27",
                        "subjects": {"BPM-501": True}, "milestones": {},
                        "comprehensive_exam_passed": False, "note": "",
                    },
                    {
                        "row": 12, "idno": "", "first_name": "Needs", "last_name": "Identifier",
                        "course": "BPM", "year": "1", "ay_entry": "26-27",
                        "subjects": {"BPM-501": True}, "milestones": {},
                        "comprehensive_exam_passed": False, "note": "",
                    },
                    {
                        "row": 13, "idno": "GS-DUP-001", "first_name": "Duplicate", "last_name": "One",
                        "course": "BPM", "year": "1", "ay_entry": "26-27",
                        "subjects": {"BPM-501": False}, "milestones": {},
                        "comprehensive_exam_passed": False, "note": "",
                    },
                    {
                        "row": 14, "idno": "GS-DUP-001", "first_name": "Duplicate", "last_name": "Two",
                        "course": "BPM", "year": "1", "ay_entry": "26-27",
                        "subjects": {"BPM-501": False}, "milestones": {},
                        "comprehensive_exam_passed": False, "note": "",
                    },
                ],
            }
            result = import_ac_monitoring(parsed, upload=upload)
            upload.result_json = "{}"
            db.session.commit()

            self.assertEqual(result["created"], 1)
            self.assertEqual(result["unresolved_count"], 4)
            issue_types = {item.issue_type for item in MonitoringValidationIssue.query.filter_by(upload_id=upload.id).all()}
            self.assertTrue(any("Subject mismatch" in value for value in issue_types))
            self.assertTrue(any("Missing required fields" in value for value in issue_types))
            self.assertTrue(any("Duplicate student ID in upload" in value for value in issue_types))
            multi_issue = next(item for item in result["validation_issues"] if item["incoming_student_number"] == "GS-2026-TEST")
            self.assertGreaterEqual(multi_issue["discrepancy_count"], 3)
            self.assertEqual(
                {item["type"] for item in multi_issue["discrepancies"]},
                {"Existing ID with different name", "School year mismatch", "Subject mismatch"},
            )
            new_student = Student.query.filter_by(student_number="GS-NEW-001").one()
            self.assertTrue(new_student.monitoring_new_student)
            self.assertEqual(new_student.monitoring_upload_id, upload.id)
            self.assertEqual(
                CourseRecord.query.filter_by(student_id=additive_student.id, course_id=self.course_id).one().status,
                "Completed",
            )

            staff = self._staff_client()
            subject_issue = MonitoringValidationIssue.query.filter(
                MonitoringValidationIssue.upload_id == upload.id,
                MonitoringValidationIssue.issue_type.like("%Subject mismatch%"),
            ).one()
            used = staff.patch(f"/api/monitoring/issues/{subject_issue.id}/resolve", json={"action": "use_uploaded"})
            self.assertEqual(used.status_code, 200, used.get_json())
            self.assertEqual(CourseRecord.query.filter_by(student_id=self.student_id, course_id=self.course_id).one().status, "Missing")

            missing_issue = MonitoringValidationIssue.query.filter(
                MonitoringValidationIssue.upload_id == upload.id,
                MonitoringValidationIssue.issue_type.like("%Missing required fields%"),
            ).one()
            edited = staff.patch(f"/api/monitoring/issues/{missing_issue.id}/resolve", json={
                "action": "edit",
                "edited": {"student_number": "GS-EDIT-001"},
            })
            self.assertEqual(edited.status_code, 200, edited.get_json())
            edited_student = Student.query.filter_by(student_number="GS-EDIT-001").one()
            self.assertTrue(edited_student.monitoring_new_student)
            self.assertTrue(TransactionLog.query.filter_by(student_id=edited_student.id, new_status="Resolved").first())

            duplicate_issue = MonitoringValidationIssue.query.filter(
                MonitoringValidationIssue.upload_id == upload.id,
                MonitoringValidationIssue.issue_type.like("%Duplicate student ID in upload%"),
            ).first()
            kept = staff.patch(f"/api/monitoring/issues/{duplicate_issue.id}/resolve", json={"action": "keep_existing"})
            self.assertEqual(kept.status_code, 200, kept.get_json())
            self.assertEqual(kept.get_json()["issue"]["status"], "Resolved")

    def test_corrected_monitoring_workbook_resolves_multiple_selected_students(self):
        from openpyxl import Workbook

        with app.app_context():
            second_student = Student(
                student_number="GS-BULK-002",
                first_name="Second",
                last_name="Student",
                email="second-bulk@example.test",
                program_id=self.program_id,
                entry_year=2025,
                academic_year_entry="25-26",
                year_level="1",
                current_stage="Coursework",
                standing="Active",
            )
            db.session.add(second_student)
            db.session.flush()
            db.session.add_all([
                CourseRecord(student_id=self.student_id, course_id=self.course_id, status="Completed"),
                CourseRecord(student_id=second_student.id, course_id=self.course_id, status="Completed"),
            ])
            source_upload = MonitoringSheetUpload(
                original_name="source-errors.xlsx",
                stored_name="test-source-errors.xlsx",
                program_code="BPM",
                row_count=2,
                subject_count=1,
                snapshot_json="{}",
                result_json="{}",
            )
            db.session.add(source_upload)
            db.session.flush()
            parsed = {
                "program_code": "BPM",
                "subjects": ["BPM-501"],
                "subject_categories": {"BPM-501": "Major"},
                "subject_titles": {"BPM-501": "Completion Course"},
                "issues": [],
                "rows": [
                    {"row": 5, "idno": "GS-2026-TEST", "first_name": "Wrong", "last_name": "Name", "course": "BPM", "year": "2", "ay_entry": "24-25", "subjects": {"BPM-501": False}, "milestones": {}, "comprehensive_exam_passed": False, "note": ""},
                    {"row": 6, "idno": "GS-BULK-002", "first_name": "Also Wrong", "last_name": "Name", "course": "BPM", "year": "2", "ay_entry": "24-25", "subjects": {"BPM-501": False}, "milestones": {}, "comprehensive_exam_passed": False, "note": ""},
                ],
            }
            result = import_ac_monitoring(parsed, upload=source_upload)
            db.session.commit()
            issue_ids = [item["id"] for item in result["validation_issues"]]
            self.assertEqual(len(issue_ids), 2)
            self.assertTrue(all(item["discrepancy_count"] >= 3 for item in result["validation_issues"]))

            workbook = Workbook()
            sheet = workbook.active
            sheet["A1"] = "PROGRAM: BPM"
            sheet.append([])
            sheet.append(["", "IDNO", "COURSE", "YR", "BASIC", "", "", "NOTE"])
            sheet.append(["AY ENTRY", "", "", "", "BPM-501", "SN", "FN", ""])
            sheet.append(["25-26", "GS-2026-TEST", "BPM", "1", "X", "Student", "Workflow", ""])
            sheet.append(["25-26", "GS-BULK-002", "BPM", "1", "X", "Student", "Second", ""])
            corrected_file = BytesIO()
            workbook.save(corrected_file)
            corrected_file.seek(0)

            response = self._staff_client().post(
                "/api/monitoring/issues/resolve-upload",
                data={
                    "issue_ids": json.dumps(issue_ids),
                    "file": (corrected_file, "corrected-monitoring.xlsx"),
                },
                content_type="multipart/form-data",
            )
            self.assertEqual(response.status_code, 200, response.get_json())
            payload = response.get_json()
            self.assertEqual(payload["resolved_count"], 2)
            self.assertEqual(payload["remaining_count"], 0)
            resolved = MonitoringValidationIssue.query.filter(MonitoringValidationIssue.id.in_(issue_ids)).all()
            self.assertTrue(all(item.status == "Resolved" for item in resolved))
            self.assertEqual(len({item.resolution_upload_id for item in resolved}), 1)
            correction_upload = db.session.get(MonitoringSheetUpload, resolved[0].resolution_upload_id)
            self.created_files.append(MONITORING_UPLOAD_ROOT / correction_upload.stored_name)

    def test_owner_role_labels_are_canonicalised_and_system_actors_excluded(self):
        # Task ownership and transaction-log actor labels use different
        # vocabularies; the workload report must not split one person across
        # several rows, and must not report the system as carrying load.
        self.assertEqual(canonical_owner_role("Graduate School Staff"), "GS Staff")
        self.assertEqual(canonical_owner_role("GS Staff"), "GS Staff")
        self.assertEqual(canonical_owner_role("Graduate School Staff · Demo Account"), "GS Staff")
        self.assertEqual(canonical_owner_role("academic coordinator"), "Academic Coordinator")
        self.assertEqual(canonical_owner_role("Adviser"), "Faculty")
        for system_actor in ("Demo Data", "Workflow System", "System · LOA RAG", ""):
            self.assertIsNone(canonical_owner_role(system_actor), system_actor)

    def test_analytics_reports_all_generate_and_kpis_are_measured(self):
        with app.app_context():
            payload = analytics_payload({})
            for key, _label in ANALYTICS_REPORTS:
                report = payload[key]
                self.assertNotIn("error", report, f"{key} raised: {report.get('error')}")
                self.assertIsInstance(report["rows"], list, key)
                self.assertTrue(report["columns"], f"{key} declared no columns")
                # Every declared column must be resolvable on every row.
                for row in report["rows"]:
                    for column in report["columns"]:
                        self.assertIn(column["key"], row, f"{key} row missing {column['key']}")

            kpis = {item["kpi"]: item for item in payload["kpis"]}
            self.assertEqual(
                kpis["Analytics report generation rate"]["value"], 100.0,
                kpis["Analytics report generation rate"]["basis"],
            )
            self.assertEqual(kpis["Required report availability"]["value"], 100.0)
            # With no flagged cases the closure rate is not measurable rather
            # than being reported as a passing or failing figure.
            self.assertIsNone(kpis["Follow-up closure logging rate"]["met"])

    def test_analytics_generation_kpi_falls_when_a_report_fails(self):
        # The generation-rate KPI must be a real measurement, not a constant.
        with app.app_context():
            with patch("app.queue_aging_report", side_effect=RuntimeError("boom")):
                payload = analytics_payload({})
            self.assertIn("error", payload["queue_aging"])
            kpis = {item["kpi"]: item for item in payload["kpis"]}
            generation = kpis["Analytics report generation rate"]
            self.assertLess(generation["value"], 100.0)
            self.assertFalse(generation["met"])
            self.assertIn("Queue aging", generation["basis"])

    def test_residency_watchlist_flags_students_past_their_limit(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            # A master's programme carries a 5-year normal and 7-year absolute
            # residency limit; a doctoral programme 7 and 9.
            student.entry_year = date.today().year - 8
            db.session.commit()
            report = residency_watchlist_report([student])
            self.assertEqual(report["count"], 1)
            row = report["rows"][0]
            self.assertEqual(row["student_number"], "GS-2026-TEST")
            self.assertEqual(row["years_in_program"], 8)
            self.assertEqual(row["normal_years"], 5)
            self.assertEqual(row["absolute_years"], 7)
            self.assertEqual(row["residency_state"], "Exceeded absolute limit")

            # A student inside the normal period is counted but not listed.
            student.entry_year = date.today().year - 1
            db.session.commit()
            report = residency_watchlist_report([student])
            self.assertEqual(report["count"], 0)
            self.assertEqual(report["totals"]["Within normal period"], 1)

    def test_completion_and_attrition_rates_are_computed_per_program(self):
        with app.app_context():
            program_id = self.program_id
            entry = date.today().year - 3

            def make(number, stage, standing="Active"):
                item = Student(
                    student_number=number, first_name="A", last_name=number,
                    email=f"{number}@example.test", program_id=program_id,
                    entry_year=entry, current_stage=stage, standing=standing,
                )
                db.session.add(item)
                return item

            students = [
                make("C1", "Completed"), make("C2", "Completed"),
                make("W1", "Withdrawn"), make("A1", "AWOL", "AWOL"),
                make("L1", "LOA", "On Leave"), make("N1", "Coursework"),
            ]
            db.session.commit()
            report = completion_attrition_report(students)
            self.assertEqual(report["count"], 1)
            row = report["rows"][0]
            self.assertEqual(row["total_students"], 6)
            self.assertEqual(row["completed"], 2)
            self.assertEqual(row["withdrawn"], 1)
            self.assertEqual(row["awol"], 1)
            self.assertEqual(row["on_leave"], 1)
            self.assertEqual(row["active"], 1)
            self.assertEqual(row["completion_rate"], 33.3)
            # Attrition counts withdrawn plus AWOL; an approved leave is not
            # attrition.
            self.assertEqual(row["attrition_rate"], 33.3)
            self.assertEqual(row["average_years_to_finish"], 3)

    def test_curriculum_tagging_requires_a_rationale_when_versions_compete(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            program_id = self.program_id
            # Publish two curriculum versions for this programme.
            for year in ("2024-2025", "2025-2026"):
                db.session.add(CurriculumOffering(
                    program_id=program_id, course_id=self.course_id,
                    academic_year=year, semester="1st Semester"))
            db.session.commit()
            self.assertEqual(available_curriculum_versions(student), ["2024-2025", "2025-2026"])

            with self.assertRaises(ValueError) as ctx:
                tag_student_curriculum(student, "2025-2026", "", None)
            self.assertIn("Record the reason", str(ctx.exception))

            # An unpublished version is refused outright.
            with self.assertRaises(ValueError):
                tag_student_curriculum(student, "1999-2000", "because", None)

            tag = tag_student_curriculum(student, "2025-2026", "Entered under this curriculum.", None)
            db.session.commit()
            self.assertEqual(tag.curriculum_version, "2025-2026")
            self.assertTrue(tag.active)

            # Re-tagging retires the previous tag rather than leaving two active.
            tag_student_curriculum(student, "2024-2025", "Corrected after review.", None)
            db.session.commit()
            active = StudentCurriculumTag.query.filter_by(student_id=student.id, active=True).all()
            self.assertEqual(len(active), 1)
            self.assertEqual(active[0].curriculum_version, "2024-2025")

    def test_study_plan_draft_is_derived_from_the_curriculum_audit(self):
        with app.app_context():
            student = db.session.get(Student, self.student_id)
            draft = generate_study_plan_draft(student, "", None)
            db.session.commit()
            audit = compute_course_audit(student)
            # Everything proposed must be a subject the student still requires.
            outstanding = {row["course"].code for row in audit["missing"] + audit["incomplete"]}
            proposed = [c for c in (draft.subject_codes or "").split("|") if c]
            self.assertTrue(proposed, "expected at least one proposed subject")
            for code in proposed:
                self.assertIn(code, outstanding)
            self.assertEqual(draft.remaining_count, len(outstanding))
            self.assertEqual(draft.status, "Draft")

    def test_onboarding_report_is_blocked_until_the_checklist_passes(self):
        with app.app_context():
            upload = MonitoringSheetUpload(
                original_name="MAED.xlsx", stored_name="maed-test.xlsx", program_code="BPM",
                row_count=1, subject_count=1, snapshot_json="{}", result_json="{}")
            db.session.add(upload)
            db.session.flush()
            student = db.session.get(Student, self.student_id)
            student.monitoring_upload_id = upload.id
            # An unresolved validation issue must hold the gate shut.
            db.session.add(MonitoringValidationIssue(
                upload_id=upload.id, row_number=1, student_id=student.id,
                issue_type="Mismatch", issue_summary="Name differs from the source workbook.",
                status="Unresolved"))
            db.session.commit()

            review = build_onboarding_review(upload)
            db.session.commit()
            checklist = onboarding_checklist_state(upload)
            self.assertFalse(all(item["passed"] for item in checklist))
            self.assertEqual(review.unresolved_issue_count, 1)

            response = self._staff_client().post(f"/api/onboarding/reviews/{review.id}/submit")
            self.assertEqual(response.status_code, 400)
            self.assertIn("Validation issues resolved", response.get_json()["outstanding"])

    def test_onboarding_completion_requires_the_dean(self):
        with app.app_context():
            upload = MonitoringSheetUpload(
                original_name="MAED.xlsx", stored_name="maed-clean.xlsx", program_code="BPM",
                row_count=1, subject_count=1, snapshot_json="{}", result_json="{}")
            db.session.add(upload)
            db.session.flush()
            student = db.session.get(Student, self.student_id)
            student.monitoring_upload_id = upload.id
            db.session.commit()
            review = build_onboarding_review(upload)
            db.session.commit()
            review_id = review.id

            staff = self._staff_client()
            self.assertEqual(staff.post(f"/api/onboarding/reviews/{review_id}/submit").status_code, 200)
            # Staff cannot decide their own submission.
            self.assertEqual(
                staff.post(f"/api/onboarding/reviews/{review_id}/decision",
                           json={"decision": "approve"}).status_code, 403)

            dean = self._dean_client()
            # Returning requires a reason.
            self.assertEqual(
                dean.post(f"/api/onboarding/reviews/{review_id}/decision",
                          json={"decision": "return"}).status_code, 400)
            response = dean.post(f"/api/onboarding/reviews/{review_id}/decision",
                                 json={"decision": "approve", "remarks": "Verified."})
            self.assertEqual(response.status_code, 200)
            item = response.get_json()["item"]
            self.assertEqual(item["status"], "Completed")
            self.assertEqual(item["dean_decision"], "Approved")
            self.assertIsNotNone(item["admission_completed_at"])

    def test_coursework_report_reaches_the_dean_and_records_the_decision(self):
        with app.app_context():
            report = build_coursework_report(None, "AY 2026-2027 1st Semester", None)
            db.session.commit()
            report_id = report.id
            self.assertGreaterEqual(report.students_reviewed, 1)

            dean = self._dean_client()
            # The Dean cannot decide a report that has not been sent.
            self.assertEqual(
                dean.post(f"/api/coursework/reports/{report_id}/decision",
                          json={"decision": "acknowledge"}).status_code, 400)
            # Nor generate one.
            self.assertEqual(dean.post("/api/coursework/reports", json={}).status_code, 403)

            self.assertEqual(
                self._staff_client().post(f"/api/coursework/reports/{report_id}/send").status_code, 200)

            response = dean.post(f"/api/coursework/reports/{report_id}/decision",
                                 json={"decision": "acknowledge", "remarks": "Noted."})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.get_json()["item"]["dean_decision"], "Acknowledged")


if __name__ == "__main__":
    unittest.main()
