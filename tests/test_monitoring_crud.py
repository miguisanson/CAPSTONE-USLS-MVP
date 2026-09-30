"""Portal CRUD for the Monitoring Sheet (owner decision 2026-10-01).

Staff and coordinators must be able to add a student, add / edit / remove
subject rows and change statuses directly in the portal, without a workbook.
Every portal value carries provenance ("Manual entry", who, when, reason) and
an audit entry; deletes are soft; a later sheet import must never silently
destroy a portal edit.
"""
import os
import tempfile
import unittest
from datetime import date
from io import BytesIO

from openpyxl import load_workbook

_DB_FILE = tempfile.NamedTemporaryFile(suffix=".sqlite3", delete=False)
_DB_FILE.close()
os.environ["DATABASE_URL"] = f"sqlite:///{_DB_FILE.name}"

from app import (  # noqa: E402
    AcademicTerm,
    Course,
    CourseRecord,
    MonitoringEdit,
    MonitoringSheetUpload,
    MonitoringValidationIssue,
    Program,
    Student,
    TermEnrollment,
    TransactionLog,
    UserAccount,
    app,
    compute_course_audit,
    db,
    generate_password_hash,
    import_ac_monitoring,
    monitoring_curriculum_courses,
    parse_ac_monitoring,
)


class MonitoringPortalCrudTests(unittest.TestCase):
    @classmethod
    def tearDownClass(cls):
        try:
            with app.app_context():
                db.session.remove()
                db.engine.dispose()
            os.unlink(_DB_FILE.name)
        except (FileNotFoundError, PermissionError):
            pass

    def setUp(self):
        app.config.update(TESTING=True)
        with app.app_context():
            db.drop_all()
            db.create_all()
            program = Program(code="MPT", name="Monitoring Portal Test", college="Graduate School")
            other = Program(code="OTH", name="Other Program", college="Graduate School")
            db.session.add_all([program, other])
            db.session.flush()
            for code, title, category, units in (
                ("MPT-501", "Research Methods", "Basic", 3),
                ("MPT-502", "Statistics", "Basic", 3),
                ("MPT-503", "Seminar", "Major", 3),
            ):
                db.session.add(Course(program_id=program.id, code=code, title=title, category=category, units=units))
            db.session.add(Course(program_id=other.id, code="OTH-900", title="Elsewhere", category="Major", units=3))
            db.session.add(AcademicTerm(
                label="AY 2026-2027 1st Semester",
                start_date=date(2026, 8, 1),
                end_date=date(2026, 12, 15),
                is_active_planning_term=True,
            ))
            db.session.flush()
            self.program_id = program.id
            self.other_program_id = other.id
            self.accounts = {}
            for role in ("staff", "academic_coordinator", "research_coordinator", "dean", "admin"):
                self.accounts[role] = self._account(role, f"{role}@example.test")
            db.session.commit()
            self.accounts = {role: account.id for role, account in self.accounts.items()}

    def _account(self, role, email, student_id=None):
        account = UserAccount(
            email=email,
            full_name=f"Test {role}",
            password_hash=generate_password_hash("test-password"),
            role=role,
            student_id=student_id,
            active=True,
        )
        db.session.add(account)
        db.session.flush()
        return account

    def _client(self, role):
        client = app.test_client()
        with client.session_transaction() as session:
            session["account_id"] = self.accounts[role]
            session["role"] = role
        return client

    def _student_payload(self, **overrides):
        payload = {
            "student_number": "9000001",
            "first_name": "maria",
            "last_name": "dela cruz",
            "program_id": self.program_id,
            "academic_year_entry": "26-27",
        }
        payload.update(overrides)
        return payload

    def _create_student(self, **overrides):
        response = self._client("staff").post("/api/monitoring/students", json=self._student_payload(**overrides))
        self.assertEqual(response.status_code, 201, response.get_json())
        return response.get_json()

    def _course_id(self, code):
        with app.app_context():
            return Course.query.filter_by(code=code).one().id

    def _import_row(self, number, first, last, subjects, ay="26-27"):
        with app.app_context():
            upload = MonitoringSheetUpload(
                original_name="sheet.xlsx", stored_name=f"sheet-{number}-{len(subjects)}.xlsx",
                program_code="MPT", row_count=1, subject_count=len(subjects),
                snapshot_json="{}", result_json="{}",
            )
            db.session.add(upload)
            db.session.flush()
            parsed = {
                "program_code": "MPT",
                "subjects": list(subjects),
                "subject_categories": {code: "Basic" for code in subjects},
                "subject_titles": {code: code for code in subjects},
                "issues": [],
                "rows": [{
                    "row": 5, "idno": number, "first_name": first, "last_name": last,
                    "course": "MPT", "year": "1", "ay_entry": ay,
                    "subjects": dict(subjects), "milestones": {},
                    "comprehensive_exam_passed": False, "note": "",
                }],
            }
            result = import_ac_monitoring(parsed, upload=upload)
            db.session.commit()
            return result

    # -- create student ---------------------------------------------------
    def test_staff_can_create_a_student_in_the_portal_like_the_import_does(self):
        body = self._create_student(email="", adviser_name="Dr. Reyes")
        student = body["student"]
        self.assertEqual(student["student_number"], "9000001")
        self.assertEqual(student["first_name"], "Maria")
        self.assertEqual(student["last_name"], "Dela Cruz")
        self.assertEqual(student["academic_year_entry"], "26-27")
        self.assertEqual(student["source"]["label"], "Manual entry")
        self.assertTrue(body["account"]["email"])
        with app.app_context():
            row = Student.query.filter_by(student_number="9000001").one()
            self.assertEqual(row.entry_year, 2026)
            self.assertEqual(row.program_id, self.program_id)
            self.assertEqual(row.adviser_name, "Dr. Reyes")
            account = UserAccount.query.filter_by(student_id=row.id, role="student").one()
            self.assertEqual(account.email, row.email)
            self.assertEqual(
                TermEnrollment.query.filter_by(student_id=row.id).count(), 1,
                "an admission-term enrollment signal is created like the import does",
            )
            self.assertEqual(CourseRecord.query.filter_by(student_id=row.id).count(), 3)
            edit = MonitoringEdit.query.filter_by(student_id=row.id, action="create_student").one()
            self.assertEqual(edit.source, "Manual entry")
            self.assertEqual(edit.actor_name, "Test staff")
            log = TransactionLog.query.filter_by(student_id=row.id).order_by(TransactionLog.id.desc()).first()
            self.assertIn("portal", (log.source_reference + log.result).lower())

    def test_create_student_applies_the_import_validations(self):
        staff = self._client("staff")
        missing = staff.post("/api/monitoring/students", json={"program_id": self.program_id})
        self.assertEqual(missing.status_code, 400)
        message = missing.get_json()["error"]
        for label in ("Student ID", "First name", "Last name", "School Year / AY Entry"):
            self.assertIn(label, message)

        self._create_student()
        duplicate = staff.post("/api/monitoring/students", json=self._student_payload(first_name="Someone", last_name="Else"))
        self.assertEqual(duplicate.status_code, 409)
        self.assertIn("already belongs to", duplicate.get_json()["error"])

        unknown = staff.post("/api/monitoring/students", json=self._student_payload(student_number="9000002", program_id=9999))
        self.assertEqual(unknown.status_code, 400)
        self.assertIn("program", unknown.get_json()["error"].lower())

        same_person = staff.post("/api/monitoring/students", json=self._student_payload(student_number="9000003"))
        self.assertEqual(same_person.status_code, 409, same_person.get_json())
        self.assertEqual(same_person.get_json()["possible_match"]["student_number"], "9000001")
        forced = staff.post("/api/monitoring/students", json=self._student_payload(
            student_number="9000003", confirm_possible_duplicate=True))
        self.assertEqual(forced.status_code, 201, forced.get_json())

    # -- who may do what --------------------------------------------------
    def test_write_roles_read_only_roles_and_students(self):
        body = self._create_student()
        student_id = body["student"]["id"]
        course_id = self._course_id("MPT-501")
        subject_payload = {"course_id": course_id, "status": "Planned"}

        for role in ("research_coordinator", "dean"):
            client = self._client(role)
            self.assertEqual(client.get(f"/api/monitoring/students/{student_id}").status_code, 200, role)
            self.assertEqual(client.get(f"/api/monitoring/students/{student_id}/history").status_code, 200, role)
            self.assertEqual(client.get(f"/api/monitoring/grid?program_id={self.program_id}").status_code, 200, role)
            self.assertEqual(client.post("/api/monitoring/students", json=self._student_payload(student_number="9000009")).status_code, 403, role)
            self.assertEqual(client.post(f"/api/monitoring/students/{student_id}/subjects", json=subject_payload).status_code, 403, role)
            self.assertEqual(client.patch(f"/api/monitoring/students/{student_id}", json={"first_name": "X", "reason": "test"}).status_code, 403, role)

        for role in ("academic_coordinator", "admin"):
            response = self._client(role).post(f"/api/monitoring/students/{student_id}/subjects", json={
                "course_id": self._course_id("MPT-502" if role == "admin" else "MPT-501"), "status": "Planned"})
            self.assertEqual(response.status_code, 201, (role, response.get_json()))

        with app.app_context():
            account = self._account("student", "student-own@example.test", student_id=student_id)
            other_student = Student(
                student_number="9000777", first_name="Other", last_name="Person", email="other@example.test",
                program_id=self.program_id, entry_year=2026, academic_year_entry="26-27",
            )
            db.session.add(other_student)
            db.session.commit()
            self.accounts["student"] = account.id
            other_id = other_student.id
        student = self._client("student")
        self.assertEqual(student.get(f"/api/monitoring/students/{student_id}").status_code, 403)
        self.assertEqual(student.post(f"/api/monitoring/students/{student_id}/subjects", json=subject_payload).status_code, 403)
        own = student.get("/api/student-portal/monitoring")
        self.assertEqual(own.status_code, 200, own.get_json())
        self.assertEqual(own.get_json()["student"]["id"], student_id)
        self.assertNotEqual(own.get_json()["student"]["id"], other_id)
        self.assertFalse(own.get_json()["permissions"]["can_edit"])
        self.assertEqual(student.post("/api/student-portal/monitoring", json=subject_payload).status_code, 405)

    # -- subject rows -----------------------------------------------------
    def test_add_subject_row_from_the_curriculum_stamps_provenance_and_audit(self):
        student_id = self._create_student()["student"]["id"]
        course_id = self._course_id("MPT-501")
        staff = self._client("staff")
        response = staff.post(f"/api/monitoring/students/{student_id}/subjects", json={
            "course_id": course_id, "status": "Enrolled",
            "term_label": "AY 2026-2027 1st Semester", "remarks": "Section A",
        })
        self.assertEqual(response.status_code, 201, response.get_json())
        subject = response.get_json()["subject"]
        self.assertEqual(subject["code"], "MPT-501")
        self.assertEqual(subject["status"], "Enrolled")
        self.assertEqual(subject["source"], "Manual entry")
        self.assertEqual(subject["source_by"], "Test staff")
        self.assertTrue(subject["source_at"])
        self.assertTrue(subject["in_curriculum"])

        with app.app_context():
            record = CourseRecord.query.filter_by(student_id=student_id, course_id=course_id).one()
            self.assertEqual(record.status, "Enrolled")
            self.assertEqual(record.term_label, "AY 2026-2027 1st Semester")
            self.assertEqual(record.source, "Manual entry")
            changes = MonitoringEdit.query.filter_by(student_id=student_id, course_id=course_id).all()
            self.assertTrue(any(c.field == "status" and c.old_value == "Missing" and c.new_value == "Enrolled" for c in changes))
            log = TransactionLog.query.filter_by(student_id=student_id).order_by(TransactionLog.id.desc()).first()
            self.assertIn("MPT-501", log.notes)
            self.assertIn("Missing", log.notes)
            self.assertIn("Enrolled", log.notes)

        again = staff.post(f"/api/monitoring/students/{student_id}/subjects", json={"course_id": course_id, "status": "Completed"})
        self.assertEqual(again.status_code, 409, "the row already exists - edit it instead")
        bad = staff.post(f"/api/monitoring/students/{student_id}/subjects", json={
            "course_id": self._course_id("MPT-502"), "status": "Graded 1.25"})
        self.assertEqual(bad.status_code, 400)
        foreign = staff.post(f"/api/monitoring/students/{student_id}/subjects", json={
            "course_id": self._course_id("OTH-900"), "status": "Planned"})
        self.assertEqual(foreign.status_code, 400)

    def test_free_form_subject_row_is_kept_but_never_counts_toward_the_curriculum(self):
        student_id = self._create_student()["student"]["id"]
        staff = self._client("staff")
        missing_details = staff.post(f"/api/monitoring/students/{student_id}/subjects", json={"code": "ELEC-1", "status": "Completed"})
        self.assertEqual(missing_details.status_code, 400)
        response = staff.post(f"/api/monitoring/students/{student_id}/subjects", json={
            "code": "ELEC-1", "title": "Special Elective", "units": 2, "status": "Completed",
            "term_label": "AY 2025-2026 2nd Semester"})
        self.assertEqual(response.status_code, 201, response.get_json())
        subject = response.get_json()["subject"]
        self.assertFalse(subject["in_curriculum"])
        self.assertEqual(subject["units"], 2)
        with app.app_context():
            student = Student.query.get(student_id)
            self.assertNotIn("ELEC-1", [c.code for c in monitoring_curriculum_courses(student.program)])
            self.assertEqual(compute_course_audit(student)["completed_units"], 0)
        detail = staff.get(f"/api/monitoring/students/{student_id}").get_json()
        self.assertIn("ELEC-1", [row["code"] for row in detail["subjects"]])

    def test_editing_an_imported_value_needs_a_reason_and_flips_provenance(self):
        self._import_row("9100001", "Import", "Student", {"MPT-501": True, "MPT-502": False})
        with app.app_context():
            student_id = Student.query.filter_by(student_number="9100001").one().id
            record = CourseRecord.query.filter_by(student_id=student_id, course_id=self._course_id("MPT-501")).one()
            record_id = record.id
            self.assertEqual(record.status, "Completed")
        staff = self._client("staff")
        detail = staff.get(f"/api/monitoring/students/{student_id}").get_json()
        row = next(item for item in detail["subjects"] if item["code"] == "MPT-501")
        self.assertEqual(row["source"], "Imported")

        url = f"/api/monitoring/students/{student_id}/subjects/{record_id}"
        no_reason = staff.patch(url, json={"status": "Withdrawn"})
        self.assertEqual(no_reason.status_code, 400)
        self.assertIn("reason", no_reason.get_json()["error"].lower())

        ok = staff.patch(url, json={"status": "INC", "remarks": "Missing final paper", "reason": "Registrar email 2026-10-01"})
        self.assertEqual(ok.status_code, 200, ok.get_json())
        self.assertEqual(ok.get_json()["subject"]["status"], "INC")
        self.assertEqual(ok.get_json()["subject"]["source"], "Manual entry")
        self.assertEqual(ok.get_json()["subject"]["source_reason"], "Registrar email 2026-10-01")
        with app.app_context():
            edits = MonitoringEdit.query.filter_by(student_id=student_id, field="status").all()
            change = next(e for e in edits if e.new_value == "INC")
            self.assertEqual(change.old_value, "Completed")
            self.assertEqual(change.reason, "Registrar email 2026-10-01")
            # INC is not "completed": the audit must agree.
            audit = compute_course_audit(Student.query.get(student_id))
            self.assertEqual(audit["completed_units"], 0)
        history = staff.get(f"/api/monitoring/students/{student_id}/history").get_json()["items"]
        self.assertTrue(any(item["new_value"] == "INC" and item["old_value"] == "Completed" for item in history))

        unchanged = staff.patch(url, json={"status": "INC", "reason": "again"})
        self.assertEqual(unchanged.status_code, 200)
        self.assertFalse(unchanged.get_json()["changed"], "a no-op edit records nothing")

    def test_delete_is_soft_and_needs_a_reason(self):
        student_id = self._create_student()["student"]["id"]
        course_id = self._course_id("MPT-503")
        staff = self._client("staff")
        created = staff.post(f"/api/monitoring/students/{student_id}/subjects", json={"course_id": course_id, "status": "Completed"})
        record_id = created.get_json()["subject"]["record_id"]
        url = f"/api/monitoring/students/{student_id}/subjects/{record_id}"

        self.assertEqual(staff.delete(url).status_code, 400)
        removed = staff.delete(url, json={"reason": "Entered against the wrong student"})
        self.assertEqual(removed.status_code, 200, removed.get_json())

        with app.app_context():
            record = db.session.get(CourseRecord, record_id)
            self.assertIsNotNone(record, "the row is kept")
            self.assertIsNotNone(record.removed_at)
            self.assertEqual(record.removal_reason, "Entered against the wrong student")
            self.assertEqual(record.status, "Missing")
            self.assertEqual(compute_course_audit(Student.query.get(student_id))["completed_units"], 0)
            self.assertTrue(MonitoringEdit.query.filter_by(student_id=student_id, action="remove_subject").count())
        detail = staff.get(f"/api/monitoring/students/{student_id}").get_json()
        removed_rows = [row for row in detail["subjects"] if row["removed"]]
        self.assertEqual([row["code"] for row in removed_rows], ["MPT-503"])
        self.assertEqual(removed_rows[0]["removed_status"], "Completed")

        self.assertEqual(staff.delete(url, json={"reason": "again"}).status_code, 409)
        readded = staff.post(f"/api/monitoring/students/{student_id}/subjects", json={"course_id": course_id, "status": "Planned"})
        self.assertEqual(readded.status_code, 201, readded.get_json())
        with app.app_context():
            self.assertIsNone(db.session.get(CourseRecord, record_id).removed_at)

    def test_student_header_edit_needs_a_reason_and_records_old_and_new(self):
        first = self._create_student()["student"]["id"]
        self._create_student(student_number="9000050", first_name="Jose", last_name="Rizal")
        staff = self._client("staff")
        url = f"/api/monitoring/students/{first}"
        self.assertEqual(staff.patch(url, json={"first_name": "Marie"}).status_code, 400)
        clash = staff.patch(url, json={"student_number": "9000050", "reason": "typo"})
        self.assertEqual(clash.status_code, 409)
        blank = staff.patch(url, json={"last_name": "  ", "reason": "typo"})
        self.assertEqual(blank.status_code, 400)

        ok = staff.patch(url, json={
            "first_name": "marie", "academic_year_entry": "25-26", "email": "marie@example.test",
            "reason": "Corrected against the AIMS printout"})
        self.assertEqual(ok.status_code, 200, ok.get_json())
        with app.app_context():
            student = Student.query.get(first)
            self.assertEqual(student.first_name, "Marie")
            self.assertEqual(student.entry_year, 2025)
            self.assertEqual(student.academic_year_entry, "25-26")
            self.assertEqual(student.email, "marie@example.test")
            edits = {e.field: e for e in MonitoringEdit.query.filter_by(student_id=first, action="update_student")}
            self.assertEqual((edits["first_name"].old_value, edits["first_name"].new_value), ("Maria", "Marie"))
            self.assertEqual((edits["academic_year_entry"].old_value, edits["academic_year_entry"].new_value), ("26-27", "25-26"))
            self.assertEqual(edits["email"].reason, "Corrected against the AIMS printout")

    def test_bulk_add_remaining_curriculum_subjects(self):
        student_id = self._create_student()["student"]["id"]
        staff = self._client("staff")
        done = staff.post(f"/api/monitoring/students/{student_id}/subjects", json={
            "course_id": self._course_id("MPT-501"), "status": "Completed"})
        gone = staff.post(f"/api/monitoring/students/{student_id}/subjects", json={
            "course_id": self._course_id("MPT-503"), "status": "Planned"}).get_json()["subject"]
        staff.delete(f"/api/monitoring/students/{student_id}/subjects/{gone['record_id']}", json={"reason": "not for this student"})
        self.assertEqual(done.status_code, 201)

        response = staff.post("/api/monitoring/subjects/bulk-add-remaining", json={
            "program_id": self.program_id, "student_ids": [student_id], "status": "Planned",
            "term_label": "AY 2026-2027 2nd Semester"})
        self.assertEqual(response.status_code, 200, response.get_json())
        self.assertEqual(response.get_json()["added"], 1, "only MPT-502: 501 is done, 503 was deliberately removed")
        with app.app_context():
            rows = {r.course.code: r for r in CourseRecord.query.filter_by(student_id=student_id)}
            self.assertEqual(rows["MPT-502"].status, "Planned")
            self.assertEqual(rows["MPT-502"].source, "Manual entry")
            self.assertEqual(rows["MPT-501"].status, "Completed")
            self.assertEqual(rows["MPT-503"].status, "Missing")
        again = staff.post("/api/monitoring/subjects/bulk-add-remaining", json={
            "program_id": self.program_id, "student_ids": [student_id]})
        self.assertEqual(again.get_json()["added"], 0)
        # program-wide (no student_ids) covers every active student of the program
        self._create_student(student_number="9000060", first_name="Ana", last_name="Lopez")
        program_wide = staff.post("/api/monitoring/subjects/bulk-add-remaining", json={"program_id": self.program_id})
        self.assertEqual(program_wide.get_json()["added"], 3)

    # -- import must not destroy portal edits -----------------------------
    def test_later_import_keeps_portal_edit_and_raises_an_issue_for_staff(self):
        student_id = self._create_student(student_number="9200001", first_name="Kept", last_name="Edit")["student"]["id"]
        course_id = self._course_id("MPT-501")
        staff = self._client("staff")
        staff.post(f"/api/monitoring/students/{student_id}/subjects", json={
            "course_id": course_id, "status": "Enrolled", "term_label": "AY 2026-2027 1st Semester"})

        result = self._import_row("9200001", "Kept", "Edit", {"MPT-501": True, "MPT-502": True})
        self.assertEqual(result["conflict_count"], 1, result)
        with app.app_context():
            record = CourseRecord.query.filter_by(student_id=student_id, course_id=course_id).one()
            self.assertEqual(record.status, "Enrolled", "portal value survives the import")
            self.assertEqual(record.source, "Manual entry")
            issue = MonitoringValidationIssue.query.filter_by(student_id=student_id).one()
            self.assertEqual(issue.status, "Unresolved")
            self.assertIn("Portal edit conflict", issue.issue_type)
            self.assertIn("MPT-501", issue.issue_summary)
            self.assertIn("Enrolled", issue.existing_json, "both facts are kept on the issue")
            self.assertIn("true", issue.uploaded_json.lower())
            issue_id = issue.id

        resolved = staff.patch(f"/api/monitoring/issues/{issue_id}/resolve", json={"action": "use_uploaded"})
        self.assertEqual(resolved.status_code, 200, resolved.get_json())
        with app.app_context():
            record = CourseRecord.query.filter_by(student_id=student_id, course_id=course_id).one()
            self.assertEqual(record.status, "Completed")
            self.assertEqual(record.source, "Imported", "provenance follows the applied value")
            self.assertTrue(
                MonitoringEdit.query.filter_by(student_id=student_id, field="status", old_value="Enrolled", new_value="Completed").count(),
                "overwriting a portal value through the issue is itself recorded",
            )

    def test_import_that_does_not_contradict_the_portal_still_applies(self):
        student_id = self._create_student(student_number="9200002", first_name="Calm", last_name="Sheet")["student"]["id"]
        self._client("staff").post(f"/api/monitoring/students/{student_id}/subjects", json={
            "course_id": self._course_id("MPT-501"), "status": "Enrolled"})
        # the sheet is silent about MPT-501 and only adds MPT-502
        result = self._import_row("9200002", "Calm", "Sheet", {"MPT-501": False, "MPT-502": True})
        self.assertEqual(result["conflict_count"], 0, result)
        with app.app_context():
            rows = {r.course.code: r for r in CourseRecord.query.filter_by(student_id=student_id)}
            self.assertEqual(rows["MPT-501"].status, "Enrolled")
            self.assertEqual(rows["MPT-502"].status, "Completed")
            self.assertEqual(rows["MPT-502"].source, "Imported")

    # -- export -----------------------------------------------------------
    def test_workbook_export_includes_portal_entered_students_and_rows(self):
        student_id = self._create_student(first_name="Portal", last_name="Only")["student"]["id"]
        staff = self._client("staff")
        staff.post(f"/api/monitoring/students/{student_id}/subjects", json={"course_id": self._course_id("MPT-501"), "status": "Completed"})
        staff.post(f"/api/monitoring/students/{student_id}/subjects", json={"course_id": self._course_id("MPT-502"), "status": "Planned"})
        staff.post(f"/api/monitoring/students/{student_id}/subjects", json={
            "code": "ELEC-1", "title": "Special Elective", "units": 2, "status": "Completed"})
        self._import_row("9300001", "Sheet", "Student", {"MPT-503": True})

        response = self._client("dean").get(f"/api/monitoring/export?program_id={self.program_id}")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response.headers["Content-Type"])
        workbook = load_workbook(BytesIO(response.data), data_only=True)

        # Sheet 1 is the AC template, so the export can be uploaded straight back.
        parsed = parse_ac_monitoring(BytesIO(response.data))
        self.assertEqual(parsed["program_code"], "MPT")
        rows = {row["idno"]: row for row in parsed["rows"]}
        self.assertEqual(set(rows), {"9000001", "9300001"})
        self.assertTrue(rows["9000001"]["subjects"]["MPT-501"])
        self.assertFalse(rows["9000001"]["subjects"]["MPT-502"], "Planned is not completed")
        self.assertTrue(rows["9300001"]["subjects"]["MPT-503"])

        # A second sheet carries every subject row with its provenance.
        detail = workbook["Subject rows"]
        header = [cell.value for cell in detail[1]]
        for column in ("Student ID", "Subject code", "Status", "Source", "Entered by", "Entered at", "Reason"):
            self.assertIn(column, header)
        records = [
            {header[index]: cell.value for index, cell in enumerate(row)}
            for row in detail.iter_rows(min_row=2)
        ]
        by_code = {(r["Student ID"], r["Subject code"]): r for r in records}
        self.assertEqual(by_code[("9000001", "MPT-502")]["Status"], "Planned")
        self.assertEqual(by_code[("9000001", "MPT-502")]["Source"], "Manual entry")
        self.assertEqual(by_code[("9000001", "ELEC-1")]["Status"], "Completed")
        self.assertEqual(by_code[("9300001", "MPT-503")]["Source"], "Imported")

    # -- downstream readers ----------------------------------------------
    def test_a_portal_only_student_flows_into_enrollment_reports_and_the_student_portal(self):
        body = self._create_student(first_name="Flow", last_name="Through")
        student_id = body["student"]["id"]
        email = body["account"]["email"]
        staff = self._client("staff")
        staff.post(f"/api/monitoring/students/{student_id}/subjects", json={"course_id": self._course_id("MPT-501"), "status": "Completed"})
        staff.post(f"/api/monitoring/students/{student_id}/subjects", json={"course_id": self._course_id("MPT-502"), "status": "Planned"})
        staff.post(f"/api/monitoring/students/{student_id}/subjects", json={"course_id": self._course_id("MPT-503"), "status": "INC"})

        with app.app_context():
            term_id = AcademicTerm.query.one().id
        enrollment = staff.get(f"/api/enrollment?program_id={self.program_id}&student_id={student_id}&term_id={term_id}")
        self.assertEqual(enrollment.status_code, 200, enrollment.get_json())
        states = {item["code"]: item for item in enrollment.get_json()["curriculum_subjects"]}
        self.assertEqual(states["MPT-501"]["enrollment_state"], "completed")
        self.assertFalse(states["MPT-501"]["selectable"], "a completed subject cannot be enrolled again")
        self.assertEqual(states["MPT-502"]["enrollment_state"], "available", "Planned is not taken yet")
        self.assertTrue(states["MPT-502"]["selectable"])
        self.assertEqual(states["MPT-503"]["enrollment_state"], "requires_resolution", "INC must be resolved first")
        self.assertIn(student_id, [item["id"] for item in enrollment.get_json()["students"]])

        grid = staff.get(f"/api/monitoring/grid?program_id={self.program_id}").get_json()
        row = next(item for item in grid["students"] if item["id"] == student_id)
        course_ids = {code: self._course_id(code) for code in ("MPT-501", "MPT-502", "MPT-503")}
        self.assertEqual(row["cells"][str(course_ids["MPT-501"])], "Completed")
        self.assertEqual(row["completed"], 1)

        listing = staff.get("/api/students?q=Through").get_json()
        self.assertIn(student_id, [item["id"] for item in listing["items"]])
        reports = staff.get("/api/reports")
        self.assertEqual(reports.status_code, 200)
        self.assertIn("Through", str(reports.get_json()) + str(staff.get(f"/api/students/{student_id}").get_json()))

        with app.app_context():
            account = UserAccount.query.filter_by(email=email).one()
            self.accounts["student"] = account.id
        portal = self._client("student").get("/api/student-portal/context")
        self.assertEqual(portal.status_code, 200, portal.get_json())
        records = {item["code"]: item for item in portal.get_json()["course_records"]}
        self.assertEqual(records["MPT-501"]["status"], "Completed")
        self.assertEqual(records["MPT-501"]["source"], "Manual entry")


if __name__ == "__main__":
    unittest.main()
