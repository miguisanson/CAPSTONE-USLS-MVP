"""Normalize every program monitoring workbook to the USLS sample layout.

The supplied AC monitoring template uses:
- a seven-digit numeric IDNO (for example, 1860101);
- AY ENTRY on every student row; and
- a separate YR column.

The generated demo workbooks previously used ``GS-2026-####`` identifiers and
blank repeated AY cells. This script converts only that known legacy ID format,
fills down AY ENTRY, preserves the existing YR value, and validates every row.
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK_ROOT = PROJECT_ROOT / "Documents" / "Monitoring_Sheets"
LEGACY_ID = re.compile(r"GS-2026-(\d{4})", re.IGNORECASE)
SEVEN_DIGIT_ID = re.compile(r"\d{7}")


def text(value) -> str:
    return "" if value is None else str(value).strip()


def find_column(sheet, row: int, label: str) -> int | None:
    expected = label.strip().upper()
    for column in range(1, sheet.max_column + 1):
        if text(sheet.cell(row, column).value).upper() == expected:
            return column
    return None


def academic_year_start(ay_entry: str) -> int:
    match = re.search(r"(\d{2,4})", ay_entry)
    if not match:
        raise ValueError(f"AY ENTRY {ay_entry!r} does not contain a year")
    year = int(match.group(1))
    return year % 100


def normalized_idno(current, ay_entry: str) -> int:
    value = text(current)
    legacy = LEGACY_ID.fullmatch(value)
    if legacy:
        # The official sample's shape is YY + 6 + four-digit sequence.
        # Example: AY 2022-2023 and legacy serial 0007 becomes 2260007.
        return int(f"{academic_year_start(ay_entry):02d}6{legacy.group(1)}")
    if SEVEN_DIGIT_ID.fullmatch(value):
        return int(value)
    raise ValueError(
        f"IDNO {value!r} is neither the known legacy format nor a seven-digit numeric ID"
    )


def update_workbook(path: Path) -> tuple[int, dict[str, str]]:
    workbook = load_workbook(path)
    sheet = workbook.active

    header_row = None
    for row in range(1, min(sheet.max_row, 15) + 1):
        if find_column(sheet, row, "IDNO") and find_column(sheet, row, "YR"):
            header_row = row
            break
    if not header_row:
        raise ValueError(f"{path}: could not find the IDNO/YR header row")

    subheader_row = header_row + 1
    idno_column = find_column(sheet, header_row, "IDNO")
    year_level_column = find_column(sheet, header_row, "YR")
    ay_entry_column = find_column(sheet, subheader_row, "AY ENTRY")
    surname_column = find_column(sheet, subheader_row, "SN")
    first_name_column = find_column(sheet, subheader_row, "FN")
    if not all(
        [idno_column, year_level_column, ay_entry_column, surname_column, first_name_column]
    ):
        raise ValueError(f"{path}: incomplete AC monitoring headers")

    last_ay_entry = ""
    changed_ids: dict[str, str] = {}
    student_rows = 0
    for row in range(subheader_row + 1, sheet.max_row + 1):
        current_id = text(sheet.cell(row, idno_column).value)
        if not current_id:
            continue
        surname = text(sheet.cell(row, surname_column).value)
        first_name = text(sheet.cell(row, first_name_column).value)
        if not surname and not first_name:
            continue

        ay_entry = text(sheet.cell(row, ay_entry_column).value)
        if ay_entry:
            last_ay_entry = ay_entry
        elif last_ay_entry:
            ay_entry = last_ay_entry
            sheet.cell(row, ay_entry_column).value = ay_entry
        else:
            raise ValueError(f"{path}: row {row} has no AY ENTRY to fill down")

        year_level = text(sheet.cell(row, year_level_column).value)
        if not year_level:
            raise ValueError(f"{path}: row {row} has no YR value")

        new_id = normalized_idno(current_id, ay_entry)
        if str(new_id) != current_id:
            changed_ids[current_id] = str(new_id)
        id_cell = sheet.cell(row, idno_column)
        id_cell.value = new_id
        id_cell.number_format = "0"
        student_rows += 1

    workbook.save(path)
    return student_rows, changed_ids


def main() -> None:
    paths = sorted(
        path
        for path in WORKBOOK_ROOT.rglob("*.xlsx")
        if not path.name.startswith("~$")
    )
    if not paths:
        raise SystemExit(f"No monitoring workbooks found under {WORKBOOK_ROOT}")

    total_rows = 0
    total_ids = 0
    all_ids: set[str] = set()
    for path in paths:
        rows, changed = update_workbook(path)
        total_rows += rows
        total_ids += len(changed)

        # Reload saved output and validate what Excel will actually read.
        workbook = load_workbook(path, data_only=True, read_only=True)
        sheet = workbook.active
        header_row = next(
            row
            for row in range(1, min(sheet.max_row, 15) + 1)
            if find_column(sheet, row, "IDNO")
        )
        idno_column = find_column(sheet, header_row, "IDNO")
        year_level_column = find_column(sheet, header_row, "YR")
        ay_entry_column = find_column(sheet, header_row + 1, "AY ENTRY")
        for row in range(header_row + 2, sheet.max_row + 1):
            value = text(sheet.cell(row, idno_column).value)
            if not value:
                continue
            if not SEVEN_DIGIT_ID.fullmatch(value):
                raise ValueError(f"{path}: row {row} saved invalid IDNO {value!r}")
            if value in all_ids:
                raise ValueError(f"Duplicate IDNO {value} across monitoring workbooks")
            all_ids.add(value)
            if not text(sheet.cell(row, ay_entry_column).value):
                raise ValueError(f"{path}: row {row} saved without AY ENTRY")
            if not text(sheet.cell(row, year_level_column).value):
                raise ValueError(f"{path}: row {row} saved without YR")

        relative = path.relative_to(PROJECT_ROOT)
        print(f"{relative}: {rows} student row(s), {len(changed)} IDNO value(s) normalized")

    print(
        f"Updated {len(paths)} workbook(s): {total_rows} student row(s), "
        f"{total_ids} legacy IDNO value(s), {len(all_ids)} unique seven-digit IDs."
    )


if __name__ == "__main__":
    main()
