"""Case 05b - Student Handoff: import an Excel monitoring sheet, result, upload history, re-import is safe."""
from __future__ import annotations

import re
import tempfile
import time
from pathlib import Path

from .common import Runtime
from .helpers_a1 import asc, goto_nav, scroll_to, scroll_top, wait_loaded, safe_notes

SLUG = "student-handoff"
SAMPLE = "Documents/Monitoring_Sheets/Ready_To_Import/MSN_Monitoring_Sheet.xlsx"


def s(n: int) -> str:
    return f"{SLUG}-{n:02d}"


def build_workbook(root: Path, stamp: str, base_id: int) -> Path:
    """Copy the MSN sample sheet but keep only two NEW students (time-suffixed names, new IDNOs)."""
    import openpyxl

    wb = openpyxl.load_workbook(root / SAMPLE)
    ws = wb.active
    ws.delete_rows(7, ws.max_row)
    for row, (given, offset) in zip((5, 6), (("ALMA", 0), ("BEN", 1))):
        ws.cell(row, 3).value = f"WALKTHRU{stamp}"
        ws.cell(row, 4).value = given
        ws.cell(row, 5).value = base_id + offset
    out = Path(tempfile.gettempdir()) / f"MSN_Monitoring_Sheet_walkthrough_{stamp}.xlsx"
    wb.save(out)
    return out


def run(rt: Runtime) -> None:
    safe_notes(rt)
    root = Path(__file__).resolve().parents[4]
    stamp = time.strftime("%H%M%S")
    base_id = int("26" + time.strftime("%H%M%S"))
    workbook = build_workbook(root, stamp, base_id)
    rt.note(f"workbook {workbook.name}: students WALKTHRU{stamp}, ALMA and BEN, IDNO {base_id} and {base_id + 1}")
    page = rt.new_page("staff@usls.edu.ph")

    def open_handoff():
        goto_nav(rt, page, "Student Handoff", "/workflow/student-handoff")
        page.get_by_text("Import from monitoring sheet").first.wait_for()
        rt.shot(page, s(1), full=True)
    rt.step("open Student Handoff", open_handoff)

    def choose_file():
        page.locator("input[type=file]").first.set_input_files(str(workbook))
        page.wait_for_timeout(800)
        page.get_by_text("Import from monitoring sheet").first.scroll_into_view_if_needed()
        rt.shot(page, s(2))
    rt.step("choose the workbook", choose_file)

    def do_import():
        page.get_by_role("button", name=re.compile(r"^Import students", re.I)).click()
        page.get_by_text(re.compile(r"Imported \d+ new student")).first.wait_for(timeout=60000)
        wait_loaded(page)
        rt.note("import result: " + asc(page.get_by_text(re.compile(r"Imported \d+ new student")).first.inner_text()))
        scroll_to(page.get_by_text(re.compile(r"Imported \d+ new student")).first, 140)
        rt.shot(page, s(3))
    rt.step("import students", do_import)

    def history():
        page.get_by_role("button", name=re.compile("View upload history", re.I)).click()
        page.wait_for_timeout(1200)
        rt.note("history: " + asc(page.locator("main").inner_text().split("Upload history")[1][:700].replace("\n", " | ")))
        scroll_to(page.get_by_text("Upload history", exact=True).first, 90)
        rt.shot(page, s(4))
    rt.step("open Upload history", history)

    def onboarding():
        scroll_to(page.get_by_text(re.compile("Onboarding report & Dean approval")).first, 90)
        rt.note("onboarding: " + asc(page.locator("main").inner_text().split("Onboarding report")[1][:900].replace("\n", " | ")))
        rt.shot(page, s(5))
    rt.step("onboarding report for the Dean (not sent)", onboarding)

    def reimport():
        rt.go(page, "/workflow/student-handoff")
        wait_loaded(page)
        page.locator("input[type=file]").first.set_input_files(str(workbook))
        page.wait_for_timeout(800)
        page.get_by_role("button", name=re.compile(r"^Import students", re.I)).click()
        page.wait_for_timeout(2500)
        page.get_by_text(re.compile(r"Imported \d+ new student")).first.wait_for(timeout=60000)
        wait_loaded(page)
        rt.note("re-import result: " + asc(page.get_by_text(re.compile(r"Imported \d+ new student")).first.inner_text()))
        scroll_to(page.get_by_text(re.compile(r"Imported \d+ new student")).first, 140)
        rt.shot(page, s(6))
    rt.step("import the same workbook again", reimport)

    def monitoring():
        try:
            page.get_by_text(re.compile(r"Open Monitoring Sheet", re.I)).first.click(timeout=8000)
        except Exception:  # noqa: BLE001
            goto_nav(rt, page, "Monitoring Sheet", "/monitoring-sheet")
            sel = page.get_by_label("Program")
            sel.select_option(label=[o for o in sel.locator("option").all_inner_texts() if o.startswith("MSN")][0])
        wait_loaded(page)
        page.get_by_role("button", name=re.compile(f"monitoring record of Walkthru{stamp}, Alma", re.I)).first.wait_for()
        page.get_by_role("button", name=re.compile(f"monitoring record of Walkthru{stamp}, Alma", re.I)).first.scroll_into_view_if_needed()
        page.wait_for_timeout(500)
        rt.note("url after Open Monitoring Sheet: " + page.url)
        rt.shot(page, s(7))
        page.get_by_role("button", name=re.compile(f"monitoring record of Walkthru{stamp}, Alma", re.I)).first.click()
        d = page.get_by_role("dialog")
        d.get_by_text(re.compile("curriculum subjects completed")).first.wait_for()
        page.wait_for_timeout(600)
        rt.note("imported record: " + asc(d.inner_text()[:500].replace("\n", " | ")))
        rt.shot(page, s(8))
        page.keyboard.press("Escape")
    rt.step("open the imported students on the Monitoring Sheet", monitoring)

    page.context.close()
