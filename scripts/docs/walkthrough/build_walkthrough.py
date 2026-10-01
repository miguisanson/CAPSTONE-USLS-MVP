"""Build the demonstration walkthrough .docx from content + screenshots.

    E:/Temp/claude/wt-venv/Scripts/python.exe scripts/docs/walkthrough/build_walkthrough.py
    (needs python-docx; PyMuPDF + LibreOffice are used once to read page numbers for the contents table)

Inputs : front.json (front and back matter), cases/NN_*.json (one per case, in
         running order), shots/*.png (from capture_walkthrough.py)
Output : Documents/CAPSTONE_ONLY/CAP-2521-IT-WALKTHROUGH.docx
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
REPO = HERE.parent.parent.parent
OUT = REPO / "Documents" / "CAPSTONE_ONLY" / "CAP-2521-IT-WALKTHROUGH.docx"
SHOTS = HERE / "shots"
GREEN = RGBColor(0x0F, 0x6B, 0x3E)
GREY = RGBColor(0x55, 0x5B, 0x63)
SHOT_WIDTH = 5.4  # inches; --shot-width overrides
SOFFICE = Path("E:/LibreOffice/program/soffice.exe")


# ---------------------------------------------------------------- low-level helpers
def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def fld_char(kind: str):
    el = OxmlElement("w:fldChar")
    el.set(qn("w:fldCharType"), kind)
    return el


def instr_text(text: str):
    el = OxmlElement("w:instrText")
    el.set(qn("xml:space"), "preserve")
    el.text = f" {text} "
    return el


def add_field(paragraph, instruction: str, placeholder: str = "1") -> None:
    for child in (fld_char("begin"), instr_text(instruction), fld_char("separate")):
        paragraph.add_run()._r.append(child)
    paragraph.add_run(placeholder)
    paragraph.add_run()._r.append(fld_char("end"))


def set_cell_margins(table, top=50, bottom=50, left=90, right=90) -> None:
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def repeat_header(row) -> None:
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


def no_split(row) -> None:
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(el)


CACHE = Path("E:/Temp/claude/walkthrough/img_cache")


def compact(path: Path) -> Path:
    """Embed a palette-quantised copy (UI screenshots shrink about 3x); originals stay in shots/.
    Falls back to the original when Pillow is missing."""
    try:
        from PIL import Image
    except ImportError:
        return path
    CACHE.mkdir(parents=True, exist_ok=True)
    out = CACHE / f"{path.stem}-{int(path.stat().st_mtime)}.png"
    if not out.exists():
        with Image.open(path) as im:
            im.convert("RGB").quantize(colors=96, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE).save(out, optimize=True)
    return out


# ---------------------------------------------------------------- document model
class Builder:
    def __init__(self, toc_pages: dict[str, int] | None = None):
        self.doc = Document()
        self.toc_pages = toc_pages or {}
        self.figure = 0
        self.headings: list[tuple[int, str]] = []
        self.missing: list[str] = []
        self.used: set[str] = set()
        self._styles()
        self._page()

    def _styles(self) -> None:
        styles = self.doc.styles
        normal = styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        normal.paragraph_format.space_after = Pt(5)
        normal.paragraph_format.line_spacing = 1.1
        for name, size, before, after in (("Heading 1", 18, 18, 8), ("Heading 2", 14, 14, 6), ("Heading 3", 11.5, 10, 4)):
            style = styles[name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = GREEN
            for attr in ("eastAsia", "ascii", "hAnsi"):
                style.element.rPr.rFonts.set(qn(f"w:{attr}"), "Calibri")
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
        cap = styles["Caption"]
        cap.font.name = "Calibri"
        cap.font.size = Pt(9)
        cap.font.italic = True
        cap.font.bold = False
        cap.font.color.rgb = GREY
        cap.paragraph_format.space_after = Pt(9)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _page(self) -> None:
        section = self.doc.sections[0]
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = section.right_margin = Inches(1.0)
        section.top_margin = Inches(0.95)
        section.bottom_margin = Inches(0.85)
        section.different_first_page_header_footer = True
        header = section.header.paragraphs[0]
        header.text = "USLS Graduate School Lifecycle Portal  |  Demonstration Walkthrough"
        header.runs[0].font.size = Pt(8.5)
        header.runs[0].font.color.rgb = GREY
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Page ")
        add_field(footer, "PAGE")
        footer.add_run(" of ")
        add_field(footer, "NUMPAGES", "1")
        for run in footer.runs:
            run.font.size = Pt(8.5)
            run.font.color.rgb = GREY

    # primitives --------------------------------------------------------
    def heading(self, text: str, level: int):
        self.headings.append((level, text))
        return self.doc.add_heading(text, level=level)

    def para(self, text: str = "", *, italic=False, size=None, color=None, align=None, after=None, bold=False):
        p = self.doc.add_paragraph()
        self._runs(p, text, italic=italic, size=size, color=color, bold=bold)
        if align is not None:
            p.alignment = align
        if after is not None:
            p.paragraph_format.space_after = Pt(after)
        return p

    def _runs(self, p, text: str, italic=False, size=None, color=None, bold=False) -> None:
        """``**bold**`` spans are supported in text."""
        for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if not chunk:
                continue
            r = p.add_run(chunk)
            r.bold = bool(i % 2) or bold
            r.italic = italic
            if size:
                r.font.size = Pt(size)
            if color is not None:
                r.font.color.rgb = color

    def label(self, text: str):
        p = self.para(text, size=11.5, color=GREEN, after=2, bold=True)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(6)
        return p

    def bullet(self, text: str):
        p = self.doc.add_paragraph(style="List Bullet")
        self._runs(p, text)
        p.paragraph_format.space_after = Pt(2)
        return p

    def note(self, text: str, label: str = "Note"):
        table = self.doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.rows[0].cells[0]
        shade(cell, "EAF5EE")
        p = cell.paragraphs[0]
        r = p.add_run(f"{label}:  ")
        r.bold = True
        r.font.color.rgb = GREEN
        self._runs(p, text)
        p.paragraph_format.space_after = Pt(0)
        set_cell_margins(table, 70, 70, 120, 120)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def table(self, header: list[str], rows: list[list[str]], widths: list[float] | None = None, small=False):
        table = self.doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, text in enumerate(header):
            cell = table.rows[0].cells[i]
            shade(cell, "0F6B3E")
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(9.5 if small else 10)
            p.paragraph_format.space_after = Pt(0)
        repeat_header(table.rows[0])
        for row in rows:
            cells = table.add_row().cells
            for i, text in enumerate(row):
                p = cells[i].paragraphs[0]
                self._runs(p, str(text), size=9.5 if small else 10)
                p.paragraph_format.space_after = Pt(0)
            no_split(table.rows[-1])
        set_cell_margins(table)
        if widths:
            table.autofit = False
            grid = table._tbl.tblGrid
            for i, col in enumerate(grid.findall(qn("w:gridCol"))):
                col.set(qn("w:w"), str(int(widths[i] * 1440)))
            for row in table.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return table

    def image(self, shot_id: str, caption: str, width: float | None = None):
        width = width or SHOT_WIDTH
        path = SHOTS / f"{shot_id}.png"
        if not path.exists():
            self.missing.append(shot_id)
            self.para(f"[missing screenshot: {shot_id}]", italic=True, color=GREY)
            return None
        self.figure += 1
        self.used.add(shot_id)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        import struct

        with open(path, "rb") as fh:  # PNG header: width and height are big-endian ints at byte 16
            head = fh.read(24)
        w_px, h_px = struct.unpack(">II", head[16:24])
        w_in = width
        if w_in * h_px / w_px > 7.4:
            w_in = 7.4 * w_px / h_px
        p.add_run().add_picture(str(compact(path)), width=Inches(w_in))
        return self.doc.add_paragraph(f"Figure {self.figure}. {caption}", style="Caption")

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # sections ----------------------------------------------------------
    def cover(self, front: dict) -> None:
        for _ in range(4):
            self.para("", after=18)
        cover = front["cover"]
        self.para(cover["institution"], size=14, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER, after=2, bold=True)
        self.para(cover["title"], size=30, color=GREEN, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, bold=True)
        self.para(cover["subtitle"], size=17, align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
        self.para(cover["tagline"], size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=30)
        self.note(cover["password_line"], label="Demo password")
        self.para(
            f"Generated {date.today():%d %B %Y} from the running application: every step was performed "
            "and every screenshot captured by script against the current build.",
            italic=True, size=9.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        self.page_break()

    def toc(self) -> None:
        self.para("Contents", size=18, color=GREEN, after=6, bold=True)
        p = self.doc.add_paragraph()
        for child in (fld_char("begin"), instr_text('TOC \\o "1-2" \\h \\z \\u'), fld_char("separate")):
            p.add_run()._r.append(child)
        self._toc_anchor = p

    def fill_toc(self, headings: list[tuple[int, str]]) -> None:
        """Insert the TOC result (entries, with page numbers when known) inside the field."""
        last = self._toc_anchor
        for level, text in [h for h in headings if h[0] <= 2]:
            entry = self.doc.add_paragraph()
            pf = entry.paragraph_format
            pf.left_indent = Inches(0.0 if level == 1 else 0.3)
            pf.space_after = Pt(1 if level == 2 else 2)
            pf.space_before = Pt(5 if level == 1 else 0)
            pf.tab_stops.add_tab_stop(Inches(6.5), alignment=2, leader=1)
            r = entry.add_run(text)
            r.bold = level == 1
            r.font.size = Pt(10.5 if level == 1 else 9.5)
            page = self.toc_pages.get(text)
            entry.add_run("\t" + (str(page) if page else "")).font.size = Pt(9.5)
            last._p.addnext(entry._p)
            last = entry
        end_p = self.doc.add_paragraph()
        end_p.add_run()._r.append(fld_char("end"))
        end_p.paragraph_format.space_after = Pt(0)
        last._p.addnext(end_p._p)
        br = self.doc.add_paragraph()
        br.add_run().add_break(WD_BREAK.PAGE)
        end_p._p.addnext(br._p)


# ---------------------------------------------------------------- content -> document
def load_cases() -> list[dict]:
    cases = []
    for path in sorted((HERE / "cases").glob("*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        data["_file"] = path.name
        cases.append(data)
    return cases


def render_blocks(b: Builder, sections: list[dict]) -> None:
    for sec in sections:
        b.heading(sec["heading"], 1)
        for block in sec["blocks"]:
            kind = block["type"]
            if kind == "p":
                b.para(block["text"])
            elif kind == "label":
                b.label(block["text"])
            elif kind == "bullets":
                for item in block["items"]:
                    b.bullet(item)
            elif kind == "note":
                b.note(block["text"], block.get("label", "Note"))
            elif kind == "table":
                b.table(block["header"], block["rows"], block.get("widths"), small=block.get("small", False))
            elif kind == "image":
                b.image(block["id"], block["caption"])


def render_case(b: Builder, number: int, case: dict) -> None:
    suffix = "  ·  reference only" if case.get("kind") == "reference" else ""
    b.heading(f"Case {number} — {case['title']}{suffix}", 2)
    bits = [x for x in (f"Time: {case['duration']}" if case.get("duration") else "", case.get("owner", "")) if x]
    if bits:
        b.para("  ·  ".join(bits), italic=True, size=9.5, color=GREY, after=3)
    b.label("Story")
    for text in case["story"]:
        b.para(text)
    if case.get("accounts"):
        b.label("Accounts used")
        b.table(["Role", "Account", "Why"], [[a[0], a[1], a[2] if len(a) > 2 else ""] for a in case["accounts"]], [1.7, 2.6, 2.2], small=True)
    if case.get("data"):
        b.label("Data to use")
        b.table(["Field", "Value / instruction"], case["data"], [1.9, 4.6], small=True)
    step_no = 0
    for part in case["parts"]:
        b.label(part["title"])
        if part.get("intro"):
            b.para(part["intro"])
        for step in part["steps"]:
            if isinstance(step, str):
                step = {"text": step}
            step_no += 1
            sp = b.doc.add_paragraph()
            sp.paragraph_format.left_indent = Inches(0.62)
            sp.paragraph_format.first_line_indent = Inches(-0.62)
            sp.paragraph_format.space_after = Pt(4)
            r = sp.add_run(f"Step {step_no}.  ")
            r.bold = True
            r.font.color.rgb = GREEN
            b._runs(sp, step["text"])
            if step.get("shots"):
                sp.paragraph_format.keep_with_next = True
            for shot in step.get("shots", []):
                b.image(shot["id"], shot["caption"], shot.get("width"))
    for text in case.get("notes", []):
        b.note(text)


def _substitute(value, numbers: dict[str, int]):
    if isinstance(value, str):
        return re.sub(r"\{case:([a-z0-9_\-]+)\}", lambda m: f"Case {numbers[m.group(1).replace('_', '-')]}" if m.group(1).replace("_", "-") in numbers else m.group(0), value)
    if isinstance(value, list):
        return [_substitute(v, numbers) for v in value]
    if isinstance(value, dict):
        return {k: _substitute(v, numbers) for k, v in value.items()}
    return value


def ordered_cases() -> tuple[list[dict], dict[str, int]]:
    cases = load_cases()
    ordered = [c for c in cases if c.get("kind", "live") == "live"] + [c for c in cases if c.get("kind") == "reference"]
    numbers = {c["slug"].replace("_", "-"): n for n, c in enumerate(ordered, 1) if c.get("slug")}
    return [_substitute(c, numbers) for c in ordered], numbers


def build(toc_pages: dict[str, int] | None, out: Path) -> Builder:
    cases, numbers = ordered_cases()
    front = _substitute(json.loads((HERE / "front.json").read_text(encoding="utf-8")), numbers)
    b = Builder(toc_pages)
    b.cover(front)
    b.toc()
    render_blocks(b, front["before_cases"])
    b.heading("6.  Demonstration Running Order", 1)
    b.para(front["running_order_intro"])
    rows = []
    for n, c in enumerate(cases, 1):
        rows.append([f"Case {n}", c["title"] + ("  ·  reference only" if c.get("kind") == "reference" else ""), c.get("owner", ""), c.get("duration", "")])
    b.table(["Case", "What is demonstrated", "Who", "Time"], rows, [0.7, 3.3, 1.8, 0.7], small=True)
    for text in front.get("running_order_notes", []):
        b.para(text)
    b.heading("7.  Demonstration Cases", 1)
    b.para(front["cases_intro"])
    for n, c in enumerate(cases, 1):
        render_case(b, n, c)
    render_blocks(b, front["after_cases"])
    b.para("")
    b.para("End of Demonstration Walkthrough", italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
    b.fill_toc(b.headings)
    leftovers = sorted(set(re.findall(r"\{case:[^}]*\}", json.dumps([front, cases]))))
    if leftovers:
        print("UNRESOLVED CASE REFERENCES:", ", ".join(leftovers))
    out.parent.mkdir(parents=True, exist_ok=True)
    b.doc.save(out)
    return b


def render_pdf(docx_path: Path, outdir: Path) -> Path:
    outdir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        [str(SOFFICE), "-env:UserInstallation=file:///E:/Temp/claude/lo-profile-wt", "--headless", "--convert-to", "pdf",
         "--outdir", str(outdir), str(docx_path)],
        check=True, capture_output=True, timeout=900,
    )
    return outdir / (docx_path.stem + ".pdf")


def pdf_page_numbers(pdf: Path, headings: list[tuple[int, str]]) -> tuple[dict[str, int], int]:
    import fitz  # PyMuPDF, available in the project venv

    doc = fitz.open(pdf)
    texts = [re.sub(r"\s+", " ", doc[i].get_text()) for i in range(len(doc))]
    total = len(doc)
    doc.close()
    # the contents table lists every heading too, so start after the last contents page
    start = 1
    for i, t in enumerate(texts[:12]):
        if "Demonstration Running Order" in t and "Case 1" in t and i > 0:
            pass
    toc_end = 0
    for i, t in enumerate(texts[:14]):
        if "About the System" in t and t.count("....") < 3 and i > 0:
            toc_end = i
            break
    pages: dict[str, int] = {}
    cursor = toc_end
    for level, text in headings:
        key = re.sub(r"\s+", " ", text)
        for i in range(cursor, total):
            if key in texts[i]:
                pages[text] = i + 1
                cursor = i
                break
    return pages, total


def main() -> int:
    global SHOT_WIDTH
    parser = argparse.ArgumentParser()
    parser.add_argument("--no-render", action="store_true", help="skip the LibreOffice pass (contents table without page numbers)")
    parser.add_argument("--out", default=str(OUT))
    parser.add_argument("--shot-width", type=float, default=SHOT_WIDTH, help="screenshot width in inches (page text width is 6.5)")
    args = parser.parse_args()
    SHOT_WIDTH = args.shot_width
    out = Path(args.out)
    b = build(None, out)
    if b.missing:
        print("MISSING SCREENSHOTS:", ", ".join(sorted(set(b.missing))))
    if not args.no_render and SOFFICE.exists():
        renders = out.parent.parent.parent / "outputs" / "walkthrough_render" if False else Path("E:/Temp/claude/walkthrough/render")
        pdf = render_pdf(out, renders)
        pages, total = pdf_page_numbers(pdf, b.headings)
        build(pages, out)
        pdf = render_pdf(out, renders)
        print(f"rendered pages: {pdf_page_numbers(pdf, b.headings)[1]}  (pdf: {pdf})")
    print("wrote", out)
    # keep a copy of the screenshots next to the document (used later to refresh the proposal's screenshot appendix)
    import shutil

    dest = REPO / "Documents" / "CAPSTONE_ONLY" / "walkthrough_screenshots"
    dest.mkdir(parents=True, exist_ok=True)
    count = 0
    for old in dest.glob("*.png"):
        if old.stem not in b.used:
            old.unlink()  # no longer referenced by the document
    for png in sorted(SHOTS.glob("*.png")):
        if png.stem in b.used:
            try:  # 256-colour copy (UI screenshots lose nothing visible and shrink about 3x); originals stay in shots/
                from PIL import Image

                with Image.open(png) as im:
                    im.convert("RGB").quantize(colors=256, method=Image.Quantize.MEDIANCUT, dither=Image.Dither.NONE).save(dest / png.name, optimize=True)
            except ImportError:
                shutil.copy2(png, dest / png.name)
            count += 1
    print(f"copied {count} screenshots to {dest}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
