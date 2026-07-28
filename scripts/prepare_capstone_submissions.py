from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


WORKSPACE = Path(r"E:\Github_Projects\CAPSTONE-USLS-MVP")
SOURCE_DIR = WORKSPACE / "Documents" / "CAPSTONE_ONLY"
SUBMISSIONS_DIR = SOURCE_DIR / "Submissions"
CAPSTONE_DOCUMENT_DIR = SUBMISSIONS_DIR / "CAPSTONE DOCUMENT"

GROUP_CODE = "CAP-2521-IT"

CAPSTONE_SOURCE = SOURCE_DIR / "Final Proposal Document - CAP-IT1.docx"
PROPOSAL_FORM_SOURCE = SOURCE_DIR / "Defense Form - CAPIT0 PROPOSAL.docx"
CAP1_FORM_SOURCE = SOURCE_DIR / "Defense Form - CAP1.docx"
WALKTHROUGH_SOURCE = SOURCE_DIR / "USLS GS Portal Demo Walkthrough.docx"

PROJECT_TITLE = (
    "DEVELOPING A GRADUATE STUDENT LIFECYCLE MONITORING AND ANALYTICS "
    "PLATFORM FOR THE UNIVERSITY OF ST. LA SALLE GRADUATE SCHOOL"
)
GROUP_MEMBERS = [
    "MENDOZA, NATALIE U.",
    "OSEÑA, KHLOE CASSANDRA M.",
    "SANSON, MIGUEL JOAQUIN A.",
    "TUMBAGA, MYRINE RENA PRINCESS B.",
]
ADVISER = "MALABANAN, OLIVER A."
LEAD_PANEL = "ARCILLA, MARY JANE"
PANEL_MEMBER = "TANGKEKO, MARIVIC"


def check_sources() -> None:
    missing = [
        path
        for path in (
            CAPSTONE_SOURCE,
            PROPOSAL_FORM_SOURCE,
            CAP1_FORM_SOURCE,
            WALKTHROUGH_SOURCE,
        )
        if not path.is_file()
    ]
    if missing:
        joined = "\n".join(str(path) for path in missing)
        raise FileNotFoundError(f"Missing required source documents:\n{joined}")


def body_paragraph_text(element) -> str:
    if element.tag != qn("w:p"):
        return ""
    return "".join(element.xpath(".//w:t/text()")).strip()


def locate_exact_heading(body_elements, heading: str) -> int:
    matches = [
        index
        for index, element in enumerate(body_elements)
        if body_paragraph_text(element) == heading
    ]
    if len(matches) != 1:
        raise ValueError(
            f"Expected one body-level heading {heading!r}; found {len(matches)}"
        )
    return matches[0]


def locate_appendix_starts(body_elements) -> list[tuple[str, int, str]]:
    starts: list[tuple[str, int, str]] = []
    pattern = re.compile(r"^Appendix ([A-Z])\.(?:\s|$)")
    for index, element in enumerate(body_elements):
        text = body_paragraph_text(element)
        match = pattern.match(text)
        if match:
            starts.append((match.group(1), index, text))

    expected_letters = [chr(code) for code in range(ord("A"), ord("Z") + 1)]
    found_letters = [letter for letter, _, _ in starts]
    if found_letters != expected_letters:
        raise ValueError(
            "Top-level appendix headings do not run from A through Z: "
            f"{found_letters}"
        )
    return starts


def prune_unused_main_document_images(document: Document) -> None:
    relationship_namespace = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    referenced_relationship_ids: set[str] = set()
    for element in document._element.iter():
        for attribute_name, value in element.attrib.items():
            if attribute_name.startswith(f"{{{relationship_namespace}}}"):
                referenced_relationship_ids.add(value)

    removable = [
        relationship_id
        for relationship_id, relationship in document.part.rels.items()
        if relationship.reltype == RT.IMAGE
        and relationship_id not in referenced_relationship_ids
    ]
    for relationship_id in removable:
        del document.part.rels[relationship_id]


def save_body_slice(
    source: Path,
    destination: Path,
    start_index: int,
    end_index: int,
    document_title: str,
) -> None:
    shutil.copy2(source, destination)
    document = Document(destination)
    body = document._element.body
    original_elements = list(body.iterchildren())

    for index, element in enumerate(original_elements):
        keep_section_properties = element.tag == qn("w:sectPr")
        keep_requested_range = start_index <= index < end_index
        if not keep_section_properties and not keep_requested_range:
            body.remove(element)

    document.core_properties.title = document_title
    document.core_properties.subject = f"{GROUP_CODE} CAP1 submission"
    prune_unused_main_document_images(document)
    document.save(destination)


def replace_paragraph_text(paragraph, text: str) -> None:
    runs = list(paragraph.runs)
    if not runs:
        paragraph.add_run(text)
        return

    runs[0].text = text
    for run in runs[1:]:
        paragraph._p.remove(run._r)


def fill_cap1_defense_form(destination: Path) -> None:
    shutil.copy2(CAP1_FORM_SOURCE, destination)
    document = Document(destination)
    table = document.tables[0]

    replace_paragraph_text(
        document.paragraphs[6],
        "\u2612 BSIT      \u2610 BSIS",
    )

    title_cell = table.cell(0, 0)
    replace_paragraph_text(title_cell.paragraphs[1], PROJECT_TITLE)

    schedule_cell = table.cell(1, 0)
    replace_paragraph_text(
        schedule_cell.paragraphs[1],
        "Pending official CAP1 defense schedule",
    )

    group_cell = table.cell(4, 0)
    replace_paragraph_text(group_cell.paragraphs[0], "\n".join(GROUP_MEMBERS))

    adviser_cell = table.cell(6, 0)
    replace_paragraph_text(adviser_cell.paragraphs[0], f"Adviser:      {ADVISER}")

    lead_panel_cell = table.cell(7, 0)
    replace_paragraph_text(
        lead_panel_cell.paragraphs[0],
        f"Lead Panel:  {LEAD_PANEL}",
    )

    panel_member_cell = table.cell(9, 0)
    replace_paragraph_text(
        panel_member_cell.paragraphs[0],
        f"Panel Member:  {PANEL_MEMBER}",
    )

    document.core_properties.title = f"{GROUP_CODE} CAP1 Defense Form"
    document.core_properties.subject = (
        "Known project, program, group-member, adviser, and panel information filled; "
        "official defense outcomes remain pending."
    )
    document.save(destination)


def build_submission_files() -> None:
    check_sources()
    SUBMISSIONS_DIR.mkdir(exist_ok=True)
    CAPSTONE_DOCUMENT_DIR.mkdir(exist_ok=True)

    shutil.copy2(
        CAPSTONE_SOURCE,
        CAPSTONE_DOCUMENT_DIR / f"{GROUP_CODE}-CAP1DOC.docx",
    )
    shutil.copy2(
        PROPOSAL_FORM_SOURCE,
        SUBMISSIONS_DIR / f"{GROUP_CODE}-PROPFORM.docx",
    )
    shutil.copy2(
        WALKTHROUGH_SOURCE,
        CAPSTONE_DOCUMENT_DIR / f"{GROUP_CODE}-WALKTHROUGH.docx",
    )
    fill_cap1_defense_form(SUBMISSIONS_DIR / f"{GROUP_CODE}-CAP1DEF.docx")

    source_document = Document(CAPSTONE_SOURCE)
    body_elements = list(source_document._element.body.iterchildren())

    chapter_headings = [
        "Chapter 1 – Project Background",
        "Chapter 2 – Review of Related Literature",
        "Chapter 3 – Methodology",
        "CHAPTER 4 – The Existing System",
        "CHAPTER 5 – The Proposed System",
        "CHAPTER 6 – System Testing",
    ]
    chapter_starts = [
        locate_exact_heading(body_elements, heading) for heading in chapter_headings
    ]
    references_start = locate_exact_heading(body_elements, "REFERENCES")
    chapter_ends = chapter_starts[1:] + [references_start]

    for chapter_number, (heading, start, end) in enumerate(
        zip(chapter_headings, chapter_starts, chapter_ends),
        start=1,
    ):
        destination = (
            CAPSTONE_DOCUMENT_DIR
            / f"{GROUP_CODE}-CAP1CH{chapter_number:02d}.docx"
        )
        save_body_slice(
            CAPSTONE_SOURCE,
            destination,
            start,
            end,
            f"{GROUP_CODE} CAP1 Chapter {chapter_number}: {heading}",
        )

    appendix_starts = locate_appendix_starts(body_elements)
    section_properties_index = next(
        (
            index
            for index, element in enumerate(body_elements)
            if element.tag == qn("w:sectPr")
        ),
        len(body_elements),
    )

    for appendix_number, (letter, start, heading) in enumerate(
        appendix_starts,
        start=1,
    ):
        end = (
            appendix_starts[appendix_number][1]
            if appendix_number < len(appendix_starts)
            else section_properties_index
        )
        destination = (
            CAPSTONE_DOCUMENT_DIR
            / f"{GROUP_CODE}-CAP1APP{appendix_number:02d}.docx"
        )
        save_body_slice(
            CAPSTONE_SOURCE,
            destination,
            start,
            end,
            f"{GROUP_CODE} CAP1 Appendix {letter}: {heading}",
        )


if __name__ == "__main__":
    build_submission_files()
